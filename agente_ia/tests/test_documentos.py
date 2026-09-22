"""
tests/test_documentos.py
Pruebas de `core/documentos.py` — REQ-040, producir documentos de verdad.

Lo que protege esta suite, además de que los archivos salgan bien formados:

- **Un documento con relleno no se produce.** "lorem ipsum", "[insertar aquí]" o "por
  definir" se rechazan: un archivo con huecos parece trabajo hecho y no lo es.
- **Nunca se pisa un archivo del usuario.** Si `informe.docx` existe, sale `informe (2).docx`.
- **Los archivos van solo a donde corresponde.** Ni a `AppData`, ni fuera de la carpeta
  personal salvo un espacio de trabajo habilitado.
- **Las plantillas se respetan.** Membrete, pies y estilos se conservan; los marcadores se
  rellenan conservando el formato, también dentro de tablas.

Office no se toca en las pruebas: la conversión a PDF y los temas van simulados.
"""

import os
from pathlib import Path

import pytest

from core import documentos


@pytest.fixture
def carpeta(tmp_path, monkeypatch):
    """Una carpeta de salida permitida, fuera de la carpeta personal del que corre la suite."""
    monkeypatch.setattr(documentos, "_bases_permitidas", lambda: [tmp_path])
    # `tmp_path` vive bajo AppData, que en producción se rechaza con razón.
    monkeypatch.setattr(documentos, "CARPETAS_VEDADAS", frozenset())
    monkeypatch.setattr(documentos, "_con_office", lambda *a, **k: False)
    monkeypatch.setattr(documentos, "_aplicar_tema", lambda *a, **k: False)
    return str(tmp_path)


_INFORME = {
    "titulo": "Informe de campo",
    "autor": "Johan",
    "bloques": [
        {"tipo": "titulo", "nivel": 1, "texto": "Visita al lote 14"},
        {"tipo": "parrafo", "texto": ("Se encontró un avance del 60 % en la fertilización "
                                      "programada, con 1.200 palmas atendidas y 800 pendientes "
                                      "para la semana siguiente según el cronograma.")},
        {"tipo": "lista", "items": ["Aplicar segunda dosis", ["el 20 de septiembre"]]},
        {"tipo": "tabla", "encabezados": ["Lote", "Palmas"], "filas": [["14", "1.200"], ["15", "800"]]},
        {"tipo": "cita", "texto": "Lo que no se mide no se mejora."},
    ],
}


# ---------------------------------------------------------------- relleno y contenido

def test_un_documento_con_lorem_ipsum_no_se_produce(carpeta):
    with pytest.raises(documentos.DocumentoRechazado, match="relleno"):
        documentos.crear_documento("x.docx", {"bloques": [
            {"tipo": "parrafo", "texto": "Lorem ipsum dolor sit amet " * 10}]}, carpeta=carpeta)


def test_un_hueco_por_completar_tampoco(carpeta):
    with pytest.raises(documentos.DocumentoRechazado, match="relleno"):
        documentos.crear_hoja("x.xlsx", {"hojas": [{"encabezados": ["A"],
                                                    "filas": [["[insertar dato]"]]}]}, carpeta=carpeta)


def test_un_titulo_sin_nada_debajo_no_es_un_documento(carpeta):
    with pytest.raises(documentos.DocumentoRechazado, match="contenido suficiente"):
        documentos.crear_documento("x.docx", {"titulo": "Informe", "bloques": []}, carpeta=carpeta)


def test_una_presentacion_de_puros_titulos_no_es_una_presentacion(carpeta):
    with pytest.raises(documentos.DocumentoRechazado, match="contenido suficiente"):
        documentos.crear_presentacion("x.pptx", {"diapositivas": [{"titulo": "Uno"}, {"titulo": "Dos"}]},
                                      carpeta=carpeta)


def test_el_contenido_puede_venir_como_json_en_texto(carpeta):
    """Los modelos mandan a veces el JSON como cadena: se acepta en vez de gastar una vuelta."""
    import json

    r = documentos.crear_documento("nota.md", json.dumps(_INFORME), carpeta=carpeta)
    assert Path(r.ruta).read_text(encoding="utf-8").startswith("# Informe de campo")


def test_el_contenido_puede_venir_en_markdown(carpeta):
    md = ("# Acta\n\n## Asistentes\n\n- Johan\n- María\n\n## Acuerdos\n\n"
          "1. Revisar el drenaje del bloque C antes del viernes de esta semana.\n"
          "2. Enviar la cotización de fertilizante al proveedor con copia a compras.\n\n"
          "| Tema | Responsable |\n|---|---|\n| Drenaje | Johan |\n")
    r = documentos.crear_documento("acta.docx", md, carpeta=carpeta)

    from docx import Document

    doc = Document(r.ruta)
    estilos = [p.style.name for p in doc.paragraphs if p.text.strip()]
    assert "Heading 2" in estilos or "Heading 1" in estilos
    assert any("List" in e for e in estilos)
    assert len(doc.tables) == 1


# ---------------------------------------------------------------------- destino

def test_no_se_pisa_un_archivo_existente(carpeta):
    r1 = documentos.crear_documento("informe.md", _INFORME, carpeta=carpeta)
    r2 = documentos.crear_documento("informe.md", _INFORME, carpeta=carpeta)

    assert Path(r1.ruta).name == "informe.md"
    assert Path(r2.ruta).name == "informe (2).md"
    assert Path(r1.ruta).exists() and Path(r2.ruta).exists()


def test_no_se_escribe_fuera_de_las_carpetas_permitidas(tmp_path, monkeypatch):
    monkeypatch.setattr(documentos, "_bases_permitidas", lambda: [tmp_path / "permitida"])
    monkeypatch.setattr(documentos, "CARPETAS_VEDADAS", frozenset())
    with pytest.raises(documentos.DocumentoRechazado, match="carpeta personal"):
        documentos.crear_documento("x.md", _INFORME, carpeta=str(tmp_path / "otra"))


def test_no_se_escribe_en_carpetas_de_configuracion(tmp_path, monkeypatch):
    monkeypatch.setattr(documentos, "_bases_permitidas", lambda: [tmp_path])
    with pytest.raises(documentos.DocumentoRechazado, match="configuración"):
        documentos.crear_documento("x.md", _INFORME, carpeta=str(tmp_path / ".ssh"))


def test_el_nombre_se_limpia_de_caracteres_prohibidos(carpeta):
    r = documentos.crear_documento('in:for<me>.md', _INFORME, carpeta=carpeta)
    assert Path(r.ruta).name == "in_for_me_.md"


# -------------------------------------------------------------------------- Word

def test_el_word_sale_con_titulo_encabezados_lista_tabla_y_cita(carpeta):
    from docx import Document

    r = documentos.crear_documento("informe.docx", _INFORME, carpeta=carpeta)
    doc = Document(r.ruta)

    textos = [p.text for p in doc.paragraphs if p.text.strip()]
    estilos = {p.style.name for p in doc.paragraphs if p.text.strip()}
    assert textos[0] == "Informe de campo"
    assert "Heading 1" in estilos
    assert "List Bullet" in estilos
    assert any("Quote" in e for e in estilos)
    assert doc.tables[0].rows[0].cells[0].text == "Lote"
    assert doc.tables[0].rows[1].cells[1].text == "1.200"


def test_una_imagen_que_no_existe_no_tumba_el_documento(carpeta):
    spec = {**_INFORME, "bloques": _INFORME["bloques"] + [{"tipo": "imagen", "ruta": "C:/no/hay.png"}]}
    r = documentos.crear_documento("con_imagen.docx", spec, carpeta=carpeta)

    assert Path(r.ruta).exists()
    assert any("No encontré la imagen" in a for a in r.avisos)


def test_pdf_sin_office_se_maqueta_igual(carpeta):
    """Sin Word, un .pdf se produce con PyMuPDF: peor que Word, mejor que nada."""
    r = documentos.crear_documento("informe.pdf", _INFORME, carpeta=carpeta)

    import pymupdf

    pdf = pymupdf.open(r.ruta)
    assert pdf.page_count >= 1
    assert "Visita al lote 14" in pdf[0].get_text()
    assert "sin Office" in r.detalle


def test_pdf_con_office_se_convierte_desde_word(carpeta, monkeypatch):
    def _convertir(aplicacion, origen, destino):
        assert aplicacion == "Word" and origen.endswith(".docx")
        Path(destino).write_bytes(b"%PDF-1.4 simulado")
        return True

    monkeypatch.setattr(documentos, "_con_office", _convertir)
    r = documentos.crear_documento("informe.pdf", _INFORME, carpeta=carpeta)

    assert Path(r.ruta).read_bytes().startswith(b"%PDF")
    assert "convertido con Word" in r.detalle


def test_html_md_y_txt_salen_del_mismo_contenido(carpeta):
    for ext in ("html", "md", "txt"):
        r = documentos.crear_documento(f"informe.{ext}", _INFORME, carpeta=carpeta)
        texto = Path(r.ruta).read_text(encoding="utf-8")
        assert "Visita al lote 14" in texto
        assert "1.200" in texto


# --------------------------------------------------------------- Word con plantilla

@pytest.fixture
def plantilla_docx(tmp_path):
    from docx import Document
    from docx.shared import RGBColor

    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "AGRO ORIENTAL S.A.S."
    doc.styles["Heading 1"].font.color.rgb = RGBColor(0x1F, 0x6E, 0x43)
    doc.add_paragraph("Cotización {{numero}} para {{cliente}}", style="Heading 1")
    t = doc.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "{{concepto}}"
    t.rows[0].cells[1].text = "{{valor}}"
    ruta = tmp_path / "plantilla.docx"
    doc.save(str(ruta))
    return str(ruta)


def test_rellenar_una_plantilla_conserva_membrete_y_rellena_las_tablas(carpeta, plantilla_docx):
    from docx import Document

    r = documentos.crear_documento(
        "cotizacion.docx", {"bloques": []}, plantilla=plantilla_docx,
        reemplazos={"{{numero}}": "2026-18", "{{cliente}}": "Palmar Ltda",
                    "{{concepto}}": "Urea", "{{valor}}": "$ 1.200.000"},
        carpeta=carpeta)
    doc = Document(r.ruta)

    assert doc.sections[0].header.paragraphs[0].text == "AGRO ORIENTAL S.A.S."
    assert doc.paragraphs[0].text == "Cotización 2026-18 para Palmar Ltda"
    assert doc.paragraphs[0].style.name == "Heading 1"          # el formato se conserva
    assert [c.text for c in doc.tables[0].rows[0].cells] == ["Urea", "$ 1.200.000"]
    assert "4 marcadores" in r.detalle


def test_contenido_nuevo_sobre_la_plantilla_hereda_estilos_y_quita_el_cuerpo_viejo(
        carpeta, plantilla_docx):
    from docx import Document

    r = documentos.crear_documento("informe.docx", _INFORME, plantilla=plantilla_docx, carpeta=carpeta)
    doc = Document(r.ruta)

    assert doc.sections[0].header.paragraphs[0].text == "AGRO ORIENTAL S.A.S."
    assert not any("{{" in p.text for p in doc.paragraphs)   # el cuerpo viejo se fue
    h1 = [p for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert h1 and str(h1[0].style.font.color.rgb) == "1F6E43"  # y el estilo de la plantilla queda


def test_marcadores_que_no_estan_en_la_plantilla_se_rechaza_y_se_explica(carpeta, plantilla_docx):
    # REQ-065: antes solo se avisaba y el archivo salía con {{cliente}} impreso.
    with pytest.raises(documentos.DocumentoRechazado) as e:
        documentos.crear_documento("x.docx", {"bloques": []}, plantilla=plantilla_docx,
                                   reemplazos={"{{inexistente}}": "nada"}, carpeta=carpeta)
    assert "sin rellenar" in str(e.value) and "{{cliente}}" in str(e.value)
    assert "ninguno de los marcadores" in str(e.value)
    assert not os.path.exists(os.path.join(carpeta, "x.docx"))


def test_una_plantilla_que_no_existe_lo_dice(carpeta):
    with pytest.raises(documentos.DocumentoRechazado, match="No encuentro la plantilla"):
        documentos.crear_documento("x.docx", _INFORME, plantilla="C:/no/hay.docx", carpeta=carpeta)


def test_inspeccionar_una_plantilla_cuenta_sus_marcadores_y_su_membrete(plantilla_docx):
    salida = documentos.inspeccionar(plantilla_docx)

    assert "AGRO ORIENTAL" in salida
    assert "{{cliente}}" in salida and "{{valor}}" in salida
    assert "Heading 1" in salida


# -------------------------------------------------------------------------- Excel

_VENTAS = {"hojas": [{
    "nombre": "Ventas", "encabezados": ["Mes", "Ventas", "Costos", "Margen"],
    "filas": [["Enero", 1200000, "800.000", "=B2-C2"], ["Febrero", "1.350.000", 870000, "=B3-C3"]],
    "totales": True, "formatos": {"B": "#,##0"},
    "grafico": {"tipo": "barras", "titulo": "Ventas", "columnas": ["Ventas"], "categorias": "Mes"},
}]}


def test_el_excel_sale_con_numeros_formulas_totales_y_grafico(carpeta):
    from openpyxl import load_workbook

    r = documentos.crear_hoja("ventas.xlsx", _VENTAS, carpeta=carpeta)
    ws = load_workbook(r.ruta)["Ventas"]

    assert ws["C2"].value == 800000                 # "800.000" a la colombiana es un número
    assert ws["B3"].value == 1350000
    assert ws["D2"].value == "=B2-C2"               # las fórmulas se quedan como fórmulas
    assert ws["A4"].value == "Total"
    assert ws["B4"].value == "=SUM(B2:B3)"
    assert ws["B2"].number_format == "#,##0"
    assert ws.freeze_panes == "A2"
    assert len(ws._charts) == 1


def test_los_numeros_se_leen_a_la_colombiana():
    v = documentos._valor_de_celda
    assert v("870.000") == 870000
    assert v("1.250.000,50") == 1250000.5
    assert v("3,4") == 3.4
    assert v("12.5") == 12.5                        # uno o dos decimales tras el punto
    assert v("2026") == 2026
    assert v("=SUM(A1:A3)") == "=SUM(A1:A3)"
    assert v("Enero") == "Enero"


def test_una_lista_de_registros_se_vuelve_hoja(carpeta):
    from openpyxl import load_workbook

    registros = [{"cliente": "Agro SAS", "cupo": 45000000}, {"cliente": "Palmar", "cupo": 32000000}]
    r = documentos.crear_hoja("clientes.xlsx", registros, carpeta=carpeta)
    ws = load_workbook(r.ruta).active

    assert [c.value for c in ws[1]] == ["cliente", "cupo"]
    assert ws["B3"].value == 32000000


def test_csv_sale_con_punto_y_coma_para_excel_en_espanol(carpeta):
    r = documentos.crear_hoja("ventas.csv", _VENTAS, carpeta=carpeta)
    lineas = Path(r.ruta).read_text(encoding="utf-8-sig").splitlines()
    assert lineas[0] == "Mes;Ventas;Costos;Margen"
    assert lineas[1].startswith("Enero;1200000")


def test_excel_sobre_plantilla_escribe_debajo_del_encabezado_que_ya_tiene(carpeta, tmp_path):
    from openpyxl import Workbook, load_workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario"
    ws.append(["Producto", "Existencias"])
    plantilla = tmp_path / "plantilla.xlsx"
    wb.save(str(plantilla))

    r = documentos.crear_hoja("inventario.xlsx", {"hojas": [{
        "nombre": "Inventario", "encabezados": ["Producto", "Existencias"],
        "filas": [["Urea", 120], ["KCl", 80]]}]}, plantilla=str(plantilla), carpeta=carpeta)
    ws = load_workbook(r.ruta)["Inventario"]

    assert [c.value for c in ws[1]] == ["Producto", "Existencias"]   # el de la plantilla
    assert ws["A2"].value == "Urea" and ws["B3"].value == 80
    assert ws.max_row == 3                                             # sin encabezado repetido


def test_una_hoja_vacia_se_rechaza(carpeta):
    with pytest.raises(documentos.DocumentoRechazado, match="vacía"):
        documentos.crear_hoja("x.xlsx", {"hojas": [{"nombre": "A", "filas": []}]}, carpeta=carpeta)


# --------------------------------------------------------------------- PowerPoint

_PRESENTACION = {"titulo": "Resultados 2025", "subtitulo": "Junta", "diapositivas": [
    {"titulo": "Producción", "puntos": ["1.850.000 t", ["Oriental +15 %"]], "notas": "Riego."},
    {"titulo": "Indicadores", "tabla": {"encabezados": ["Ind", "2025"], "filas": [["Agua", "3,1"]]}},
    {"titulo": "Cierre", "texto": "Auditoría en marzo."},
]}


def test_la_presentacion_sale_con_portada_vinetas_tabla_y_notas(carpeta):
    from pptx import Presentation

    r = documentos.crear_presentacion("resultados.pptx", _PRESENTACION, carpeta=carpeta)
    prs = Presentation(r.ruta)

    titulos = [s.shapes.title.text for s in prs.slides]
    assert titulos == ["Resultados 2025", "Producción", "Indicadores", "Cierre"]
    cuerpo = [ph for ph in prs.slides[1].placeholders if ph.placeholder_format.idx == 1][0]
    niveles = [(p.text, p.level) for p in cuerpo.text_frame.paragraphs]
    assert niveles == [("1.850.000 t", 0), ("Oriental +15 %", 1)]
    assert prs.slides[1].notes_slide.notes_text_frame.text == "Riego."
    assert any(sh.has_table for sh in prs.slides[2].shapes)
    assert prs.slide_width > prs.slide_height * 1.5          # panorámica, no 4:3


def test_sin_powerpoint_el_tema_no_se_aplica_y_se_avisa(carpeta):
    r = documentos.crear_presentacion("x.pptx", _PRESENTACION, carpeta=carpeta)
    assert any("tema" in a.lower() for a in r.avisos)


def test_con_powerpoint_se_aplica_el_tema_pedido(carpeta, monkeypatch):
    aplicados = []
    monkeypatch.setattr(documentos, "_ruta_de_tema", lambda n: "C:/temas/Facet.thmx")
    monkeypatch.setattr(documentos, "_aplicar_tema", lambda ruta, tema: aplicados.append(tema) or True)

    r = documentos.crear_presentacion("x.pptx", _PRESENTACION, carpeta=carpeta, tema="Facet")

    assert aplicados == ["Facet"]
    assert "tema Facet" in r.detalle


def test_un_tema_que_no_existe_lo_dice_y_lista_los_que_hay(carpeta, monkeypatch):
    monkeypatch.setattr(documentos, "_ruta_de_tema", lambda n: None)
    monkeypatch.setattr(documentos, "temas_disponibles", lambda: ["Facet", "Ion"])

    r = documentos.crear_presentacion("x.pptx", _PRESENTACION, carpeta=carpeta, tema="Inventado")

    assert any("Inventado" in a and "Facet" in a for a in r.avisos)


def test_pptx_sobre_plantilla_rellena_y_conserva_el_tema(carpeta, tmp_path):
    from pptx import Presentation

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Propuesta para {{cliente}}"
    plantilla = tmp_path / "plantilla.pptx"
    prs.save(str(plantilla))

    r = documentos.crear_presentacion("propuesta.pptx", {"diapositivas": []},
                                      plantilla=str(plantilla),
                                      reemplazos={"{{cliente}}": "Palmar"}, carpeta=carpeta)

    assert Presentation(r.ruta).slides[0].shapes.title.text == "Propuesta para Palmar"
    assert "con el tema de plantilla.pptx" in r.detalle


def test_pdf_de_presentacion_sin_powerpoint_se_maqueta_y_lo_dice(carpeta):
    """Sin PowerPoint no hay tema, pero sí PDF: una página apaisada por diapositiva."""
    r = documentos.crear_presentacion("x.pdf", _PRESENTACION, carpeta=carpeta)

    import pymupdf

    pdf = pymupdf.open(r.ruta)
    assert pdf.page_count == 4                          # portada + 3
    assert "Producción" in pdf[1].get_text()
    assert pdf[0].rect.width > pdf[0].rect.height       # apaisado
    assert any("Sin PowerPoint" in a for a in r.avisos)


def test_pdf_de_hoja_sin_excel_se_maqueta_y_avisa_de_las_formulas(carpeta):
    r = documentos.crear_hoja("ventas.pdf", _VENTAS, carpeta=carpeta)

    import pymupdf

    pdf = pymupdf.open(r.ruta)
    texto = pdf[0].get_text()
    assert "Ventas" in texto and "Enero" in texto
    assert "1.200.000" in texto                          # los números, a la colombiana
    assert "2.550.000" in texto                          # =SUM(B2:B3), calculada sin Excel
    assert "400.000" in texto                            # =B2-C2
    assert not any("fórmula" in a for a in r.avisos)     # nada quedó sin calcular
    assert any("gráfico" in a for a in r.avisos)


def test_una_formula_que_no_se_sabe_calcular_se_muestra_escrita_y_se_avisa(carpeta):
    hoja = {"hojas": [{"nombre": "X", "encabezados": ["A", "B"],
                       "filas": [[1, "=VLOOKUP(A2,A:B,2)"], [2, 3]]}]}
    r = documentos.crear_hoja("x.pdf", hoja, carpeta=carpeta)

    import pymupdf

    assert "VLOOKUP" in pymupdf.open(r.ruta)[0].get_text()
    assert any("fórmula" in a for a in r.avisos)


def test_el_evaluador_de_formulas_no_es_eval():
    """Cualquier cosa fuera de la gramática devuelve None: nunca se ejecuta texto ajeno."""
    f = documentos._calcular_formula
    celdas = {"A1": 2, "A2": 3, "A3": "=A1*A2"}
    assert f("=A1+A2", celdas) == 5
    assert f("=A3", celdas) == 6                         # una fórmula que referencia otra
    assert f("=SUMA(A1:A2)*2", celdas) == 10
    assert f("=-A1", celdas) == -2
    assert f("=A1/0", celdas) is None
    assert f("=__import__('os')", celdas) is None
    assert f("=A1+texto", celdas) is None


def test_pdf_de_hoja_con_excel_se_convierte(carpeta, monkeypatch):
    def _convertir(aplicacion, origen, destino):
        assert aplicacion == "Excel" and origen.endswith(".xlsx")
        Path(destino).write_bytes(b"%PDF-1.4 simulado")
        return True

    monkeypatch.setattr(documentos, "_con_office", _convertir)
    r = documentos.crear_hoja("ventas.pdf", _VENTAS, carpeta=carpeta)

    assert Path(r.ruta).read_bytes().startswith(b"%PDF")
    assert "convertido con Excel" in r.detalle


# ------------------------------------------------------------------- otros archivos

def test_un_json_sale_con_sangria(carpeta):
    r = documentos.crear_archivo("datos.json", {"a": 1, "b": [1, 2]}, carpeta=carpeta)
    assert Path(r.ruta).read_text(encoding="utf-8") == '{\n  "a": 1,\n  "b": [\n    1,\n    2\n  ]\n}'


def test_cualquier_extension_de_texto(carpeta):
    r = documentos.crear_archivo("consulta.sql", "SELECT * FROM ventas WHERE anio = 2026;", carpeta=carpeta)
    assert Path(r.ruta).read_text(encoding="utf-8").startswith("SELECT")


def test_un_archivo_sin_extension_se_rechaza(carpeta):
    with pytest.raises(documentos.DocumentoRechazado, match="extensión"):
        documentos.crear_archivo("sinextension", "hola", carpeta=carpeta)


def test_convertir_a_pdf_un_archivo_que_no_es_de_office_lo_dice(carpeta, tmp_path):
    origen = tmp_path / "nota.txt"
    origen.write_text("hola", encoding="utf-8")
    with pytest.raises(documentos.DocumentoRechazado, match="solo Word, Excel y PowerPoint"):
        documentos.convertir_a_pdf(str(origen))


# --------------------------------------------------------------- las herramientas

def test_las_herramientas_estan_registradas_con_su_nivel():
    from agents.tool_registry import get_tool
    from core.security_manager import ChannelType, RiskLevel, security_manager

    for amarilla in ("document_create", "spreadsheet_create", "presentation_create",
                     "file_create", "document_to_pdf"):
        assert get_tool(amarilla) is not None, amarilla
        assert security_manager.classify_action(amarilla) == RiskLevel.YELLOW, amarilla
    assert security_manager.classify_action("document_inspect") == RiskLevel.GREEN
    # Leer una plantilla es leer un archivo del usuario: solo delante del computador.
    assert not security_manager.is_action_allowed("document_inspect", ChannelType.TELEGRAM)


def test_la_confirmacion_dice_que_archivo_se_va_a_crear_y_donde():
    from core.security_manager import format_details

    detalle = format_details("document_create", {"nombre": "renuncia.docx", "carpeta": "Escritorio",
                                                 "plantilla": "membrete.docx"})
    assert "renuncia.docx" in detalle and "Escritorio" in detalle and "membrete.docx" in detalle


def test_la_herramienta_devuelve_el_rechazo_en_palabras_del_modelo(monkeypatch, carpeta):
    from agents.tool_registry import get_tool

    salida = get_tool("document_create").invoke(
        {"nombre": "x.docx", "contenido": {"bloques": [{"tipo": "parrafo", "texto": "lorem ipsum " * 20}]},
         "carpeta": carpeta})

    assert "relleno" in salida


# =============================================================================== gráficos
#
# En PowerPoint el gráfico es NATIVO —el usuario lo puede editar—; en Word y en el PDF sin
# Office es una imagen dibujada con matplotlib. Ninguna de las dos formas se prueba por
# píxeles: se prueba que el gráfico exista, que las cifras salgan a la colombiana y que un
# gráfico sin datos no reviente el documento.

_GRAFICO = {"tipo": "grafico", "tipo_grafico": "barras", "titulo": "Producción por región",
            "categorias": ["Oriental", "Norte"],
            "series": [{"nombre": "2025", "valores": ["828.000", "551.000"]}], "pie": "Figura 1"}


def test_un_grafico_en_word_es_una_imagen_con_su_pie(carpeta):
    from docx import Document

    spec = {"titulo": "Producción", "bloques": [
        {"tipo": "parrafo", "texto": "La producción creció en las dos regiones durante el año "
                                     "gracias a las hectáreas nuevas con riego por goteo."},
        _GRAFICO]}
    r = documentos.crear_documento("produccion.docx", spec, carpeta=carpeta)
    doc = Document(r.ruta)

    assert len(doc.inline_shapes) == 1
    assert any(p.text == "Figura 1" for p in doc.paragraphs)


def test_un_grafico_sin_datos_no_tumba_el_word(carpeta):
    spec = {"titulo": "X", "bloques": [
        {"tipo": "parrafo", "texto": "Un párrafo suficientemente largo como para que el documento "
                                     "tenga contenido real y pase la validación de tamaño mínimo, "
                                     "que exige más de ciento veinte caracteres de prosa."},
        {"tipo": "grafico", "titulo": "Vacío"}]}
    r = documentos.crear_documento("x.docx", spec, carpeta=carpeta)

    assert Path(r.ruta).exists()
    assert any("no traía datos" in a for a in r.avisos)


def test_un_grafico_en_powerpoint_es_nativo(carpeta):
    from pptx import Presentation

    pres = {"titulo": "Resultados", "diapositivas": [
        {"titulo": "Por región", "grafico": _GRAFICO},
        {"titulo": "Cierre", "texto": "Gracias por su atención a todos los asistentes."}]}
    r = documentos.crear_presentacion("g.pptx", pres, carpeta=carpeta)
    prs = Presentation(r.ruta)

    graficos = [sh for sh in prs.slides[1].shapes if sh.has_chart]
    assert len(graficos) == 1
    chart = graficos[0].chart
    assert [c for c in chart.plots[0].categories] == ["Oriental", "Norte"]
    assert list(chart.plots[0].series[0].values) == [828000, 551000]     # a la colombiana


def test_una_torta_en_powerpoint_lleva_porcentajes(carpeta):
    from pptx import Presentation

    pres = {"titulo": "R", "diapositivas": [
        {"titulo": "Participación", "grafico": {"tipo_grafico": "torta", "categorias": ["A", "B"],
                                                "valores": [3, 1]}},
        {"titulo": "Cierre", "texto": "Gracias por su atención a todos los asistentes."}]}
    r = documentos.crear_presentacion("t.pptx", pres, carpeta=carpeta)
    chart = [sh for sh in Presentation(r.ruta).slides[1].shapes if sh.has_chart][0].chart

    assert chart.plots[0].has_data_labels
    assert chart.plots[0].data_labels.show_percentage


def test_el_grafico_como_imagen_se_dibuja(tmp_path):
    destino = str(tmp_path / "g.png")
    assert documentos._grafico_como_imagen(_GRAFICO, destino) is True
    assert Path(destino).stat().st_size > 1000


def test_un_grafico_sin_series_no_se_dibuja(tmp_path):
    assert documentos._grafico_como_imagen({"titulo": "nada"}, str(tmp_path / "g.png")) is False


# ============================================================ la región decide los números

def test_la_convencion_sigue_a_la_region(monkeypatch):
    import config_manager

    monkeypatch.setattr(config_manager, "get_region", lambda: "es-CO")
    assert documentos._valor_de_celda("870.000") == 870000
    assert documentos.convencion().separador_csv == ";"
    assert documentos.convencion().formatear(1250000.5, 1) == "1.250.000,5"

    monkeypatch.setattr(config_manager, "get_region", lambda: "en-US")
    assert documentos._valor_de_celda("870.000") == 870.0
    assert documentos._valor_de_celda("870,000") == 870000
    assert documentos.convencion().separador_csv == ","
    assert documentos.convencion().formatear(1250000.5, 1) == "1,250,000.5"

    monkeypatch.setattr(config_manager, "get_region", lambda: "es-MX")   # México escribe a la inglesa
    assert documentos._valor_de_celda("1,250,000") == 1250000


def test_una_region_desconocida_cae_en_su_idioma(monkeypatch):
    import config_manager

    monkeypatch.setattr(config_manager, "get_region", lambda: "es-XX")
    assert documentos.convencion().decimal == ","
    monkeypatch.setattr(config_manager, "get_region", lambda: "xx-YY")
    assert documentos.convencion().decimal == "."


def test_el_csv_usa_el_separador_de_la_region(carpeta, monkeypatch):
    import config_manager

    monkeypatch.setattr(config_manager, "get_region", lambda: "en-US")
    r = documentos.crear_hoja("v.csv", _VENTAS, carpeta=carpeta)
    assert Path(r.ruta).read_text(encoding="utf-8-sig").splitlines()[0] == "Mes,Ventas,Costos,Margen"


# ======================================================================= editar existente

def _word_base(carpeta):
    return documentos.crear_documento("base.docx", {"titulo": "Informe", "bloques": [
        {"tipo": "titulo", "nivel": 1, "texto": "1. Antecedentes"},
        {"tipo": "parrafo", "texto": "El proyecto arrancó en enero de 2025 con un presupuesto "
                                     "inicial de 320 millones de pesos aprobado por la junta."},
        {"tipo": "titulo", "nivel": 1, "texto": "2. Conclusiones"},
        {"tipo": "parrafo", "texto": "Se recomienda continuar con la segunda fase durante el "
                                     "primer trimestre de 2026, ajustando el cronograma."},
        {"tipo": "parrafo", "texto": "BORRADOR — no distribuir."}]}, carpeta=carpeta).ruta


def test_editar_word_reemplaza_agrega_donde_se_pide_y_quita(carpeta):
    from docx import Document

    r = documentos.editar_documento(_word_base(carpeta), {
        "reemplazos": {"320 millones": "345 millones"},
        "quitar": ["BORRADOR"],
        "agregar": [{"tipo": "parrafo", "texto": "En marzo se aprobó una adición de 25 millones "
                                                 "para el componente de riego, que explica la "
                                                 "diferencia con el presupuesto inicial."}],
        "despues_de": "arrancó en enero"})
    textos = [p.text for p in Document(r.ruta).paragraphs if p.text.strip()]

    assert "345 millones" in textos[2]
    assert textos[3].startswith("En marzo")                 # justo después del ancla
    assert textos[4] == "2. Conclusiones"
    assert not any("BORRADOR" in t for t in textos)
    assert "1 reemplazos" in r.detalle and "1 párrafos quitados" in r.detalle


def test_editar_deja_el_original_intacto_por_defecto(carpeta):
    from docx import Document

    original = _word_base(carpeta)
    r = documentos.editar_documento(original, {"reemplazos": {"320": "999"}})

    assert Path(r.ruta).name == "base (2).docx"
    assert any("320 millones" in p.text for p in Document(original).paragraphs)
    assert "quedó intacto" in r.detalle


def test_editar_en_sitio_sobrescribe_el_original(carpeta):
    from docx import Document

    original = _word_base(carpeta)
    r = documentos.editar_documento(original, {"reemplazos": {"320": "999"}}, en_sitio=True)

    assert r.ruta == original
    assert any("999 millones" in p.text for p in Document(original).paragraphs)


def test_editar_word_con_un_ancla_que_no_existe_deja_el_contenido_al_final_y_avisa(carpeta):
    from docx import Document

    r = documentos.editar_documento(_word_base(carpeta), {
        "agregar": "Un párrafo nuevo con contenido suficiente para que no lo tome por relleno "
                   "ni por un documento vacío en la validación.",
        "despues_de": "sección inexistente"})
    textos = [p.text for p in Document(r.ruta).paragraphs if p.text.strip()]

    assert textos[-1].startswith("Un párrafo nuevo")
    assert any("No encontré" in a for a in r.avisos)


def test_editar_sin_cambios_lo_dice(carpeta):
    with pytest.raises(documentos.DocumentoRechazado, match="ningún cambio"):
        documentos.editar_documento(_word_base(carpeta), {})


def test_editar_con_relleno_se_rechaza(carpeta):
    with pytest.raises(documentos.DocumentoRechazado, match="relleno"):
        documentos.editar_documento(_word_base(carpeta), {"reemplazos": {"320": "[insertar]"}})


def test_editar_fuera_de_las_carpetas_permitidas_se_rechaza(carpeta, tmp_path, monkeypatch):
    original = _word_base(carpeta)
    monkeypatch.setattr(documentos, "_bases_permitidas", lambda: [tmp_path / "otra"])
    with pytest.raises(documentos.DocumentoRechazado, match="no lo edito"):
        documentos.editar_documento(original, {"reemplazos": {"a": "b"}})


def test_editar_excel_celdas_y_filas(carpeta):
    from openpyxl import load_workbook

    origen = documentos.crear_hoja("v.xlsx", {"hojas": [{"nombre": "Ventas",
                                                          "encabezados": ["Mes", "Ventas"],
                                                          "filas": [["Enero", 100], ["Febrero", 200]]}]},
                                   carpeta=carpeta).ruta
    r = documentos.editar_documento(origen, {"celdas": {"Ventas!B3": "250"},
                                             "agregar": {"Ventas": [["Marzo", "1.300"]]},
                                             "reemplazos": {"Febrero": "Feb"}}, en_sitio=True)
    ws = load_workbook(origen)["Ventas"]

    assert ws["B3"].value == 250
    assert ws["A3"].value == "Feb"
    assert [c.value for c in ws[4]] == ["Marzo", 1300]
    assert "1 celdas escritas" in r.detalle


def test_editar_powerpoint_quita_por_titulo_y_agrega_sin_corromper(carpeta):
    """Caso real: borrar la 2 de tres y agregar una chocaba con el nombre de la 3."""
    from pptx import Presentation

    origen = documentos.crear_presentacion("p.pptx", {"titulo": "R", "diapositivas": [
        {"titulo": "Producción", "puntos": ["1.850.000 t"]},
        {"titulo": "Borrador", "texto": "quitar esto"},
        {"titulo": "Cierre", "texto": "Gracias."}]}, carpeta=carpeta).ruta
    r = documentos.editar_documento(origen, {"reemplazos": {"Gracias.": "Muchas gracias."},
                                             "quitar": ["Borrador"],
                                             "agregar": [{"titulo": "Anexo", "puntos": ["Datos", "Método"]}]})
    prs = Presentation(r.ruta)

    assert [s.shapes.title.text for s in prs.slides] == ["R", "Producción", "Cierre", "Anexo"]
    partes = sorted(str(s.part.partname) for s in prs.slides)
    assert len(partes) == len(set(partes)) == 4


def test_editar_powerpoint_quita_por_numero(carpeta):
    from pptx import Presentation

    origen = documentos.crear_presentacion("p.pptx", {"titulo": "R", "diapositivas": [
        {"titulo": "Uno", "texto": "Contenido uno."}, {"titulo": "Dos", "texto": "Contenido dos."}]},
        carpeta=carpeta).ruta
    r = documentos.editar_documento(origen, {"quitar": ["2"]})

    assert [s.shapes.title.text for s in Presentation(r.ruta).slides] == ["R", "Dos"]


def test_document_edit_registrada_y_amarilla():
    from agents.tool_registry import get_tool
    from core.security_manager import RiskLevel, security_manager

    assert get_tool("document_edit") is not None
    assert security_manager.classify_action("document_edit") == RiskLevel.YELLOW


# =============================================================================== REQ-065
# Segunda vuelta de documentos: lo que falló al validar REQ-040 con archivos reales, y el
# estilo guiado por el usuario. Cada test nombra el criterio de la SPEC que protege.


# ---------------------------------------------------------------- CA-01: fórmulas en español

@pytest.mark.parametrize("en_espanol, como_la_guarda_excel", [
    ("=D5*0,19", "=D5*0.19"),
    ('=SI(B2>5;"al;to";"bajo")', '=IF(B2>5,"al;to","bajo")'),
    ("=SUMA(B2:B4)", "=SUM(B2:B4)"),
    ("=REDONDEAR(B5*1,19;2)", "=ROUND(B5*1.19,2)"),
    ("=MÁX(A1:A3)+AÑO(HOY())", "=MAX(A1:A3)+YEAR(TODAY())"),
    ('=CONTAR.SI(A:A;"x")', '=COUNTIF(A:A,"x")'),
    ("=VERDADERO", "=TRUE"),
    ("=PROMEDIO(B2:B4)*1,5", "=AVERAGE(B2:B4)*1.5"),
    ("=si(a1>0;1;0)", "=IF(a1>0,1,0)"),
    ("=SUMA(A1;{1;2})", "=SUM(A1,{1;2})"),             # el «;» de una matriz no es separador
])
def test_una_formula_en_espanol_se_guarda_como_excel_la_entiende(en_espanol, como_la_guarda_excel):
    assert documentos.normalizar_formula(en_espanol) == como_la_guarda_excel


@pytest.mark.parametrize("en_ingles", [
    "=ROUND(B5*1.19,2)", "=SUM(B2:B4)/2", "=Clientes!C4", "='Hoja dos'!B2*2", "=B5*C5",
    "=IF(A1>0,1,0)",
])
def test_una_formula_que_ya_venia_en_ingles_no_se_toca(en_ingles):
    assert documentos.normalizar_formula(en_ingles) == en_ingles


def test_las_formulas_del_excel_salen_normalizadas_en_el_archivo(carpeta):
    from openpyxl import load_workbook

    r = documentos.crear_hoja("f.xlsx", {"hojas": [{"nombre": "H", "encabezados": ["a", "b", "c"],
                                                     "filas": [["x", "10", "=B2*0,19"],
                                                               ["y", "20", '=SI(B3>5;"alto";"bajo")'],
                                                               ["z", "30", "=SUMA(B2:B3)"]]}]},
                              carpeta=carpeta)
    ws = load_workbook(r.ruta)["H"]

    assert ws["C2"].value == "=B2*0.19"
    assert ws["C3"].value == '=IF(B3>5,"alto","bajo")'
    assert ws["C4"].value == "=SUM(B2:B3)"


# ---------------------------------------------------------------- CA-02 / CA-03 / CA-07: plantilla Excel

@pytest.fixture
def plantilla_xlsx_con_titulo(tmp_path):
    """Título en la 1, subtítulo en la 2, encabezado verde en la 4 y nada más."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Ventas"
    ws["A1"] = "AGRO SAS — Reporte de ventas"
    ws["A2"] = "Área comercial"
    for i, h in enumerate(["Producto", "Cantidad", "Precio", "Subtotal"], 1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1B5E20")
    ruta = tmp_path / "plantilla_ventas.xlsx"
    wb.save(str(ruta))
    return str(ruta)


def test_desde_no_repite_el_encabezado_de_la_plantilla_y_los_datos_van_donde_el_modelo_cree(
        carpeta, plantilla_xlsx_con_titulo):
    from openpyxl import load_workbook

    r = documentos.crear_hoja("v.xlsx", {"hojas": [{
        "nombre": "Ventas", "desde": "A5",
        "encabezados": ["Producto", "Cantidad", "Precio", "Subtotal"],
        "filas": [["Abono", "120", "35.500", "=B5*C5"], ["Semilla", "40", "870.000", "=B6*C6"]],
        "totales": True}]}, plantilla=plantilla_xlsx_con_titulo, carpeta=carpeta)
    ws = load_workbook(r.ruta)["Ventas"]

    assert [c.value for c in ws[4]] == ["Producto", "Cantidad", "Precio", "Subtotal"]   # el suyo
    assert ws["A4"].fill.fgColor.rgb.endswith("1B5E20")                                  # con su verde
    assert ws["A5"].value == "Abono" and ws["D5"].value == "=B5*C5"   # justo donde el modelo lo escribió
    assert ws["A7"].value == "Total" and ws["D7"].value == "=SUM(D5:D6)"
    assert "datos desde la fila 5" in r.detalle


def test_desde_respeta_la_columna(carpeta, plantilla_xlsx_con_titulo):
    from openpyxl import load_workbook

    r = documentos.crear_hoja("v.xlsx", {"hojas": [{
        "nombre": "Ventas", "desde": "C8", "encabezados": ["Mes", "Total"],
        "filas": [["Enero", "100"], ["Febrero", "200"]]}]},
        plantilla=plantilla_xlsx_con_titulo, carpeta=carpeta)
    ws = load_workbook(r.ruta)["Ventas"]

    assert ws["C8"].value == "Mes" and ws["D8"].value == "Total"
    assert ws["C9"].value == "Enero" and ws["D10"].value == 200
    assert ws["A8"].value is None


def test_el_grafico_sobre_plantilla_grafica_la_columna_pedida_con_su_nombre(carpeta, plantilla_xlsx_con_titulo):
    from openpyxl import load_workbook

    r = documentos.crear_hoja("v.xlsx", {"hojas": [{
        "nombre": "Ventas", "desde": "A5",
        "encabezados": ["Producto", "Cantidad", "Precio", "Subtotal"],
        "filas": [["Abono", "120", "35.500", "4.260.000"], ["Semilla", "40", "870.000", "34.800.000"]],
        "grafico": {"tipo": "barras", "titulo": "Subtotal", "columnas": ["Subtotal"],
                    "categorias": "Producto"}}]},
        plantilla=plantilla_xlsx_con_titulo, carpeta=carpeta)
    ws = load_workbook(r.ruta)["Ventas"]
    serie = ws._charts[0].series[0]

    assert "$D$5:$D$6" in serie.val.numRef.f            # Subtotal, no la columna 2
    assert serie.tx.v == "Subtotal"                      # y con su nombre, no «Series1»
    categorias = (serie.cat.numRef.f if serie.cat.numRef is not None else serie.cat.strRef.f)
    assert "$A$5:$A$6" in categorias


def test_un_encabezado_nuevo_en_la_plantilla_copia_el_estilo_del_que_ya_tiene(carpeta, plantilla_xlsx_con_titulo):
    from openpyxl import load_workbook

    r = documentos.crear_hoja("v.xlsx", {"hojas": [{
        "nombre": "Ventas", "encabezados": ["Vendedor", "Zona", "Comisión"],   # otros nombres
        "filas": [["Ana", "Norte", "1.200.000"]]}]},
        plantilla=plantilla_xlsx_con_titulo, carpeta=carpeta)
    ws = load_workbook(r.ruta)["Ventas"]

    fila = next(f for f in range(5, ws.max_row + 1) if ws.cell(row=f, column=1).value == "Vendedor")
    assert ws.cell(row=fila, column=1).fill.fgColor.rgb.endswith("1B5E20")   # el verde de la plantilla
    assert ws.cell(row=fila + 1, column=1).value == "Ana"


# ---------------------------------------------------------------- CA-04: agregar filas con totales

def test_agregar_filas_en_una_hoja_con_totales_las_mete_antes_y_la_suma_las_incluye(carpeta):
    from openpyxl import load_workbook

    origen = documentos.crear_hoja("c.xlsx", {"hojas": [
        {"nombre": "Clientes", "encabezados": ["Nombre", "Ciudad", "Deuda"],
         "filas": [["Finca", "Ibagué", "4.350.000"], ["Roble", "Neiva", "250.000"]], "totales": True,
         "formatos": {"C": '"$" #,##0'},
         "grafico": {"tipo": "barras", "columnas": ["Deuda"], "categorias": "Nombre"}},
        {"nombre": "Resumen", "encabezados": ["Concepto", "Valor"],
         "filas": [["Cartera total", "=Clientes!C4"]]}]}, carpeta=carpeta).ruta
    r = documentos.editar_documento(origen, {"agregar": {"Clientes": [["Vivero", "Cali", "1.100.000"]]}},
                                    en_sitio=True)
    wb = load_workbook(origen)
    ws = wb["Clientes"]

    assert [c.value for c in ws[4]] == ["Vivero", "Cali", 1100000]      # antes del total
    assert ws["A5"].value == "Total" and ws["C5"].value == "=SUM(C2:C4)"   # y la suma la incluye
    assert ws["C4"].number_format == '"$" #,##0'                         # con el formato de la fila anterior
    assert wb["Resumen"]["B2"].value == "=Clientes!C5"                  # la otra hoja sigue apuntando al total
    assert "$C$2:$C$4" in ws._charts[0].series[0].val.numRef.f           # y el gráfico también
    assert "antes de la fila de totales" in r.detalle


def test_agregar_filas_sin_totales_sigue_agregando_al_final(carpeta):
    from openpyxl import load_workbook

    origen = documentos.crear_hoja("s.xlsx", {"hojas": [{"nombre": "H", "encabezados": ["a", "b"],
                                                          "filas": [["x", "1"]]}]}, carpeta=carpeta).ruta
    documentos.editar_documento(origen, {"agregar": {"H": [["y", "2"]]}}, en_sitio=True)
    ws = load_workbook(origen)["H"]

    assert [c.value for c in ws[3]] == ["y", 2]


# ---------------------------------------------------------------- CA-05: marcadores sin rellenar

def test_contenido_nuevo_sobre_una_plantilla_con_marcador_en_el_pie_se_rechaza_y_dice_donde(carpeta, tmp_path):
    from docx import Document

    doc = Document()
    doc.sections[0].footer.paragraphs[0].text = "Cliente: {{cliente}}"
    plantilla = tmp_path / "membrete.docx"
    doc.save(str(plantilla))

    with pytest.raises(documentos.DocumentoRechazado, match=r"\{\{cliente\}\} \(pie de página\)"):
        documentos.crear_documento("informe.docx", _INFORME, plantilla=str(plantilla), carpeta=carpeta)
    assert not os.path.exists(os.path.join(carpeta, "informe.docx"))

    # Con el valor, sale.
    r = documentos.crear_documento("informe.docx", _INFORME, plantilla=str(plantilla),
                                   reemplazos={"{{cliente}}": "Palmar"}, carpeta=carpeta)
    assert Document(r.ruta).sections[0].footer.paragraphs[0].text == "Cliente: Palmar"


def test_una_presentacion_sobre_plantilla_con_marcador_sin_valor_se_rechaza(carpeta):
    from pptx import Presentation

    origen = documentos.crear_presentacion("base.pptx", {"titulo": "[EMPRESA]", "diapositivas": [
        {"titulo": "Uno", "texto": "Contenido uno."}, {"titulo": "Dos", "texto": "Contenido dos."}]},
        carpeta=carpeta).ruta

    with pytest.raises(documentos.DocumentoRechazado, match=r"\[EMPRESA\] \(diapositiva 1\)"):
        documentos.crear_presentacion("nueva.pptx", {"bloques": []}, plantilla=origen,
                                      reemplazos={"{{nada}}": "x"}, carpeta=carpeta)
    r = documentos.crear_presentacion("nueva.pptx", {"bloques": []}, plantilla=origen,
                                      reemplazos={"[EMPRESA]": "Agro SAS"}, carpeta=carpeta)
    assert Presentation(r.ruta).slides[0].shapes.title.text == "Agro SAS"


# ---------------------------------------------------------------- CA-06: formas tolerantes

def test_agregar_con_despues_de_adentro_funciona_igual_que_la_forma_documentada(carpeta):
    from docx import Document

    origen = documentos.crear_documento("i.docx", _INFORME, carpeta=carpeta).ruta
    r = documentos.editar_documento(origen, {"agregar": {
        "despues_de": "Visita al lote 14",
        "bloques": [{"tipo": "parrafo", "texto": "Nota: cifras preliminares."}]}})
    textos = [p.text for p in Document(r.ruta).paragraphs]

    assert textos[textos.index("Visita al lote 14") + 1] == "Nota: cifras preliminares."
    assert "después de «Visita al lote 14»" in r.detalle
    assert "despues_de" not in textos and "bloques" not in textos     # nada de claves como párrafos


def test_bloques_que_es_un_objeto_sin_sentido_se_rechaza(carpeta):
    with pytest.raises(documentos.DocumentoRechazado, match="lista de bloques"):
        documentos.crear_documento("x.docx", {"bloques": {"seccion": "a", "otra": "b"}}, carpeta=carpeta)


def test_un_bloque_suelto_como_objeto_se_acepta(carpeta):
    r = documentos.crear_documento("x.docx", {"bloques": {"tipo": "parrafo", "texto": "a" * 150}},
                                   carpeta=carpeta)
    assert "1 bloques" in r.detalle


# ---------------------------------------------------------------- CA-08 / CA-09 / CA-12: estilo en Word

def test_el_estilo_del_usuario_se_aplica_al_word(carpeta, tmp_path):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

    logo = tmp_path / "logo.png"
    documentos._grafico_como_imagen({"categorias": ["a"], "series": [{"nombre": "s", "valores": [1]}]}, str(logo))

    r = documentos.crear_documento("e.docx", _INFORME, carpeta=carpeta, estilo={
        "fuente": "Arial", "tamaño": "11", "fuente_titulos": "Georgia", "color_titulos": "azul oscuro",
        "márgenes": 2, "orientacion": "horizontal", "encabezado": "AGRO SAS · Informe interno",
        "pie": "Confidencial", "logo": str(logo), "numeracion": True, "justificado": True})
    doc = Document(r.ruta)
    seccion = doc.sections[0]

    assert doc.styles["Normal"].font.name == "Arial" and doc.styles["Normal"].font.size.pt == 11
    assert doc.styles["Normal"].paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert doc.styles["Heading 1"].font.name == "Georgia"
    assert str(doc.styles["Heading 1"].font.color.rgb) == "1F3864"
    assert seccion.orientation == WD_ORIENT.LANDSCAPE and seccion.page_width > seccion.page_height
    assert round(seccion.left_margin.cm, 1) == 2.0
    assert any("Informe interno" in p.text for p in seccion.header.paragraphs)
    assert seccion.header._element.findall(".//" + qn("w:drawing"))          # el logo está
    assert any("Confidencial" in p.text for p in seccion.footer.paragraphs)
    instrucciones = [e.text for e in seccion.footer._element.iter(qn("w:instrText"))]
    assert any("PAGE" in i for i in instrucciones) and any("NUMPAGES" in i for i in instrucciones)
    assert "estilo:" in r.detalle and "número de página" in r.detalle


def test_un_color_o_un_tamano_que_no_se_entiende_se_avisa_y_no_tumba_el_documento(carpeta):
    r = documentos.crear_documento("e.docx", _INFORME, carpeta=carpeta,
                                   estilo={"color_titulos": "fucsia brillante", "tamano": "doce"})
    assert os.path.exists(r.ruta)
    assert any("fucsia brillante" in a for a in r.avisos) and any("doce" in a for a in r.avisos)


def test_los_colores_se_entienden_por_nombre_y_en_hex():
    c = documentos._color_hex
    assert c("azul oscuro") == "1F3864" and c("Verde") == "2E7D32"
    assert c("#2e7d32") == "2E7D32" and c("2E7D32") == "2E7D32" and c("#abc") == "AABBCC"
    assert c("fucsia brillante") is None and c("") is None


def test_el_estilo_encima_de_la_plantilla_manda(carpeta, plantilla_docx):
    from docx import Document

    r = documentos.crear_documento("i.docx", _INFORME, plantilla=plantilla_docx, carpeta=carpeta,
                                   reemplazos={"{{numero}}": "1", "{{cliente}}": "c",
                                               "{{concepto}}": "u", "{{valor}}": "1"},
                                   estilo={"color_titulos": "rojo"})
    doc = Document(r.ruta)

    assert doc.sections[0].header.paragraphs[0].text == "AGRO ORIENTAL S.A.S."   # la plantilla sigue
    assert str(doc.styles["Heading 1"].font.color.rgb) == "C00000"              # pero el color es el pedido


def test_el_estilo_no_aplica_a_md_y_se_dice(carpeta):
    r = documentos.crear_documento("n.md", _INFORME, carpeta=carpeta, estilo={"fuente": "Arial"})
    assert any("no aplica" in a for a in r.avisos)


def test_el_html_lleva_la_fuente_y_los_colores_pedidos(carpeta):
    r = documentos.crear_documento("n.html", _INFORME, carpeta=carpeta,
                                   estilo={"fuente": "Georgia", "color_titulos": "verde", "tamano": 12})
    html = Path(r.ruta).read_text(encoding="utf-8")
    assert "'Georgia'" in html and "#2E7D32" in html and "12pt" in html


# ---------------------------------------------------------------- CA-10: estilo y título en Excel

def test_el_estilo_y_el_titulo_de_hoja_se_aplican_al_excel(carpeta):
    from openpyxl import load_workbook

    r = documentos.crear_hoja("v.xlsx", {"hojas": [{
        "nombre": "Ventas", "titulo": "AGRO SAS — Ventas del trimestre",
        "encabezados": ["Producto", "Total"], "filas": [["Abono", "100"], ["Semilla", "200"]],
        "totales": True}]}, carpeta=carpeta,
        estilo={"fuente": "Calibri", "tamano": 10, "color_encabezado": "verde oscuro", "bordes": False})
    ws = load_workbook(r.ruta)["Ventas"]

    assert ws["A1"].value == "AGRO SAS — Ventas del trimestre" and "A1:B1" in ws.merged_cells
    assert ws["A3"].value == "Producto" and ws["A3"].fill.fgColor.rgb.endswith("1B5E20")
    assert ws["A4"].value == "Abono" and ws["A4"].font.name == "Calibri" and ws["A4"].font.size == 10
    assert ws["B6"].value == "=SUM(B4:B5)"
    assert ws["A4"].border.left.style is None                       # sin bordes, como se pidió
    assert ws.freeze_panes == "A4"
    assert "encabezado en la 3" in r.detalle and "datos desde la fila 4" in r.detalle


# ---------------------------------------------------------------- CA-11: estilo en PowerPoint

def test_el_estilo_se_aplica_a_la_presentacion(carpeta, tmp_path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    logo = tmp_path / "logo.png"
    documentos._grafico_como_imagen({"categorias": ["a"], "series": [{"nombre": "s", "valores": [1]}]}, str(logo))
    r = documentos.crear_presentacion("p.pptx", {"titulo": "Resultados", "diapositivas": [
        {"titulo": "Uno", "puntos": ["Primer punto", "Segundo punto"]},
        {"titulo": "Dos", "texto": "Contenido dos."}]}, carpeta=carpeta,
        estilo={"fuente": "Georgia", "color_titulos": "verde", "color_texto": "#333333",
                "logo": str(logo), "pie": "Agro SAS · Confidencial", "numeracion": True})
    prs = Presentation(r.ruta)
    segunda = prs.slides[1]
    titulo = segunda.shapes.title

    assert titulo.text_frame.paragraphs[0].runs[0].font.name == "Georgia"
    assert str(titulo.text_frame.paragraphs[0].runs[0].font.color.rgb) == "2E7D32"
    for slide in prs.slides:
        assert any(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in slide.shapes)      # logo en todas
    textos = [s.text_frame.text for s in segunda.shapes if s.has_text_frame]
    assert "Agro SAS · Confidencial" in textos and "2 / 3" in textos
    assert "estilo:" in r.detalle and "logo" in r.detalle


# ---------------------------------------------------------------- CA-13: el registro

def test_las_tres_herramientas_exponen_estilo_y_spreadsheet_explica_desde():
    from agents.tool_registry import get_tool

    for nombre in ("document_create", "spreadsheet_create", "presentation_create"):
        tool = get_tool(nombre)
        assert "estilo" in tool.parameters_schema["properties"], nombre
        assert "estilo" in tool.description.lower(), nombre
    hoja = get_tool("spreadsheet_create")
    assert "PRIMERA FILA DE DATOS" in hoja.parameters_schema["properties"]["contenido"]["description"]
    assert "SUMA" in hoja.description
