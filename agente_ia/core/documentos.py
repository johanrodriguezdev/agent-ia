"""
core/documentos.py
REQ-040 — Crear documentos de verdad: Word, Excel, PowerPoint, PDF y lo que haga falta.

**Qué había.** El agente escribía texto plano (`file_write`, `write_file_direct`) y leía
documentos de Office (`FileAnalysisSkill`), pero no podía *producir* uno: pedirle "hazme el
informe en Word" terminaba en un `.txt` en el escritorio. Y menos seguir una plantilla: el
membrete de la empresa, la cotización con el logo, la presentación con el tema corporativo.

**Cómo se reparte el trabajo.** El modelo escribe el contenido —todo el contenido: las
secciones, las filas, las diapositivas— como una estructura, y este módulo lo convierte en
un archivo bien formado. Es la misma división que en el resto del proyecto: la inteligencia
en el modelo, la mecánica acá. Lo que este módulo sí decide es que un documento vacío o de
relleno **no se produce**: "lorem ipsum", "[insertar aquí]", "por definir" se rechazan con
un mensaje que le pide al modelo el contenido real. Un archivo con marcadores de relleno
parece trabajo hecho y no lo es; peor que ningún archivo.

**Plantillas.** Se abre el archivo de plantilla y se escribe *dentro* de él, así que
conserva sus estilos, encabezados, pies, márgenes, tema y patrones. Hay dos maneras, que se
pueden combinar:

- `reemplazos`: la plantilla trae marcadores —`{{cliente}}`, `[FECHA]`— y se sustituyen,
  conservando el formato del texto que estaba. Es lo que se hace con una cotización.
- contenido nuevo: se vacía el cuerpo de la plantilla (nunca los encabezados ni los pies) y
  se escribe el contenido con los estilos que la plantilla define. Es lo que se hace con un
  membrete.

Si tras rellenar queda algún marcador (`{{cliente}}` en el pie, `[FECHA]` en una tabla), el
archivo **no se guarda**: se rechaza diciendo cuál y dónde, para que el modelo vuelva con
el valor o se lo pregunte al usuario (REQ-065).

**Estilo guiado por el usuario (REQ-065).** Sin plantilla —o encima de una— el usuario
puede pedir el aspecto en palabras y el modelo lo pasa en `estilo`: fuente, tamaño,
color de títulos, márgenes, orientación, encabezado y pie, logo, número de página; en
Excel el color y la fuente del encabezado y un título de hoja; en PowerPoint fuente,
colores, logo, pie y número de diapositiva. Un color que no se entiende se avisa y se
ignora: nunca impide que salga el archivo.

**Fórmulas (REQ-065).** El modelo escribe `=SUMA(B2:B4)`, `=SI(a;b;c)` o `=D5*0,19` con
la misma naturalidad con la que habla, y Excel guarda las fórmulas siempre en inglés y con
punto decimal: tal cual, el archivo mostraba `#¿NOMBRE?` o directamente no abría. Se
normalizan al escribir (`normalizar_formula`).

**PDF.** Con Office instalado —como en esta máquina— se convierte con el propio Word, Excel
o PowerPoint, que es la única forma de que el PDF se vea exactamente igual que el
documento. Sin Office, un documento de texto se maqueta con PyMuPDF; una hoja o una
presentación no tienen ese respaldo y se dice.

**Dónde quedan.** En el Escritorio, salvo que se pida otra carpeta dentro de la carpeta
personal del usuario o de un espacio de trabajo habilitado. Nunca se pisa un archivo que
ya exista: se numera (`informe (2).docx`), como hacen los adjuntos de correo.
"""

import csv
import json
import logging
import os
import re
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

logger = logging.getLogger(__name__)

#: Lo que este módulo sabe producir, por extensión.
EXTENSIONES_DOCUMENTO = ("docx", "pdf", "md", "html", "txt")
EXTENSIONES_HOJA = ("xlsx", "csv")
EXTENSIONES_PRESENTACION = ("pptx",)
EXTENSIONES_DATOS = ("json",)

#: Marcadores de relleno. Un documento que los trae no está hecho, está esbozado, y se
#: rechaza. La lista es corta a propósito: mejor dejar pasar un caso raro que rechazar un
#: documento real por una palabra normal.
_RELLENO = (
    "lorem ipsum", "[insertar", "[inserte", "[introducir", "[completar", "[pendiente",
    "[por definir", "por definir]", "texto de ejemplo", "texto de prueba", "contenido aquí",
    "contenido aqui", "tbd", "todo:", "xxxx", "placeholder", "{{", "}}", "<<", ">>",
    "ejemplo de párrafo", "ejemplo de parrafo",
)

#: Menos que esto y el documento es un título con nada debajo.
_MINIMO_CARACTERES_DOCUMENTO = 120

#: Ancho por defecto de una imagen dentro de un documento, en centímetros.
_ANCHO_IMAGEN_CM = 14.0

#: Tope de ancho de columna en Excel. Un texto largo no puede convertir la columna en una
#: cinta de tres pantallas.
_ANCHO_COLUMNA_MAXIMO = 60


#: Carpetas de la carpeta personal donde NO se crean documentos: son de configuración y
#: de llaves, no de trabajo del usuario.
CARPETAS_VEDADAS = frozenset({"appdata", ".ssh", ".gnupg", ".aws", ".azure", ".config"})


class DocumentoRechazado(Exception):
    """El contenido o el destino no sirven, y el motivo se le puede decir al modelo."""


@dataclass
class Resultado:
    """Qué se produjo y dónde, en términos que el modelo pueda contarle al usuario."""

    ruta: str
    formato: str
    detalle: str = ""
    avisos: List[str] = field(default_factory=list)

    def describir(self) -> str:
        texto = f"Listo: {self.ruta}"
        if self.detalle:
            texto += f" ({self.detalle})"
        texto += "."
        for aviso in self.avisos:
            texto += f" {aviso}"
        return texto


# =============================================================================== destino

def _escritorio() -> Path:
    """Return el Escritorio del usuario: el de OneDrive si lo redirige, si no el normal."""
    onedrive = Path.home() / "OneDrive" / "Desktop"
    if onedrive.exists():
        return onedrive
    escritorio = Path.home() / "Desktop"
    if escritorio.exists():
        return escritorio
    return Path.home()


def _bases_permitidas() -> List[Path]:
    """Return dónde se pueden crear archivos: la carpeta personal y los espacios de trabajo."""
    bases = [Path.home()]
    try:
        from core.workspace_config import cargar_raices

        bases += [Path(r) for r in cargar_raices()]
    except Exception as e:
        logger.debug(f"no se pudieron leer los espacios de trabajo: {e}")
    return bases


def carpeta_de_salida(carpeta: Optional[str]) -> Path:
    """Return la carpeta donde se va a escribir, ya comprobada. Levanta `DocumentoRechazado`."""
    if not carpeta or not str(carpeta).strip():
        return _escritorio()

    pedida = Path(os.path.expanduser(str(carpeta).strip()))
    if not pedida.is_absolute():
        pedida = _escritorio() / pedida
    real = Path(os.path.realpath(pedida))

    for base in _bases_permitidas():
        try:
            base_real = Path(os.path.realpath(base))
            if os.path.commonpath([str(real).lower(), str(base_real).lower()]) == \
                    str(base_real).lower():
                break
        except ValueError:
            continue
    else:
        raise DocumentoRechazado(
            f"No creo archivos en «{carpeta}»: solo dentro de tu carpeta personal o de un "
            f"espacio de trabajo habilitado."
        )

    # Dentro de la carpeta personal hay sitios que no son para documentos.
    nombre = real.name.lower()
    partes = {p.lower() for p in real.parts}
    if partes & CARPETAS_VEDADAS or nombre.startswith("."):
        raise DocumentoRechazado(
            f"«{carpeta}» es una carpeta de configuración, no de documentos. Elegí otra."
        )

    try:
        real.mkdir(parents=True, exist_ok=True)
    except OSError as e:
        raise DocumentoRechazado(f"No pude crear la carpeta «{carpeta}»: {e.strerror or e}")
    return real


def ruta_libre(carpeta: Path, nombre: str) -> Path:
    """Return una ruta que no exista: si `informe.docx` está, `informe (2).docx`.

    Nunca se pisa un archivo del usuario. Un documento que él hizo a mano y que el agente
    reescribe porque se llamaba igual no se recupera.
    """
    limpio = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(nombre or "").strip()) or "documento"
    base, ext = os.path.splitext(limpio)
    candidata = carpeta / limpio
    numero = 2
    while candidata.exists():
        candidata = carpeta / f"{base} ({numero}){ext}"
        numero += 1
    return candidata


def _extension(nombre: str) -> str:
    return os.path.splitext(str(nombre or ""))[1].lower().lstrip(".")


def _plantilla_valida(plantilla: Optional[str], extensiones: Tuple[str, ...]) -> Optional[Path]:
    """Return la ruta de la plantilla comprobada, o None si no se pidió."""
    if not plantilla or not str(plantilla).strip():
        return None
    ruta = Path(os.path.expanduser(str(plantilla).strip()))
    if not ruta.is_absolute():
        for base in [_escritorio(), Path.home() / "Documents", Path.home() / "OneDrive" / "Documentos",
                     Path.home() / "Downloads"]:
            if (base / ruta).exists():
                ruta = base / ruta
                break
    if not ruta.is_file():
        raise DocumentoRechazado(f"No encuentro la plantilla «{plantilla}».")
    if _extension(ruta.name) not in extensiones:
        raise DocumentoRechazado(
            f"«{ruta.name}» no sirve de plantilla para esto: hace falta un archivo "
            f"{', '.join('.' + e for e in extensiones)}."
        )
    return ruta


# ============================================================================== contenido

def _texto_de(valor: Any) -> str:
    """Return todo el texto que haya dentro de una estructura, para revisarlo."""
    if isinstance(valor, str):
        return valor
    if isinstance(valor, dict):
        return " ".join(_texto_de(v) for v in valor.values())
    if isinstance(valor, (list, tuple)):
        return " ".join(_texto_de(v) for v in valor)
    return "" if valor is None else str(valor)


def rechazar_relleno(contenido: Any) -> None:
    """Levanta `DocumentoRechazado` si el contenido trae marcadores de relleno.

    Se revisa lo que el modelo escribió, no la plantilla: los `{{marcadores}}` de una
    plantilla son justamente lo que se va a sustituir.
    """
    texto = _texto_de(contenido).lower()
    hallados = sorted({marca for marca in _RELLENO if marca in texto})
    if hallados:
        raise DocumentoRechazado(
            f"El contenido trae marcadores de relleno ({', '.join(repr(h) for h in hallados)}). "
            f"Un documento con huecos no es un documento: escribí el contenido completo y "
            f"real en cada sección antes de generarlo."
        )


def _como_estructura(contenido: Any) -> Any:
    """Return `contenido` como diccionario o lista, venga como venga.

    Los modelos mandan a veces el JSON como texto dentro del argumento. Se acepta, porque
    rechazarlo sería gastar una vuelta en pedirlo de otra forma.
    """
    if isinstance(contenido, str):
        recortado = contenido.strip()
        if recortado.startswith(("{", "[")):
            try:
                return json.loads(recortado)
            except json.JSONDecodeError as e:
                raise DocumentoRechazado(f"El contenido no es un JSON válido: {e}")
        return recortado
    return contenido


# ================================================================= documentos de texto

def _bloques_de(spec: Any) -> Tuple[Dict[str, Any], List[Dict[str, Any]]]:
    """Return `(cabecera, bloques)` de una especificación de documento.

    Se acepta un texto suelto —se parte en párrafos— y también Markdown sencillo: `#`
    títulos, `-` listas, `1.` numeradas. Es lo que un modelo escribe con más naturalidad.
    """
    if isinstance(spec, str):
        return {}, _bloques_desde_markdown(spec)
    if isinstance(spec, list):
        return {}, [b if isinstance(b, dict) else {"tipo": "parrafo", "texto": str(b)} for b in spec]
    if not isinstance(spec, dict):
        raise DocumentoRechazado("El contenido del documento tiene que ser un objeto con "
                                 "'bloques', una lista de bloques o un texto en Markdown.")
    bloques = spec.get("bloques") or spec.get("secciones") or spec.get("contenido") or []
    if isinstance(bloques, str):
        bloques = _bloques_desde_markdown(bloques)
    if isinstance(bloques, dict):
        # Un bloque suelto ({"tipo": "parrafo", ...}) se acepta; un objeto con otra forma
        # no: iterarlo escribiría sus claves («despues_de», «bloques») como párrafos y el
        # documento saldría con basura diciendo «Listo». Pasó en la validación.
        parece_bloque = any(bloques.get(k) for k in ("tipo", "texto", "items", "filas"))
        if parece_bloque:
            bloques = [bloques]
        else:
            raise DocumentoRechazado(
                "'bloques' tiene que ser una lista de bloques ({tipo, texto...}) o un texto en "
                f"Markdown, no un objeto con las claves {', '.join(map(str, bloques))}."
            )
    if not isinstance(bloques, (list, tuple)):
        raise DocumentoRechazado(
            "'bloques' tiene que ser una lista de bloques o un texto en Markdown.")
    cabecera = {k: spec.get(k) for k in ("titulo", "subtitulo", "autor", "fecha") if spec.get(k)}
    return cabecera, [b if isinstance(b, dict) else {"tipo": "parrafo", "texto": str(b)}
                      for b in bloques]


def _bloques_desde_markdown(texto: str) -> List[Dict[str, Any]]:
    """Return bloques a partir de Markdown sencillo. No es un parser completo ni lo pretende."""
    bloques: List[Dict[str, Any]] = []
    lista: Optional[Dict[str, Any]] = None
    parrafo: List[str] = []

    def _cerrar_parrafo():
        nonlocal parrafo
        if parrafo:
            bloques.append({"tipo": "parrafo", "texto": " ".join(parrafo)})
            parrafo = []

    def _cerrar_lista():
        nonlocal lista
        if lista:
            bloques.append(lista)
            lista = None

    for linea in (texto or "").splitlines():
        cruda = linea.rstrip()
        if not cruda.strip():
            _cerrar_parrafo()
            _cerrar_lista()
            continue
        encabezado = re.match(r"^(#{1,4})\s+(.*)$", cruda)
        if encabezado:
            _cerrar_parrafo()
            _cerrar_lista()
            bloques.append({"tipo": "titulo", "nivel": len(encabezado.group(1)),
                            "texto": encabezado.group(2).strip()})
            continue
        item = re.match(r"^\s*(?:[-*•]|\d+[.)])\s+(.*)$", cruda)
        if item:
            _cerrar_parrafo()
            numerada = bool(re.match(r"^\s*\d+[.)]", cruda))
            if lista is None or lista.get("numerada") != numerada:
                _cerrar_lista()
                lista = {"tipo": "lista", "numerada": numerada, "items": []}
            lista["items"].append(item.group(1).strip())
            continue
        if cruda.startswith("|") and cruda.endswith("|"):
            _cerrar_parrafo()
            _cerrar_lista()
            celdas = [c.strip() for c in cruda.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in celdas):
                continue                       # la línea de separación de la tabla
            if bloques and bloques[-1].get("tipo") == "tabla" and bloques[-1].get("_abierta"):
                bloques[-1]["filas"].append(celdas)
            else:
                bloques.append({"tipo": "tabla", "encabezados": celdas, "filas": [],
                                "_abierta": True})
            continue
        if bloques and bloques[-1].get("tipo") == "tabla" and bloques[-1].get("_abierta"):
            bloques[-1].pop("_abierta", None)
        _cerrar_lista()
        parrafo.append(cruda.strip())

    _cerrar_parrafo()
    _cerrar_lista()
    for b in bloques:
        b.pop("_abierta", None)
    return bloques


def _validar_documento(cabecera: Dict[str, Any], bloques: List[Dict[str, Any]],
                       con_reemplazos: bool) -> None:
    """Un documento tiene que tener contenido de verdad. Levanta si no."""
    if con_reemplazos and not bloques:
        return                                 # rellenar una plantilla es contenido en sí
    prosa = 0
    filas = 0
    for b in bloques:
        tipo = b.get("tipo", "parrafo")
        if tipo in ("parrafo", "cita", "titulo"):
            prosa += len(str(b.get("texto") or ""))
        elif tipo == "lista":
            prosa += sum(len(str(i)) for i in b.get("items") or [])
        elif tipo == "tabla":
            filas += len(b.get("filas") or [])
        elif tipo == "grafico":
            filas += len(_series_de(b)[0])
    if not bloques or (prosa < _MINIMO_CARACTERES_DOCUMENTO and filas < 2):
        raise DocumentoRechazado(
            "El documento no tiene contenido suficiente: hacen falta las secciones con su "
            "texto completo, no solo el título. Escribí el documento entero y volvé a pedirlo."
        )


def _reemplazar_en_parrafo(parrafo, reemplazos: Dict[str, str]) -> int:
    """Sustituye marcadores en un párrafo de Word conservando el formato. Return cuántos.

    Word parte el texto de un párrafo en "runs" según el formato, y un marcador como
    `{{cliente}}` puede quedar repartido en tres. Se sustituye sobre el texto completo y se
    vuelve a escribir en el primer run, que conserva su fuente y su tamaño; los demás se
    vacían. Es lo que hace que la cotización rellenada siga pareciendo la cotización.
    """
    texto = parrafo.text
    if not any(marca in texto for marca in reemplazos):
        return 0
    cuantos = 0
    for marca, valor in reemplazos.items():
        if marca in texto:
            cuantos += texto.count(marca)
            texto = texto.replace(marca, str(valor))
    if parrafo.runs:
        parrafo.runs[0].text = texto
        for run in parrafo.runs[1:]:
            run.text = ""
    else:
        parrafo.add_run(texto)
    return cuantos


def _reemplazar_en_docx(doc, reemplazos: Dict[str, str]) -> int:
    """Return cuántos marcadores se sustituyeron en cuerpo, tablas, encabezados y pies."""
    cuantos = 0
    contenedores = [doc]
    for seccion in doc.sections:
        contenedores += [seccion.header, seccion.footer, seccion.first_page_header,
                         seccion.first_page_footer]
    for contenedor in contenedores:
        try:
            for parrafo in contenedor.paragraphs:
                cuantos += _reemplazar_en_parrafo(parrafo, reemplazos)
            for tabla in contenedor.tables:
                for fila in tabla.rows:
                    for celda in fila.cells:
                        for parrafo in celda.paragraphs:
                            cuantos += _reemplazar_en_parrafo(parrafo, reemplazos)
        except Exception as e:
            logger.debug(f"no se pudo recorrer una parte de la plantilla: {e}")
    return cuantos


def _marcadores_restantes_docx(doc) -> List[Tuple[str, str]]:
    """Return `[(marcador, dónde)]` de lo que quedó sin rellenar en cuerpo, tablas,
    encabezados y pies."""
    hallados: Dict[str, str] = {}

    def _revisar(contenedor, donde: str) -> None:
        try:
            textos = [p.text for p in contenedor.paragraphs]
            for tabla in contenedor.tables:
                for fila in tabla.rows:
                    for celda in fila.cells:
                        textos += [p.text for p in celda.paragraphs]
        except Exception as e:
            logger.debug(f"no se pudo revisar {donde}: {e}")
            return
        for marcador in _marcadores_en("\n".join(textos)):
            hallados.setdefault(marcador, donde)

    _revisar(doc, "cuerpo")
    for seccion in doc.sections:
        _revisar(seccion.header, "encabezado")
        _revisar(seccion.footer, "pie de página")
        _revisar(seccion.first_page_header, "encabezado de primera página")
        _revisar(seccion.first_page_footer, "pie de primera página")
    return sorted(hallados.items())


def _marcadores_restantes_pptx(prs) -> List[Tuple[str, str]]:
    """Return `[(marcador, dónde)]` de lo que quedó sin rellenar en las diapositivas."""
    hallados: Dict[str, str] = {}
    for numero, slide in enumerate(prs.slides, 1):
        textos = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                textos.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for fila in shape.table.rows:
                    for celda in fila.cells:
                        textos.append(celda.text_frame.text)
        for marcador in _marcadores_en("\n".join(textos)):
            hallados.setdefault(marcador, f"diapositiva {numero}")
    return sorted(hallados.items())


def _vaciar_cuerpo(doc) -> None:
    """Quita el contenido del cuerpo de una plantilla, dejando encabezados, pies y estilos."""
    cuerpo = doc.element.body
    for hijo in list(cuerpo):
        if hijo.tag.endswith("}sectPr"):
            continue                           # la configuración de página se queda
        cuerpo.remove(hijo)


def _estilo(doc, *candidatos: str):
    """Return el primer estilo que exista en el documento de entre los candidatos, o None."""
    for nombre in candidatos:
        try:
            return doc.styles[nombre]
        except KeyError:
            continue
    return None


# ================================================================ estilo guiado por el usuario

#: Colores que la gente nombra al pedir un documento. Un hex («#2E7D32», «2E7D32») se
#: acepta tal cual; cualquier otra cosa se avisa y se ignora, sin tumbar el documento.
_COLORES_POR_NOMBRE = {
    "negro": "000000", "blanco": "FFFFFF", "gris": "7F7F7F", "gris oscuro": "404040",
    "gris claro": "BFBFBF", "azul": "2F5496", "azul oscuro": "1F3864", "azul claro": "5B9BD5",
    "azul marino": "1F3864", "verde": "2E7D32", "verde oscuro": "1B5E20", "verde claro": "70AD47",
    "rojo": "C00000", "rojo oscuro": "8B0000", "naranja": "ED7D31", "amarillo": "FFC000",
    "morado": "7030A0", "violeta": "7030A0", "purpura": "7030A0", "cafe": "7B3F00",
    "marron": "7B3F00", "dorado": "BF9000", "turquesa": "00B0F0", "celeste": "9DC3E6",
    "vinotinto": "722F37", "rosa": "E91E63", "rosado": "E91E63",
}

#: Sinónimos que un modelo (o un usuario) usa para las claves de `estilo`.
_ALIAS_ESTILO = {
    "tamaño": "tamano", "tamano_letra": "tamano", "size": "tamano", "font": "fuente",
    "tipografia": "fuente", "tipografía": "fuente", "letra": "fuente",
    "margenes": "margenes_cm", "márgenes": "margenes_cm", "margenes_cm": "margenes_cm",
    "orientación": "orientacion", "color_titulo": "color_titulos",
    "color_encabezados": "color_encabezado",
    "encabezado_color": "color_encabezado", "numeracion_paginas": "numeracion",
    "numerar": "numeracion", "numero_de_pagina": "numeracion", "numero_pagina": "numeracion",
    "justificar": "justificado", "interlineado": "interlineado", "footer": "pie",
    "header": "encabezado", "logo_ancho": "logo_ancho_cm", "fuente_titulo": "fuente_titulos",
}

_CLAVES_DE_COLOR = ("color_titulos", "color_texto", "color_encabezado", "color_texto_encabezado")


def _color_hex(valor: Any) -> Optional[str]:
    """Return el color como «RRGGBB», o None si no se entiende."""
    texto = _sin_tildes(str(valor or "")).strip().lower()
    if not texto:
        return None
    if texto in _COLORES_POR_NOMBRE:
        return _COLORES_POR_NOMBRE[texto]
    m = re.fullmatch(r"#?([0-9a-f]{6})", texto)
    if m:
        return m.group(1).upper()
    m = re.fullmatch(r"#?([0-9a-f])([0-9a-f])([0-9a-f])", texto)
    if m:
        return "".join(c * 2 for c in m.groups()).upper()
    return None


def _estilo_de(spec: Any, avisos: List[str]) -> Dict[str, Any]:
    """Return el `estilo` que pidió el usuario, con claves canónicas y colores en hex.

    Lo que no se entiende se avisa y se deja fuera: un color mal escrito no puede impedir
    que salga el informe.
    """
    if not spec:
        return {}
    spec = _como_estructura(spec)
    if not isinstance(spec, dict):
        avisos.append("El 'estilo' tiene que ser un objeto (fuente, tamano, color_titulos...); "
                      "se ignoró.")
        return {}
    estilo: Dict[str, Any] = {}
    for clave, valor in spec.items():
        canonica = _ALIAS_ESTILO.get(str(clave).strip().lower(), str(clave).strip().lower())
        if valor is None or valor == "":
            continue
        if canonica in _CLAVES_DE_COLOR:
            hexa = _color_hex(valor)
            if hexa is None:
                avisos.append(f"No entendí el color «{valor}» de '{canonica}'; "
                              f"se dejó el de siempre.")
                continue
            estilo[canonica] = hexa
        elif canonica == "tamano":
            try:
                estilo[canonica] = float(str(valor).replace(",", ".").replace("pt", "").strip())
            except ValueError:
                avisos.append(f"No entendí el tamaño de letra «{valor}»; se dejó el de siempre.")
        elif canonica == "orientacion":
            texto = _sin_tildes(str(valor)).lower()
            apaisado = texto in ("horizontal", "apaisado", "apaisada", "landscape")
            estilo[canonica] = "horizontal" if apaisado else "vertical"
        elif canonica in ("numeracion", "justificado", "bordes"):
            estilo[canonica] = valor if isinstance(valor, bool) else \
                _sin_tildes(str(valor)).strip().lower() in ("1", "si", "true", "verdadero", "yes")
        else:
            estilo[canonica] = valor
    return estilo


def _margenes_de(valor: Any) -> Dict[str, float]:
    """Return {superior, inferior, izquierdo, derecho} en cm a partir de un número o un objeto."""
    if isinstance(valor, (int, float)):
        return {k: float(valor) for k in ("superior", "inferior", "izquierdo", "derecho")}
    if isinstance(valor, str):
        try:
            return _margenes_de(float(valor.replace(",", ".").replace("cm", "").strip()))
        except ValueError:
            return {}
    if isinstance(valor, dict):
        alias = {"top": "superior", "arriba": "superior", "bottom": "inferior", "abajo": "inferior",
                 "left": "izquierdo", "izquierda": "izquierdo", "right": "derecho",
                 "derecha": "derecho"}
        salida = {}
        for k, v in valor.items():
            clave = alias.get(str(k).lower(), str(k).lower())
            try:
                salida[clave] = float(str(v).replace(",", ".").replace("cm", "").strip())
            except ValueError:
                continue
        return salida
    return {}


def _campo_de_word(parrafo, instruccion: str) -> None:
    """Añade un campo de Word (`PAGE`, `NUMPAGES`) al párrafo. python-docx no tiene API."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = parrafo.add_run()
    inicio = OxmlElement("w:fldChar")
    inicio.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruccion} "
    separador = OxmlElement("w:fldChar")
    separador.set(qn("w:fldCharType"), "separate")
    texto = OxmlElement("w:t")
    texto.text = "1"
    fin = OxmlElement("w:fldChar")
    fin.set(qn("w:fldCharType"), "end")
    for elemento in (inicio, instr, separador, texto, fin):
        run._r.append(elemento)


def _aplicar_estilo_docx(doc, estilo: Dict[str, Any], avisos: List[str]) -> List[str]:
    """Aplica al documento lo que el usuario pidió de aspecto. Return qué se aplicó.

    Va sobre los estilos del documento (`Normal`, `Title`, `Heading N`) y sus secciones,
    no párrafo por párrafo: así vale para lo que se escribe después y también encima de
    una plantilla, que es lo que se pidió explícitamente cuando vienen las dos cosas.
    """
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    if not estilo:
        return []
    aplicado: List[str] = []
    normal = _estilo(doc, "Normal")

    if estilo.get("fuente") and normal is not None:
        normal.font.name = str(estilo["fuente"])
        rpr = normal.element.get_or_add_rPr()
        fuentes = rpr.find(qn("w:rFonts"))
        if fuentes is not None:
            for atributo in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                fuentes.set(qn(atributo), str(estilo["fuente"]))
        aplicado.append(f"fuente {estilo['fuente']}")
    if estilo.get("tamano") and normal is not None:
        normal.font.size = Pt(float(estilo["tamano"]))
        aplicado.append(f"{estilo['tamano']:g} pt")
    if estilo.get("color_texto") and normal is not None:
        normal.font.color.rgb = RGBColor.from_string(estilo["color_texto"])
    if estilo.get("justificado") and normal is not None:
        normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        aplicado.append("justificado")
    if estilo.get("interlineado") and normal is not None:
        try:
            interlineado = float(str(estilo["interlineado"]).replace(",", "."))
            normal.paragraph_format.line_spacing = interlineado
        except ValueError:
            avisos.append(f"No entendí el interlineado «{estilo['interlineado']}».")

    if estilo.get("fuente_titulos") or estilo.get("color_titulos"):
        for nombre in ("Title", "Título", "Subtitle", "Subtítulo", "Heading 1", "Heading 2",
                       "Heading 3", "Heading 4", "Título 1", "Título 2", "Título 3", "Título 4"):
            s = _estilo(doc, nombre)
            if s is None:
                continue
            if estilo.get("fuente_titulos"):
                s.font.name = str(estilo["fuente_titulos"])
                rpr = s.element.get_or_add_rPr()
                fuentes = rpr.find(qn("w:rFonts"))
                if fuentes is not None:
                    for atributo in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                        fuentes.set(qn(atributo), str(estilo["fuente_titulos"]))
                    # Los títulos de Word heredan la fuente del tema; ese atributo manda.
                    del_tema = ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme")
                    for atributo in del_tema:
                        if fuentes.get(qn(atributo)) is not None:
                            del fuentes.attrib[qn(atributo)]
            if estilo.get("color_titulos"):
                s.font.color.rgb = RGBColor.from_string(estilo["color_titulos"])
        aplicado.append("títulos " + " ".join(
            str(estilo[k]) for k in ("fuente_titulos", "color_titulos") if estilo.get(k)))

    margenes = _margenes_de(estilo["margenes_cm"]) if estilo.get("margenes_cm") is not None else {}
    orientacion = estilo.get("orientacion")
    for seccion in doc.sections:
        for lado, atributo in (("superior", "top_margin"), ("inferior", "bottom_margin"),
                               ("izquierdo", "left_margin"), ("derecho", "right_margin")):
            if lado in margenes:
                setattr(seccion, atributo, Cm(margenes[lado]))
        if orientacion == "horizontal" and seccion.orientation != WD_ORIENT.LANDSCAPE:
            seccion.orientation = WD_ORIENT.LANDSCAPE
            seccion.page_width, seccion.page_height = seccion.page_height, seccion.page_width
        elif orientacion == "vertical" and seccion.orientation == WD_ORIENT.LANDSCAPE:
            seccion.orientation = WD_ORIENT.PORTRAIT
            seccion.page_width, seccion.page_height = seccion.page_height, seccion.page_width
    if margenes:
        aplicado.append("márgenes")
    if orientacion:
        aplicado.append(orientacion)

    seccion = doc.sections[0]
    if estilo.get("encabezado"):
        seccion.header.is_linked_to_previous = False
        encabezado = seccion.header
        parrafo = encabezado.paragraphs[0] if encabezado.paragraphs else encabezado.add_paragraph()
        parrafo.text = str(estilo["encabezado"])
        aplicado.append("encabezado")
    if estilo.get("logo"):
        ruta_logo = str(estilo["logo"])
        if os.path.isfile(ruta_logo):
            try:
                seccion.header.is_linked_to_previous = False
                parrafo = seccion.header.add_paragraph()
                posicion = _sin_tildes(str(estilo.get("logo_posicion") or "derecha")).lower()
                alineaciones = {"izquierda": WD_ALIGN_PARAGRAPH.LEFT,
                                "centro": WD_ALIGN_PARAGRAPH.CENTER}
                parrafo.alignment = alineaciones.get(posicion, WD_ALIGN_PARAGRAPH.RIGHT)
                ancho = float(str(estilo.get("logo_ancho_cm") or 3).replace(",", "."))
                parrafo.add_run().add_picture(ruta_logo, width=Cm(ancho))
                # Va primero en el encabezado: arriba del texto, no debajo.
                seccion.header._element.insert(0, parrafo._p)
                aplicado.append("logo")
            except Exception as e:
                avisos.append(f"No pude poner el logo «{os.path.basename(ruta_logo)}»: {e}.")
        else:
            avisos.append(f"No encontré el logo «{ruta_logo}»; el documento salió sin él.")
    if estilo.get("pie"):
        seccion.footer.is_linked_to_previous = False
        pie = seccion.footer
        parrafo = pie.paragraphs[0] if pie.paragraphs else pie.add_paragraph()
        parrafo.text = str(estilo["pie"])
        aplicado.append("pie")
    if estilo.get("numeracion"):
        seccion.footer.is_linked_to_previous = False
        parrafo = seccion.footer.add_paragraph()
        parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        parrafo.add_run("Página ")
        _campo_de_word(parrafo, "PAGE")
        parrafo.add_run(" de ")
        _campo_de_word(parrafo, "NUMPAGES")
        aplicado.append("número de página")
    return aplicado


# ================================================================================ gráficos

#: Paleta sobria para los gráficos que se dibujan como imagen (Word, PDF sin Office).
_COLORES_GRAFICO = ("#305496", "#C55A11", "#548235", "#7030A0", "#BF9000", "#2E75B6")


def _series_de(grafico: Dict[str, Any]) -> Tuple[List[str], List[Tuple[str, List[float]]]]:
    """Return `(categorías, [(nombre de serie, valores)])` de un bloque de gráfico.

    Se aceptan dos formas: `series` con `nombre` y `valores`, o un par `etiquetas`/`valores`
    para un gráfico de una sola serie. Los valores se leen como números a la colombiana.
    """
    categorias = [str(c) for c in (grafico.get("categorias") or grafico.get("etiquetas") or [])]
    series: List[Tuple[str, List[float]]] = []
    crudas = grafico.get("series")
    if isinstance(crudas, list) and crudas:
        for s in crudas:
            if isinstance(s, dict):
                valores = [_valor_de_celda(v) for v in (s.get("valores") or [])]
                series.append((str(s.get("nombre") or ""), [v if isinstance(v, (int, float)) else 0
                                                             for v in valores]))
    elif grafico.get("valores"):
        valores = [_valor_de_celda(v) for v in grafico["valores"]]
        series.append((str(grafico.get("titulo") or ""), [v if isinstance(v, (int, float)) else 0
                                                          for v in valores]))
    if not categorias and series:
        categorias = [str(i + 1) for i in range(len(series[0][1]))]
    return categorias, series


def _grafico_como_imagen(grafico: Dict[str, Any], destino: str) -> bool:
    """Dibuja el gráfico en un PNG con matplotlib. Return si se pudo.

    Para Word y para el PDF sin Office: python-docx no sabe crear gráficos nativos, así que
    un gráfico en Word es una imagen. Se dibuja limpio —sin marco arriba ni a la derecha,
    cuadrícula suave, etiquetas de valor— para que no parezca sacado de una consola.
    """
    categorias, series = _series_de(grafico)
    if not series:
        return False
    try:
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError:
        logger.info("matplotlib no está instalado: no se puede dibujar el gráfico")
        return False

    tipo = str(grafico.get("tipo_grafico") or grafico.get("tipo") or "barras").lower()
    conv = convencion()
    fig, ax = plt.subplots(figsize=(7.5, 4.2), dpi=150)
    try:
        if tipo in ("torta", "pastel", "pie"):
            nombre, valores = series[0]
            ax.pie(valores, labels=categorias, colors=_COLORES_GRAFICO[:len(valores)],
                   autopct=lambda p: f"{p:.0f} %" if p >= 4 else "", startangle=90,
                   wedgeprops={"linewidth": 1, "edgecolor": "white"})
            ax.axis("equal")
        elif tipo in ("lineas", "línea", "linea", "line"):
            for i, (nombre, valores) in enumerate(series):
                ax.plot(categorias[:len(valores)], valores, marker="o", linewidth=2,
                        color=_COLORES_GRAFICO[i % len(_COLORES_GRAFICO)], label=nombre or None)
        else:
            ancho = 0.8 / max(1, len(series))
            for i, (nombre, valores) in enumerate(series):
                posiciones = [x + i * ancho for x in range(len(valores))]
                barras = ax.bar(posiciones, valores, width=ancho,
                                color=_COLORES_GRAFICO[i % len(_COLORES_GRAFICO)], label=nombre or None)
                if len(valores) <= 12:
                    ax.bar_label(barras, fmt=lambda v: conv.formatear(v), fontsize=8, padding=2)
            ax.set_xticks([x + ancho * (len(series) - 1) / 2 for x in range(len(categorias))])
            ax.set_xticklabels(categorias)
        if tipo not in ("torta", "pastel", "pie"):
            from matplotlib.ticker import FuncFormatter

            ax.spines["top"].set_visible(False)
            ax.spines["right"].set_visible(False)
            ax.grid(axis="y", color="#DDDDDD", linewidth=0.8)
            ax.set_axisbelow(True)
            ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _p: conv.formatear(v)))
            if grafico.get("eje_y"):
                ax.set_ylabel(str(grafico["eje_y"]))
            if len(series) > 1 or any(n for n, _ in series):
                ax.legend(frameon=False)
        if grafico.get("titulo"):
            ax.set_title(str(grafico["titulo"]), fontsize=12, loc="left", pad=12)
        fig.tight_layout()
        fig.savefig(destino, dpi=150)
        return True
    except Exception as e:
        logger.warning(f"no se pudo dibujar el gráfico: {e}")
        return False
    finally:
        plt.close(fig)


def _grafico_nativo_pptx(slide, grafico: Dict[str, Any], x, y, ancho, alto) -> bool:
    """Pone un gráfico NATIVO de PowerPoint —editable por el usuario— en la diapositiva."""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

    categorias, series = _series_de(grafico)
    if not series:
        return False
    tipo = str(grafico.get("tipo_grafico") or grafico.get("tipo") or "barras").lower()
    clase = {"lineas": XL_CHART_TYPE.LINE_MARKERS, "linea": XL_CHART_TYPE.LINE_MARKERS,
             "línea": XL_CHART_TYPE.LINE_MARKERS, "line": XL_CHART_TYPE.LINE_MARKERS,
             "torta": XL_CHART_TYPE.PIE, "pastel": XL_CHART_TYPE.PIE, "pie": XL_CHART_TYPE.PIE,
             }.get(tipo, XL_CHART_TYPE.COLUMN_CLUSTERED)
    datos = CategoryChartData()
    datos.categories = categorias
    for nombre, valores in series:
        datos.add_series(nombre or "Serie", valores)
    try:
        forma = slide.shapes.add_chart(clase, x, y, ancho, alto, datos)
        chart = forma.chart
        if grafico.get("titulo"):
            chart.has_title = True
            chart.chart_title.text_frame.text = str(grafico["titulo"])
        chart.has_legend = len(series) > 1 or clase == XL_CHART_TYPE.PIE
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
        if clase == XL_CHART_TYPE.PIE:
            plot = chart.plots[0]
            plot.has_data_labels = True
            plot.data_labels.show_percentage = True
            plot.data_labels.show_value = False
            plot.data_labels.number_format = "0%"
            plot.data_labels.number_format_is_linked = False
        else:
            # Separador de miles en el eje: PowerPoint lo formatea según la región del
            # equipo que abra el archivo, así que "#,##0" sale 1.250.000 en Colombia.
            chart.value_axis.tick_labels.number_format = "#,##0"
            chart.value_axis.tick_labels.number_format_is_linked = False
        return True
    except Exception as e:
        logger.warning(f"no se pudo crear el gráfico en la diapositiva: {e}")
        return False


def _escribir_bloques(doc, bloques: List[Dict[str, Any]], avisos: List[str]) -> None:
    """Escribe los bloques al final del documento con sus estilos. Lo usan crear y editar."""
    from docx.shared import Cm, Pt

    for bloque in bloques:
        tipo = str(bloque.get("tipo") or "parrafo").lower()
        if tipo == "titulo":
            nivel = max(1, min(int(bloque.get("nivel") or 1), 4))
            texto = str(bloque.get("texto") or "")
            try:
                doc.add_heading(texto, level=nivel)
            except KeyError:
                p = doc.add_paragraph(texto)
                p.runs[0].bold = True
                p.runs[0].font.size = Pt(18 - 2 * nivel)
        elif tipo == "lista":
            numerada = bool(bloque.get("numerada"))
            estilo = _estilo(doc, "List Number" if numerada else "List Bullet",
                             "Lista con números" if numerada else "Lista con viñetas")
            for item in bloque.get("items") or []:
                if isinstance(item, (list, tuple)):
                    for sub in item:
                        p = doc.add_paragraph(str(sub), style=estilo)
                        p.paragraph_format.left_indent = Cm(1.25)
                    continue
                if estilo is not None:
                    doc.add_paragraph(str(item), style=estilo)
                else:
                    doc.add_paragraph(("• " if not numerada else "") + str(item))
        elif tipo == "tabla":
            encabezados = [str(c) for c in (bloque.get("encabezados") or [])]
            filas = [[("" if c is None else str(c)) for c in fila]
                     for fila in (bloque.get("filas") or [])]
            columnas = max([len(encabezados)] + [len(f) for f in filas]) if (encabezados or filas) else 0
            if not columnas:
                continue
            if bloque.get("titulo"):
                doc.add_paragraph(str(bloque["titulo"]),
                                  style=_estilo(doc, "Caption", "Descripción"))
            tabla = doc.add_table(rows=0, cols=columnas)
            estilo_tabla = _estilo(doc, "Table Grid", "Tabla con cuadrícula")
            if estilo_tabla is not None:
                tabla.style = estilo_tabla
            if encabezados:
                celdas = tabla.add_row().cells
                for i, texto in enumerate(encabezados):
                    celdas[i].text = texto
                    for p in celdas[i].paragraphs:
                        for r in p.runs:
                            r.bold = True
            for fila in filas:
                celdas = tabla.add_row().cells
                for i, texto in enumerate(fila[:columnas]):
                    celdas[i].text = texto
            doc.add_paragraph()
        elif tipo == "imagen":
            ruta_imagen = str(bloque.get("ruta") or "")
            if not os.path.isfile(ruta_imagen):
                avisos.append(f"No encontré la imagen «{ruta_imagen}»; el documento se hizo sin ella.")
                continue
            try:
                ancho = float(bloque.get("ancho_cm") or _ANCHO_IMAGEN_CM)
                doc.add_picture(ruta_imagen, width=Cm(ancho))
                if bloque.get("pie"):
                    doc.add_paragraph(str(bloque["pie"]), style=_estilo(doc, "Caption", "Descripción"))
            except Exception as e:
                avisos.append(f"No pude insertar la imagen «{os.path.basename(ruta_imagen)}»: {e}.")
        elif tipo == "salto":
            doc.add_page_break()
        elif tipo == "grafico":
            # python-docx no crea gráficos nativos: en Word, un gráfico es una imagen. Se
            # dibuja con matplotlib y se inserta con su pie.
            with tempfile.TemporaryDirectory() as temporal:
                png = os.path.join(temporal, "grafico.png")
                if _grafico_como_imagen(bloque, png):
                    doc.add_picture(png, width=Cm(float(bloque.get("ancho_cm") or 15)))
                    if bloque.get("pie"):
                        doc.add_paragraph(str(bloque["pie"]), style=_estilo(doc, "Caption", "Descripción"))
                else:
                    avisos.append("Un gráfico no traía datos (hacen falta 'categorias' y 'series') "
                                  "y se omitió.")
        elif tipo == "cita":
            doc.add_paragraph(str(bloque.get("texto") or ""),
                              style=_estilo(doc, "Intense Quote", "Quote", "Cita destacada", "Cita"))
        else:
            p = doc.add_paragraph(str(bloque.get("texto") or ""))
            if bloque.get("negrita"):
                for r in p.runs:
                    r.bold = True



def _docx(cabecera: Dict[str, Any], bloques: List[Dict[str, Any]], ruta: Path,
          plantilla: Optional[Path], reemplazos: Dict[str, str],
          estilo: Optional[Dict[str, Any]] = None) -> Resultado:
    from docx import Document
    from docx.shared import Cm, Pt

    doc = Document(str(plantilla)) if plantilla else Document()
    avisos: List[str] = []

    sustituidos = _reemplazar_en_docx(doc, reemplazos) if reemplazos else 0
    if reemplazos and not sustituidos:
        avisos.append("Ojo: ninguno de los marcadores que me diste aparece en la plantilla; "
                      "revisá cómo están escritos en ella con document_inspect.")
    if plantilla and bloques:
        _vaciar_cuerpo(doc)
    if plantilla:
        # Un marcador que quedó sin rellenar —en el pie, en una tabla— sale impreso y el
        # archivo parece hecho sin estarlo. Se rechaza antes de guardar: el modelo puede
        # volver con `reemplazos` o preguntarle al usuario el dato que falta.
        restantes = _marcadores_restantes_docx(doc)
        if restantes:
            raise DocumentoRechazado(
                "La plantilla quedó con marcadores sin rellenar: "
                + "; ".join(f"{m} ({donde})" for m, donde in restantes)
                + ". Pasá su valor en 'reemplazos' (o preguntale al usuario) y volvé a pedirlo."
                + (" " + avisos[-1] if reemplazos and not sustituidos else "")
            )
    aplicado = _aplicar_estilo_docx(doc, estilo or {}, avisos)

    if cabecera.get("titulo"):
        p = doc.add_paragraph(str(cabecera["titulo"]), style=_estilo(doc, "Title", "Título"))
        if p.style is None or p.style.name not in ("Title", "Título"):
            p.runs[0].bold = True
            p.runs[0].font.size = Pt(24)
    if cabecera.get("subtitulo"):
        doc.add_paragraph(str(cabecera["subtitulo"]), style=_estilo(doc, "Subtitle", "Subtítulo"))
    pie_portada = " · ".join(str(cabecera[k]) for k in ("autor", "fecha") if cabecera.get(k))
    if pie_portada:
        doc.add_paragraph(pie_portada)

    _escribir_bloques(doc, bloques, avisos)

    doc.save(str(ruta))
    detalle = f"{len(bloques)} bloques"
    if sustituidos:
        detalle += f", {sustituidos} marcadores rellenados"
    if plantilla:
        detalle += f", sobre la plantilla {plantilla.name}"
    if aplicado:
        detalle += ", estilo: " + ", ".join(aplicado)
    return Resultado(ruta=str(ruta), formato="docx", detalle=detalle, avisos=avisos)


def _html_de(cabecera: Dict[str, Any], bloques: List[Dict[str, Any]],
             estilo: Optional[Dict[str, Any]] = None) -> str:
    """Return el documento como HTML sencillo y limpio. Sirve para .html y para el PDF sin Office."""
    import html as _html

    estilo = estilo or {}
    fuente = _html.escape(str(estilo.get("fuente") or "Calibri"))
    tamano = float(estilo.get("tamano") or 11)
    color_texto = "#" + (estilo.get("color_texto") or "222222")
    color_titulos = "#" + estilo["color_titulos"] if estilo.get("color_titulos") else "inherit"
    fuente_titulos = "inherit"
    if estilo.get("fuente_titulos"):
        fuente_titulos = f"'{_html.escape(str(estilo['fuente_titulos']))}'"
    partes = ["<!doctype html><html><head><meta charset='utf-8'>",
              f"<style>body{{font-family:'{fuente}',Calibri,Arial,sans-serif;"
              f"font-size:{tamano:g}pt;line-height:1.4;max-width:17cm;margin:2cm auto;"
              f"color:{color_texto}}}"
              f"h1,h2,h3{{color:{color_titulos};font-family:{fuente_titulos}}}"
              "h1{font-size:20pt}h2{font-size:15pt}"
              "h3{font-size:12.5pt}table{border-collapse:collapse;margin:8pt 0}"
              "td,th{border:1px solid #999;padding:4pt 6pt;vertical-align:top}th{background:#eee}"
              "blockquote{border-left:3px solid #999;margin:8pt 0;padding-left:8pt;color:#444}"
              ".portada{color:#555}</style></head><body>"]
    e = _html.escape
    if cabecera.get("titulo"):
        partes.append(f"<h1>{e(str(cabecera['titulo']))}</h1>")
    if cabecera.get("subtitulo"):
        partes.append(f"<p class='portada'><b>{e(str(cabecera['subtitulo']))}</b></p>")
    pie = " · ".join(str(cabecera[k]) for k in ("autor", "fecha") if cabecera.get(k))
    if pie:
        partes.append(f"<p class='portada'>{e(pie)}</p>")
    for b in bloques:
        tipo = str(b.get("tipo") or "parrafo").lower()
        if tipo == "titulo":
            n = max(1, min(int(b.get("nivel") or 1), 4)) + 1
            partes.append(f"<h{n}>{e(str(b.get('texto') or ''))}</h{n}>")
        elif tipo == "lista":
            etiqueta = "ol" if b.get("numerada") else "ul"
            items = "".join(f"<li>{e(str(i))}</li>" for i in (b.get("items") or []))
            partes.append(f"<{etiqueta}>{items}</{etiqueta}>")
        elif tipo == "tabla":
            enc = "".join(f"<th>{e(str(c))}</th>" for c in (b.get("encabezados") or []))
            filas = "".join("<tr>" + "".join(f"<td>{e('' if c is None else str(c))}</td>"
                                             for c in f) + "</tr>"
                            for f in (b.get("filas") or []))
            titulo = f"<p><i>{e(str(b['titulo']))}</i></p>" if b.get("titulo") else ""
            partes.append(f"{titulo}<table>{'<tr>' + enc + '</tr>' if enc else ''}{filas}</table>")
        elif tipo == "cita":
            partes.append(f"<blockquote>{e(str(b.get('texto') or ''))}</blockquote>")
        elif tipo == "salto":
            partes.append("<div style='page-break-after:always'></div>")
        elif tipo == "grafico":
            import base64

            with tempfile.TemporaryDirectory() as temporal:
                png = os.path.join(temporal, "grafico.png")
                if _grafico_como_imagen(b, png):
                    datos = base64.b64encode(Path(png).read_bytes()).decode("ascii")
                    partes.append(f"<img src='data:image/png;base64,{datos}' style='max-width:100%'>")
                    if b.get("pie"):
                        partes.append(f"<p><i>{e(str(b['pie']))}</i></p>")
        elif tipo == "imagen":
            if os.path.isfile(str(b.get("ruta") or "")):
                partes.append(f"<img src='file:///{e(str(b['ruta']))}' style='max-width:100%'>")
        else:
            partes.append(f"<p>{e(str(b.get('texto') or ''))}</p>")
    partes.append("</body></html>")
    return "".join(partes)


def _markdown_de(cabecera: Dict[str, Any], bloques: List[Dict[str, Any]]) -> str:
    lineas: List[str] = []
    if cabecera.get("titulo"):
        lineas += [f"# {cabecera['titulo']}", ""]
    if cabecera.get("subtitulo"):
        lineas += [f"**{cabecera['subtitulo']}**", ""]
    pie = " · ".join(str(cabecera[k]) for k in ("autor", "fecha") if cabecera.get(k))
    if pie:
        lineas += [f"*{pie}*", ""]
    for b in bloques:
        tipo = str(b.get("tipo") or "parrafo").lower()
        if tipo == "titulo":
            nivel = max(1, min(int(b.get("nivel") or 1), 4)) + 1
            lineas += ["#" * nivel + " " + str(b.get("texto") or ""), ""]
        elif tipo == "lista":
            for i, item in enumerate(b.get("items") or [], 1):
                lineas.append((f"{i}. " if b.get("numerada") else "- ") + str(item))
            lineas.append("")
        elif tipo == "tabla":
            enc = [str(c) for c in (b.get("encabezados") or [])]
            filas = b.get("filas") or []
            if b.get("titulo"):
                lineas += [f"*{b['titulo']}*", ""]
            if enc:
                lineas.append("| " + " | ".join(enc) + " |")
                lineas.append("|" + "---|" * len(enc))
            for f in filas:
                lineas.append("| " + " | ".join("" if c is None else str(c) for c in f) + " |")
            lineas.append("")
        elif tipo == "cita":
            lineas += ["> " + str(b.get("texto") or ""), ""]
        elif tipo == "salto":
            lineas += ["---", ""]
        elif tipo == "grafico":
            categorias, series = _series_de(b)
            if series:
                if b.get("titulo"):
                    lineas += [f"*{b['titulo']}*", ""]
                lineas.append("| " + " | ".join([""] + [n or "Serie" for n, _ in series]) + " |")
                lineas.append("|" + "---|" * (len(series) + 1))
                for i, c in enumerate(categorias):
                    lineas.append("| " + " | ".join([c] + [str(v[i]) if i < len(v) else ""
                                                         for _, v in series]) + " |")
                lineas.append("")
        elif tipo == "imagen":
            lineas += [f"![{b.get('pie') or ''}]({b.get('ruta') or ''})", ""]
        else:
            lineas += [str(b.get("texto") or ""), ""]
    return "\n".join(lineas).rstrip() + "\n"


def _texto_plano_de(cabecera: Dict[str, Any], bloques: List[Dict[str, Any]]) -> str:
    md = _markdown_de(cabecera, bloques)
    return re.sub(r"^#+\s*", "", md, flags=re.MULTILINE).replace("**", "").replace("*", "")


def _pdf_sin_office(html: str, ruta: Path) -> None:
    """Maqueta el HTML en un PDF con PyMuPDF. Es el respaldo cuando no hay Word."""
    import pymupdf

    historia = pymupdf.Story(html=html)
    escritor = pymupdf.DocumentWriter(str(ruta))
    rect_pagina = pymupdf.paper_rect("a4")
    area = rect_pagina + (56, 56, -56, -56)         # 2 cm de margen
    hay_mas = True
    while hay_mas:
        dispositivo = escritor.begin_page(rect_pagina)
        hay_mas, _ = historia.place(area)
        historia.draw(dispositivo)
        escritor.end_page()
    escritor.close()


def crear_documento(nombre: str, contenido: Any, plantilla: Optional[str] = None,
                    reemplazos: Optional[Dict[str, str]] = None,
                    carpeta: Optional[str] = None, estilo: Any = None) -> Resultado:
    """Produce un documento de texto: .docx, .pdf, .md, .html o .txt. Levanta `DocumentoRechazado`.

    `contenido` es un objeto con `titulo`, `subtitulo`, `autor`, `fecha` y `bloques` —o un
    texto en Markdown, que se convierte—. Cada bloque tiene `tipo`: `titulo` (con `nivel`),
    `parrafo`, `lista` (`items`, `numerada`), `tabla` (`encabezados`, `filas`, `titulo`),
    `imagen` (`ruta`, `ancho_cm`, `pie`), `cita`, `salto`. `estilo` es lo que el usuario
    pidió de aspecto: `fuente`, `tamano`, `fuente_titulos`, `color_titulos`, `color_texto`,
    `margenes_cm`, `orientacion`, `encabezado`, `pie`, `logo`, `numeracion`, `justificado`.
    """
    extension = _extension(nombre)
    if extension not in EXTENSIONES_DOCUMENTO:
        raise DocumentoRechazado(
            f"«{nombre}» no es un documento de texto: hace falta "
            f"{', '.join('.' + e for e in EXTENSIONES_DOCUMENTO)}."
        )
    reemplazos = {str(k): str(v) for k, v in (reemplazos or {}).items()}
    spec = _como_estructura(contenido)
    rechazar_relleno(spec)
    cabecera, bloques = _bloques_de(spec)
    _validar_documento(cabecera, bloques, con_reemplazos=bool(reemplazos))
    destino = carpeta_de_salida(carpeta)
    ruta = ruta_libre(destino, nombre)
    ruta_plantilla = _plantilla_valida(plantilla, ("docx",)) if plantilla else None
    avisos_estilo: List[str] = []
    estilo_doc = _estilo_de(estilo, avisos_estilo)

    if extension == "docx":
        resultado = _docx(cabecera, bloques, ruta, ruta_plantilla, reemplazos, estilo_doc)
        resultado.avisos = avisos_estilo + resultado.avisos
        return resultado

    if extension == "pdf":
        # Primero el Word —con la plantilla si la hay— y de ahí el PDF con el propio Word:
        # es la única forma de que el PDF se vea exactamente igual. Sin Word, se maqueta.
        with tempfile.TemporaryDirectory() as temporal:
            intermedio = Path(temporal) / (ruta.stem + ".docx")
            resultado = _docx(cabecera, bloques, intermedio, ruta_plantilla, reemplazos, estilo_doc)
            resultado.avisos = avisos_estilo + resultado.avisos
            pdf_temporal = Path(temporal) / (ruta.stem + ".pdf")
            if _con_office("Word", str(intermedio), str(pdf_temporal)):
                os.replace(str(pdf_temporal), str(ruta))
                return Resultado(ruta=str(ruta), formato="pdf",
                                 detalle=resultado.detalle + ", convertido con Word",
                                 avisos=resultado.avisos)
        avisos = list(resultado.avisos)
        if ruta_plantilla:
            avisos.append("Sin Word no puedo aplicar la plantilla al PDF: se maquetó con el "
                          "estilo básico.")
        if estilo_doc:
            avisos.append("Sin Word, del estilo pedido solo van la fuente y los colores.")
        _pdf_sin_office(_html_de(cabecera, bloques, estilo_doc), ruta)
        return Resultado(ruta=str(ruta), formato="pdf",
                         detalle=f"{len(bloques)} bloques, maquetado sin Office", avisos=avisos)

    avisos = list(avisos_estilo)
    if extension == "html":
        ruta.write_text(_html_de(cabecera, bloques, estilo_doc), encoding="utf-8")
    elif extension == "md":
        ruta.write_text(_markdown_de(cabecera, bloques), encoding="utf-8")
        if estilo_doc:
            avisos.append("Un .md no lleva fuente ni colores: el estilo no aplica.")
    else:
        ruta.write_text(_texto_plano_de(cabecera, bloques), encoding="utf-8")
        if estilo_doc:
            avisos.append("Un .txt no lleva fuente ni colores: el estilo no aplica.")
    return Resultado(ruta=str(ruta), formato=extension, detalle=f"{len(bloques)} bloques",
                     avisos=avisos)


# ==================================================================== hojas de cálculo

def _hojas_de(spec: Any) -> List[Dict[str, Any]]:
    """Return la lista de hojas de una especificación, venga como venga."""
    if isinstance(spec, dict) and "hojas" in spec:
        hojas = spec["hojas"]
    elif isinstance(spec, dict) and ("filas" in spec or "encabezados" in spec):
        hojas = [spec]
    elif isinstance(spec, list) and spec and isinstance(spec[0], (list, tuple)):
        hojas = [{"encabezados": list(spec[0]), "filas": [list(f) for f in spec[1:]]}]
    elif isinstance(spec, list) and spec and isinstance(spec[0], dict) and "filas" in spec[0]:
        hojas = spec
    elif isinstance(spec, list) and spec and isinstance(spec[0], dict):
        # Una lista de registros: las claves son los encabezados.
        claves = list(spec[0].keys())
        hojas = [{"encabezados": claves, "filas": [[r.get(k) for k in claves] for r in spec]}]
    else:
        raise DocumentoRechazado("El contenido de la hoja tiene que traer 'hojas' o, al menos, "
                                 "'encabezados' y 'filas'.")
    salida = []
    for i, h in enumerate(hojas or [], 1):
        if not isinstance(h, dict):
            raise DocumentoRechazado("Cada hoja tiene que ser un objeto con 'encabezados' y 'filas'.")
        filas = [list(f) if isinstance(f, (list, tuple)) else [f] for f in (h.get("filas") or [])]
        if not filas and not h.get("encabezados"):
            raise DocumentoRechazado(f"La hoja {i} está vacía: no trae filas.")
        salida.append({**h, "nombre": str(h.get("nombre") or f"Hoja{i}")[:31], "filas": filas,
                       "encabezados": [str(c) for c in (h.get("encabezados") or [])]})
    if not salida:
        raise DocumentoRechazado("No hay ninguna hoja que escribir.")
    return salida


@dataclass(frozen=True)
class Convencion:
    """Cómo se escriben los números en una región: qué separa los miles y qué los decimales."""

    miles: str
    decimal: str
    separador_csv: str

    def formatear(self, valor: float, decimales: int = 0) -> str:
        """Return el número escrito como se escribe en la región: 1.250.000,50 o 1,250,000.50."""
        texto = f"{valor:,.{decimales}f}"                 # a la inglesa, y luego se traduce
        return texto.replace(",", "\x00").replace(".", self.decimal).replace("\x00", self.miles)


#: Convenciones conocidas. Las de habla hispana usan punto de miles y coma decimal, salvo
#: México y Centroamérica, que escriben como en Estados Unidos. Una región que no esté acá
#: cae en la del idioma: `es-*` a la española, todo lo demás a la inglesa.
_CONVENCIONES = {
    "es": Convencion(miles=".", decimal=",", separador_csv=";"),
    "es-MX": Convencion(miles=",", decimal=".", separador_csv=","),
    "es-GT": Convencion(miles=",", decimal=".", separador_csv=","),
    "es-SV": Convencion(miles=",", decimal=".", separador_csv=","),
    "es-HN": Convencion(miles=",", decimal=".", separador_csv=","),
    "es-NI": Convencion(miles=",", decimal=".", separador_csv=","),
    "es-PA": Convencion(miles=",", decimal=".", separador_csv=","),
    "es-DO": Convencion(miles=",", decimal=".", separador_csv=","),
    "es-PR": Convencion(miles=",", decimal=".", separador_csv=","),
    "en": Convencion(miles=",", decimal=".", separador_csv=","),
    "pt": Convencion(miles=".", decimal=",", separador_csv=";"),
    "fr": Convencion(miles=" ", decimal=",", separador_csv=";"),
    "de": Convencion(miles=".", decimal=",", separador_csv=";"),
    "it": Convencion(miles=".", decimal=",", separador_csv=";"),
}


def convencion() -> Convencion:
    """Return la convención numérica de la región configurada.

    Se consulta cada vez y no se cachea: cambiar la región en la configuración tiene que
    valer para el siguiente documento, no para el siguiente arranque.
    """
    region = ""
    try:
        from config_manager import get_region

        region = get_region()
    except Exception as e:
        logger.debug(f"no se pudo leer la región, se usa la del proyecto: {e}")
    region = (region or "es-CO").replace("_", "-")
    idioma = region.split("-")[0].lower()
    return _CONVENCIONES.get(region) or _CONVENCIONES.get(idioma) or _CONVENCIONES["en"]


def _sin_tildes(texto: str) -> str:
    """Return el texto sin tildes ni diéresis: «MÁX» → «MAX», «AÑO» → «ANO»."""
    import unicodedata

    return "".join(c for c in unicodedata.normalize("NFD", texto)
                   if unicodedata.category(c) != "Mn")


#: Funciones de Excel en español → el nombre con el que se guardan en el archivo, que es
#: siempre el inglés. El modelo escribe «=SUMA(B2:B4)» o «=SI(B2>5;"alto";"bajo")» con la
#: misma naturalidad con la que habla, y Excel, al abrir el archivo, no traduce: «SUMA»
#: sale como #¿NOMBRE? y el «;» o la coma decimal directamente impiden abrirlo ("Error en
#: el método Open"). Pasó en la validación de REQ-040 y el agente había dicho «Listo».
#: Las claves van sin tildes y en mayúsculas (`_sin_tildes`).
_FUNCIONES_ES = {
    "SUMA": "SUM", "PROMEDIO": "AVERAGE", "MAX": "MAX", "MIN": "MIN", "SI": "IF",
    "SI.ERROR": "IFERROR", "SI.CONJUNTO": "IFS", "Y": "AND", "O": "OR", "NO": "NOT",
    "CONTAR": "COUNT", "CONTARA": "COUNTA", "CONTAR.SI": "COUNTIF",
    "CONTAR.SI.CONJUNTO": "COUNTIFS", "CONTAR.BLANCO": "COUNTBLANK", "SUMAR.SI": "SUMIF",
    "SUMAR.SI.CONJUNTO": "SUMIFS", "PROMEDIO.SI": "AVERAGEIF",
    "PROMEDIO.SI.CONJUNTO": "AVERAGEIFS", "REDONDEAR": "ROUND", "REDONDEAR.MAS": "ROUNDUP",
    "REDONDEAR.MENOS": "ROUNDDOWN", "TRUNCAR": "TRUNC", "ENTERO": "INT", "RESIDUO": "MOD",
    "COCIENTE": "QUOTIENT", "RAIZ": "SQRT", "POTENCIA": "POWER", "PRODUCTO": "PRODUCT",
    "SUMAPRODUCTO": "SUMPRODUCT", "SUBTOTALES": "SUBTOTAL", "MEDIANA": "MEDIAN",
    "MODA": "MODE", "DESVEST": "STDEV", "K.ESIMO.MAYOR": "LARGE", "K.ESIMO.MENOR": "SMALL",
    "BUSCARV": "VLOOKUP", "BUSCARH": "HLOOKUP", "BUSCARX": "XLOOKUP", "BUSCAR": "LOOKUP",
    "COINCIDIR": "MATCH", "INDICE": "INDEX", "DESREF": "OFFSET", "INDIRECTO": "INDIRECT",
    "ELEGIR": "CHOOSE", "FILA": "ROW", "FILAS": "ROWS", "COLUMNA": "COLUMN",
    "COLUMNAS": "COLUMNS", "CONCATENAR": "CONCATENATE", "UNIRCADENAS": "TEXTJOIN",
    "TEXTO": "TEXT", "VALOR": "VALUE", "MONEDA": "DOLLAR", "LARGO": "LEN",
    "IZQUIERDA": "LEFT", "DERECHA": "RIGHT", "EXTRAE": "MID", "ESPACIOS": "TRIM",
    "MAYUSC": "UPPER", "MINUSC": "LOWER", "NOMPROPIO": "PROPER", "SUSTITUIR": "SUBSTITUTE",
    "REEMPLAZAR": "REPLACE", "HALLAR": "SEARCH", "ENCONTRAR": "FIND", "REPETIR": "REPT",
    "HOY": "TODAY", "AHORA": "NOW", "FECHA": "DATE", "ANO": "YEAR", "MES": "MONTH",
    "DIA": "DAY", "DIAS": "DAYS", "SIFECHA": "DATEDIF", "DIAS.LAB": "NETWORKDAYS",
    "DIA.LAB": "WORKDAY", "DIASEM": "WEEKDAY", "FIN.MES": "EOMONTH", "ESBLANCO": "ISBLANK",
    "ESNUMERO": "ISNUMBER", "ESTEXTO": "ISTEXT", "ESERROR": "ISERROR", "ALEATORIO": "RAND",
    "ALEATORIO.ENTRE": "RANDBETWEEN", "PAGO": "PMT", "VNA": "NPV", "TIR": "IRR",
    "TASA": "RATE", "VA": "PV", "VF": "FV", "CAMBIAR": "SWITCH", "ABS": "ABS",
}


def normalizar_formula(formula: str) -> str:
    """Return la fórmula como la guarda Excel: funciones en inglés, «,» entre argumentos
    y «.» decimal. Una fórmula que ya venía así no cambia.

    Se recorre con el tokenizador de openpyxl y no con regex sobre el texto: un «;» dentro
    de `"una cadena"` o de una matriz `{1;2}` no es un separador. La coma entre dos enteros
    es decimal si la fórmula lleva «;» (entonces la coma no puede ser separador) o si está
    fuera de toda función (`=D5*0,19`); dentro de una función y sin «;», es separador:
    `=ROUND(B5*1.19,2)` viene en inglés y adivinar lo contrario lo rompería.
    """
    from openpyxl.formula.tokenizer import Token, Tokenizer

    try:
        tokens = Tokenizer(formula).items
    except Exception as e:
        logger.debug(f"no se pudo tokenizar la fórmula {formula!r}: {e}")
        return formula
    hay_punto_y_coma = any(t.type == Token.SEP and t.value == ";" for t in tokens)
    salida: List[str] = []
    en_funcion = 0
    en_matriz = 0
    for i, t in enumerate(tokens):
        valor = t.value
        if t.type == Token.FUNC and t.subtype == Token.OPEN:
            nombre = valor[:-1]
            valor = _FUNCIONES_ES.get(_sin_tildes(nombre).upper(), nombre) + "("
            en_funcion += 1
        elif t.type == Token.FUNC and t.subtype == Token.CLOSE:
            en_funcion -= 1
        elif t.type == Token.ARRAY:
            en_matriz += 1 if t.subtype == Token.OPEN else -1
        elif t.type == Token.SEP and valor == ";" and not en_matriz:
            valor = ","
        elif valor == "," and t.type in (Token.SEP, Token.OP_IN):
            anterior = tokens[i - 1] if i else None
            siguiente = tokens[i + 1] if i + 1 < len(tokens) else None
            entre_enteros = (
                anterior is not None and siguiente is not None
                and anterior.type == Token.OPERAND and anterior.subtype == Token.NUMBER
                and siguiente.type == Token.OPERAND and siguiente.subtype == Token.NUMBER
                and "." not in anterior.value and "." not in siguiente.value
            )
            if entre_enteros and (hay_punto_y_coma or not en_funcion):
                valor = "."
        elif t.type == Token.OPERAND and t.subtype == Token.RANGE:
            clave = _sin_tildes(valor).upper()
            if clave == "VERDADERO":
                valor = "TRUE"
            elif clave == "FALSO":
                valor = "FALSE"
        salida.append(valor)
    return "=" + "".join(salida)


def _valor_de_celda(valor: Any) -> Any:
    """Return el valor listo para Excel: números como números, fórmulas como fórmulas."""
    if valor is None or isinstance(valor, (int, float, bool)):
        return valor
    texto = str(valor).strip()
    if texto.startswith("="):
        return normalizar_formula(texto)        # fórmula, la calcula Excel
    # "1.250.000" y "12,5" son números escritos a mano: se guardan como número para que
    # las fórmulas y los formatos funcionen, no como texto que parece número.
    #
    # Qué es de miles y qué es decimal lo decide la región configurada, no el código. La
    # primera versión tenía fija la convención inglesa y una hoja de costos quedó con
    # ochocientos setenta pesos donde iban ochocientos setenta mil ("870.000"); fijar la
    # española habría hecho lo mismo, al revés, con un usuario en México o Estados Unidos.
    conv = convencion()
    miles, decimal = re.escape(conv.miles), re.escape(conv.decimal)
    patron = rf"-?\d{{1,3}}({miles}\d{{3}})+({decimal}\d+)?|-?\d+({decimal}\d+)?"
    if re.fullmatch(patron, texto):
        limpio = texto.replace(conv.miles, "").replace(conv.decimal, ".")
        try:
            return float(limpio) if "." in limpio else int(limpio)
        except ValueError:
            return texto
    # Un número escrito con la convención contraria y sin ambigüedad —"12.5" donde el punto
    # es de miles no puede ser de miles porque no le siguen tres dígitos— se acepta igual.
    otro = "." if conv.decimal == "," else ","
    if re.fullmatch(rf"-?\d+{re.escape(otro)}\d{{1,2}}", texto):
        try:
            return float(texto.replace(otro, "."))
        except ValueError:
            return texto
    return texto


def _columna(indice: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(indice)


def _celda_a_indices(celda: Any) -> Tuple[int, int]:
    """Return `(columna, fila)` de «C3» → (3, 3). «5» → (1, 5); «C» → (3, 1); raro → (1, 1)."""
    from openpyxl.utils import column_index_from_string

    m = re.fullmatch(r"\s*\$?([A-Za-z]{0,3})\$?(\d{0,7})\s*", str(celda or ""))
    if not m:
        return 1, 1
    letras, numero = m.group(1), m.group(2)
    columna = column_index_from_string(letras.upper()) if letras else 1
    return columna, (int(numero) if numero else 1) or 1


def _textos_de_fila(ws, fila: int) -> List[str]:
    """Return los textos de una fila que ya existe, sin crear celdas nuevas."""
    if fila < 1 or fila > (ws.max_row or 0):
        return []
    return [_sin_tildes(str(v)).strip().lower()
            for v in next(ws.iter_rows(min_row=fila, max_row=fila, values_only=True))
            if isinstance(v, str) and v.strip()]


def _fila_de_encabezado_existente(ws, encabezados: List[str], candidatas) -> Optional[int]:
    """Return la primera fila candidata que ya trae (al menos la mitad de) esos encabezados.

    Es lo que evita repetir el encabezado de una plantilla debajo del suyo: si el modelo
    manda «Producto, Cantidad, Precio» y la fila 4 de la plantilla ya dice eso, los datos
    van en la 5 y el encabezado se queda como estaba, con su formato.
    """
    buscados = {_sin_tildes(e).strip().lower() for e in encabezados if str(e).strip()}
    if not buscados:
        return None
    for fila in candidatas:
        textos = set(_textos_de_fila(ws, fila))
        if textos and len(buscados & textos) >= max(1, (len(buscados) + 1) // 2):
            return fila
    return None


def _celda_de_encabezado_con_estilo(ws, hasta_fila: int):
    """Return una celda de encabezado de la plantilla (con relleno sólido) por encima de
    `hasta_fila`, o None. Su formato se copia al encabezado nuevo en vez del azul fijo."""
    for fila in range(min(hasta_fila - 1, ws.max_row or 0), 0, -1):
        for celda in next(ws.iter_rows(min_row=fila, max_row=fila)):
            if celda.value is None:
                continue
            if celda.fill is not None and celda.fill.fill_type == "solid":
                return celda
            break
    return None


def _escribir_hoja(ws, hoja: Dict[str, Any], desde_fila: int = 1, desde_columna: int = 1,
                   escribir_encabezados: bool = True, estilo: Optional[Dict[str, Any]] = None,
                   celda_modelo_encabezado=None) -> Dict[str, Any]:
    """Escribe título, encabezados, filas, totales y formato en una hoja. Return qué se hizo.

    `desde_fila`/`desde_columna` es la esquina de la tabla; `escribir_encabezados=False`
    cuando la plantilla ya los tiene; `estilo` es lo que pidió el usuario (fuente, tamaño,
    colores del encabezado, bordes); `celda_modelo_encabezado` es una celda de la plantilla
    cuyo formato se copia al encabezado nuevo.
    """
    from copy import copy

    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    estilo = estilo or {}
    encabezados = hoja["encabezados"]
    filas = hoja["filas"]
    columnas = max([len(encabezados)] + [len(f) for f in filas]) if (encabezados or filas) else 0
    col0 = max(1, int(desde_columna or 1))
    fila_actual = max(1, int(desde_fila or 1))
    con_bordes = estilo.get("bordes") is not False
    borde = Border(*(Side(style="thin", color="BBBBBB"),) * 4) if con_bordes else None

    def _fuente(**extra) -> Font:
        base = {k: v for k, v in (("name", estilo.get("fuente")),
                                  ("size", estilo.get("tamano"))) if v}
        base.update(extra)
        return Font(**base)

    con_estilo_de_texto = bool(estilo.get("fuente") or estilo.get("tamano"))

    titulo = str(hoja.get("titulo") or "").strip()
    if titulo:
        celda = ws.cell(row=fila_actual, column=col0, value=titulo)
        color_titulo = estilo.get("color_titulos") or estilo.get("color_encabezado") or "000000"
        celda.font = _fuente(bold=True, size=float(estilo.get("tamano") or 11) + 5,
                             color=color_titulo)
        celda.alignment = Alignment(horizontal="left", vertical="center")
        if columnas > 1:
            ws.merge_cells(start_row=fila_actual, start_column=col0,
                           end_row=fila_actual, end_column=col0 + columnas - 1)
        fila_actual += 2                           # el título y una fila en blanco

    fila_encabezado: Optional[int] = None
    if encabezados and escribir_encabezados:
        fila_encabezado = fila_actual
        for i, texto in enumerate(encabezados):
            celda = ws.cell(row=fila_actual, column=col0 + i, value=texto)
            if celda_modelo_encabezado is not None and not estilo.get("color_encabezado"):
                celda.font = copy(celda_modelo_encabezado.font)
                celda.fill = copy(celda_modelo_encabezado.fill)
                celda.alignment = copy(celda_modelo_encabezado.alignment)
                celda.border = copy(celda_modelo_encabezado.border)
            else:
                celda.font = _fuente(bold=True,
                                     color=estilo.get("color_texto_encabezado") or "FFFFFF")
                celda.fill = PatternFill("solid",
                                        fgColor=estilo.get("color_encabezado") or "305496")
                celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                if borde:
                    celda.border = borde
        fila_actual += 1

    primera_dato = fila_actual
    for fila in filas:
        for i, valor in enumerate(fila[:columnas], 0):
            celda = ws.cell(row=fila_actual, column=col0 + i, value=_valor_de_celda(valor))
            if borde:
                celda.border = borde
            if con_estilo_de_texto:
                celda.font = _fuente()
        fila_actual += 1
    ultima_dato = fila_actual - 1

    # Totales: una fila con SUMA en cada columna que sea numérica, y la etiqueta en la primera.
    if hoja.get("totales") and filas and ultima_dato >= primera_dato:
        numericas = []
        for c in range(col0, col0 + columnas):
            valores = [ws.cell(row=r, column=c).value for r in range(primera_dato, ultima_dato + 1)]
            if valores and all(isinstance(v, (int, float)) or
                               (isinstance(v, str) and v.startswith("=")) for v in valores if v is not None) \
                    and any(isinstance(v, (int, float)) or (isinstance(v, str) and v.startswith("="))
                            for v in valores):
                numericas.append(c)
        if numericas:
            etiqueta = ws.cell(row=fila_actual, column=col0, value="Total")
            etiqueta.font = _fuente(bold=True)
            for c in numericas:
                if c == col0:
                    continue
                letra = get_column_letter(c)
                celda = ws.cell(row=fila_actual, column=c,
                                value=f"=SUM({letra}{primera_dato}:{letra}{ultima_dato})")
                celda.font = _fuente(bold=True)
                if borde:
                    celda.border = borde
            fila_actual += 1

    # Formatos de número por columna: {"B": "#,##0", "D": "0.0%"}. Las letras son las de
    # la hoja, no las de la tabla: es lo que el usuario ve en Excel.
    for letra, formato in (hoja.get("formatos") or {}).items():
        try:
            for r in range(primera_dato, fila_actual):
                ws[f"{str(letra).upper()}{r}"].number_format = str(formato)
        except Exception as e:
            logger.debug(f"formato «{formato}» no aplicable a la columna {letra}: {e}")

    # Anchos: a la medida del contenido, con tope.
    for c in range(col0, col0 + columnas):
        letra = get_column_letter(c)
        desde_ancho = primera_dato - (1 if fila_encabezado else 0)
        largo = max((len(str(ws.cell(row=r, column=c).value or ""))
                     for r in range(desde_ancho, fila_actual)), default=8)
        ws.column_dimensions[letra].width = min(max(10, largo + 2), _ANCHO_COLUMNA_MAXIMO)

    if fila_encabezado is not None and col0 == 1:
        ws.freeze_panes = f"A{fila_encabezado + 1}"
        if ultima_dato >= primera_dato:
            ws.auto_filter.ref = f"A{fila_encabezado}:{get_column_letter(columnas)}{ultima_dato}"

    grafico = hoja.get("grafico")
    if grafico and filas and ultima_dato >= primera_dato:
        _agregar_grafico(ws, grafico, encabezados, primera_dato, ultima_dato, fila_actual + 2, col0)

    return {"filas": len(filas), "columnas": columnas, "totales": bool(hoja.get("totales")),
            "fila_encabezado": fila_encabezado, "primera_dato": primera_dato}


def _agregar_grafico(ws, grafico: Dict[str, Any], encabezados: List[str],
                     primera: int, ultima: int, fila_ancla: int, col0: int = 1) -> None:
    """Pone un gráfico de barras, líneas o torta con las columnas pedidas.

    Las columnas se buscan por su nombre en `encabezados` aunque el encabezado no se haya
    escrito (plantilla que ya lo traía): antes, en ese caso, la búsqueda fallaba y se
    graficaba la columna 2 con el nombre «Series1». El nombre de la serie se pone
    explícito (`SeriesLabel`) y no se lee de la celda de arriba, que puede no ser el
    encabezado.
    """
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.chart.series import SeriesLabel
    from openpyxl.utils import column_index_from_string

    tipo = str(grafico.get("tipo") or "barras").lower()
    clase = {"barras": BarChart, "columnas": BarChart, "lineas": LineChart, "línea": LineChart,
             "linea": LineChart, "torta": PieChart, "pastel": PieChart, "pie": PieChart}.get(tipo, BarChart)
    chart = clase()
    chart.title = str(grafico.get("titulo") or "")
    if hasattr(chart, "y_axis") and grafico.get("eje_y"):
        chart.y_axis.title = str(grafico["eje_y"])

    columnas = grafico.get("columnas") or ([encabezados[1]] if len(encabezados) > 1 else [])
    if isinstance(columnas, str):
        columnas = [columnas]
    categorias = grafico.get("categorias") or (encabezados[0] if encabezados else "A")
    normales = [_sin_tildes(e).strip().lower() for e in encabezados]

    def _indice(referencia: Any) -> int:
        texto = str(referencia).strip()
        buscado = _sin_tildes(texto).lower()
        if buscado in normales:
            return col0 + normales.index(buscado)
        try:
            return column_index_from_string(texto.upper())
        except Exception:
            return col0 + 1

    for col in columnas:
        c = _indice(col)
        chart.add_data(Reference(ws, min_col=c, min_row=primera, max_row=ultima),
                       titles_from_data=False)
        nombre = encabezados[c - col0] if 0 <= c - col0 < len(encabezados) else str(col)
        chart.series[-1].tx = SeriesLabel(v=nombre)
    chart.set_categories(Reference(ws, min_col=_indice(categorias), min_row=primera,
                                   max_row=ultima))
    chart.width, chart.height = 18, 9
    if chart.legend is not None:
        chart.legend.position = "b"          # abajo: a la derecha se corta al imprimir
    from openpyxl.utils import get_column_letter

    ws.add_chart(chart, f"{get_column_letter(col0)}{fila_ancla}")


_REF_RE = re.compile(r"\$?([A-Z]{1,3})\$?(\d{1,7})")
_RANGO_RE = re.compile(r"\$?[A-Z]{1,3}\$?\d{1,7}:\$?[A-Z]{1,3}\$?\d{1,7}")
_TOKEN_RE = re.compile(
    r"\s*(?:(?P<num>\d+(?:\.\d+)?)|(?P<fn>SUM|SUMA|AVERAGE|PROMEDIO|MIN|MAX)\(|"
    r"(?P<rango>\$?[A-Z]{1,3}\$?\d+:\$?[A-Z]{1,3}\$?\d+)|(?P<ref>\$?[A-Z]{1,3}\$?\d+)|"
    r"(?P<op>[-+*/()]))"
)


def _calcular_formula(formula: str, celdas: Dict[str, Any], profundidad: int = 0) -> Optional[float]:
    """Return el valor de una fórmula SENCILLA, o None si no se puede calcular.

    Sin Excel nadie calcula, y un PDF con `=SUM(B2:B4)` escrito en la fila de totales no es
    un informe: es una hoja a medio hacer. Esto evalúa lo que este módulo produce y lo que
    una hoja corriente trae —sumas, restas, productos, divisiones, paréntesis, SUMA,
    PROMEDIO, MÍN y MÁX sobre rangos— con un evaluador propio: nunca `eval()`, y cualquier
    cosa fuera de esa gramática devuelve None, que se muestra como la fórmula escrita.
    """
    if profundidad > 20:
        return None
    texto = formula.lstrip("=").strip().upper()
    pos = 0
    salida: List[Any] = []                              # notación polaca inversa
    operadores: List[str] = []
    precedencia = {"+": 1, "-": 1, "*": 2, "/": 2}

    def _valor_de(ref: str) -> Optional[float]:
        v = celdas.get(ref.replace("$", ""))
        if isinstance(v, bool):
            return None
        if isinstance(v, (int, float)):
            return float(v)
        if isinstance(v, str) and v.startswith("="):
            return _calcular_formula(v, celdas, profundidad + 1)
        return 0.0 if v in (None, "") else None

    def _rango(texto_rango: str) -> Optional[List[float]]:
        a, b = texto_rango.replace("$", "").split(":")
        ca, fa = _REF_RE.match(a).groups()
        cb, fb = _REF_RE.match(b).groups()
        from openpyxl.utils import column_index_from_string, get_column_letter

        valores = []
        for c in range(column_index_from_string(ca), column_index_from_string(cb) + 1):
            for f in range(int(fa), int(fb) + 1):
                v = _valor_de(f"{get_column_letter(c)}{f}")
                if v is None:
                    return None
                valores.append(v)
        return valores

    esperando_operando = True
    while pos < len(texto):
        m = _TOKEN_RE.match(texto, pos)
        if not m:
            return None
        pos = m.end()
        if m.group("num"):
            salida.append(float(m.group("num")))
            esperando_operando = False
        elif m.group("fn"):
            cierre = texto.find(")", pos)
            if cierre < 0:
                return None
            interior = texto[pos:cierre]
            pos = cierre + 1
            valores = _rango(interior) if _RANGO_RE.fullmatch(interior) else None
            if valores is None:
                return None
            fn = m.group("fn")
            if fn in ("SUM", "SUMA"):
                salida.append(float(sum(valores)))
            elif fn in ("AVERAGE", "PROMEDIO"):
                salida.append(sum(valores) / len(valores) if valores else 0.0)
            elif fn == "MIN":
                salida.append(float(min(valores)) if valores else 0.0)
            else:
                salida.append(float(max(valores)) if valores else 0.0)
            esperando_operando = False
        elif m.group("rango"):
            return None                                 # un rango suelto no es un valor
        elif m.group("ref"):
            v = _valor_de(m.group("ref"))
            if v is None:
                return None
            salida.append(v)
            esperando_operando = False
        else:
            op = m.group("op")
            if op == "(":
                operadores.append(op)
            elif op == ")":
                while operadores and operadores[-1] != "(":
                    salida.append(operadores.pop())
                if not operadores:
                    return None
                operadores.pop()
            else:
                if op == "-" and esperando_operando:
                    salida.append(0.0)                  # menos unario: 0 - x
                while operadores and operadores[-1] != "(" and \
                        precedencia.get(operadores[-1], 0) >= precedencia[op]:
                    salida.append(operadores.pop())
                operadores.append(op)
                esperando_operando = True
    while operadores:
        if operadores[-1] == "(":
            return None
        salida.append(operadores.pop())

    pila: List[float] = []
    for token in salida:
        if isinstance(token, float):
            pila.append(token)
            continue
        if len(pila) < 2:
            return None
        b, a = pila.pop(), pila.pop()
        if token == "+":
            pila.append(a + b)
        elif token == "-":
            pila.append(a - b)
        elif token == "*":
            pila.append(a * b)
        elif token == "/":
            if b == 0:
                return None
            pila.append(a / b)
    return pila[0] if len(pila) == 1 else None


def _html_de_hoja(ruta_xlsx: Path) -> Tuple[str, List[str]]:
    """Return `(html, avisos)` con cada hoja del libro como una tabla. Para el PDF sin Excel."""
    import html as _html

    from openpyxl import load_workbook

    e = _html.escape
    # En modo de solo lectura, porque es el único en que `close()` suelta el archivo de
    # verdad. En modo normal openpyxl deja el archivo abierto hasta que el objeto se
    # recolecta, y en Windows eso bloquea el borrado de la carpeta temporal: pasó.
    wb = load_workbook(str(ruta_xlsx), read_only=True)
    try:
        return _html_de_libro(wb, e)
    finally:
        wb.close()


def _html_de_libro(wb, e) -> Tuple[str, List[str]]:
    avisos: List[str] = []
    partes = ["<!doctype html><html><head><meta charset='utf-8'><style>"
              "body{font-family:Calibri,Arial,sans-serif;font-size:9pt;margin:1.2cm}"
              "h2{font-size:13pt;margin:0 0 6pt 0}table{border-collapse:collapse;margin-bottom:14pt}"
              "td,th{border:1px solid #999;padding:3pt 5pt;text-align:left;white-space:nowrap}"
              "th{background:#305496;color:#fff}td.n{text-align:right}</style></head><body>"]
    sin_calcular = False
    from openpyxl.utils import get_column_letter

    for ws in wb.worksheets:
        filas_leidas = [f for f in ws.iter_rows(values_only=True)
                        if any(c is not None for c in f)]
        if not filas_leidas:
            continue
        # Un mapa celda -> valor para calcular las fórmulas sencillas sin Excel.
        celdas_por_ref: Dict[str, Any] = {}
        for numero_fila, fila in enumerate(ws.iter_rows(values_only=True), 1):
            for numero_col, valor in enumerate(fila, 1):
                if valor is not None:
                    celdas_por_ref[f"{get_column_letter(numero_col)}{numero_fila}"] = valor
        partes.append(f"<h2>{e(ws.title)}</h2><table>")
        conv = convencion()
        for i, fila in enumerate(filas_leidas, 1):
            etiqueta = "th" if i == 1 else "td"
            celdas = []
            for valor in fila:
                if isinstance(valor, str) and valor.startswith("="):
                    calculado = _calcular_formula(valor, celdas_por_ref)
                    if calculado is None:
                        sin_calcular = True
                        celdas.append(f"<{etiqueta} class='n'><i>{e(valor)}</i></{etiqueta}>")
                    else:
                        decimales = 0 if float(calculado).is_integer() else 2
                        celdas.append(f"<{etiqueta} class='n'>{e(conv.formatear(calculado, decimales))}</{etiqueta}>")
                elif isinstance(valor, (int, float)) and not isinstance(valor, bool):
                    decimales = 0 if float(valor).is_integer() else 2
                    celdas.append(f"<{etiqueta} class='n'>{e(conv.formatear(valor, decimales))}</{etiqueta}>")
                else:
                    celdas.append(f"<{etiqueta}>{e('' if valor is None else str(valor))}</{etiqueta}>")
            partes.append(f"<tr>{''.join(celdas)}</tr>")
        partes.append("</table>")
    partes.append("</body></html>")
    if sin_calcular:
        avisos.append("Sin Excel, alguna fórmula quedó escrita y no calculada en el PDF; "
                      "el .xlsx sí las calcula al abrirlo.")
    return "".join(partes), avisos


def _pdf_apaisado_sin_office(html: str, ruta: Path) -> None:
    """Maqueta HTML en un PDF apaisado (A4 horizontal): para hojas y diapositivas."""
    import pymupdf

    historia = pymupdf.Story(html=html)
    escritor = pymupdf.DocumentWriter(str(ruta))
    rect_pagina = pymupdf.paper_rect("a4-l")
    area = rect_pagina + (36, 36, -36, -36)
    hay_mas = True
    while hay_mas:
        dispositivo = escritor.begin_page(rect_pagina)
        hay_mas, _ = historia.place(area)
        historia.draw(dispositivo)
        escritor.end_page()
    escritor.close()


def crear_hoja(nombre: str, contenido: Any, plantilla: Optional[str] = None,
               carpeta: Optional[str] = None, estilo: Any = None) -> Resultado:
    """Produce una hoja de cálculo: .xlsx o .csv. Levanta `DocumentoRechazado`.

    `contenido` trae `hojas`, y cada hoja `nombre`, `encabezados`, `filas`, y opcionalmente
    `totales` (fila de sumas), `formatos` (`{"B": "#,##0"}`), `grafico` (`tipo`, `titulo`,
    `columnas`, `categorias`) y `desde` (celda donde empezar, para una plantilla).
    Las celdas que empiezan por `=` son fórmulas y las calcula Excel; pueden venir en
    español (`=SUMA(...)`, `=SI(a;b;c)`, `0,19`) y se guardan como Excel las entiende.
    `estilo` es lo que pidió el usuario: `fuente`, `tamano`, `color_encabezado`,
    `color_texto_encabezado`, `bordes`; y cada hoja acepta `titulo`.
    """
    extension = _extension(nombre)
    if extension not in EXTENSIONES_HOJA + ("pdf",):
        raise DocumentoRechazado(f"«{nombre}» no es una hoja de cálculo: hace falta .xlsx, .csv o .pdf.")
    spec = _como_estructura(contenido)
    rechazar_relleno(spec)
    hojas = _hojas_de(spec)
    avisos_estilo: List[str] = []
    estilo_hoja = _estilo_de(estilo, avisos_estilo)
    destino = carpeta_de_salida(carpeta)
    ruta = ruta_libre(destino, nombre)

    if extension == "pdf":
        # Primero el Excel de verdad, y de ahí el PDF con Excel. Sin Excel, cada hoja se
        # maqueta como tabla apaisada: sin gráficos y con las fórmulas escritas, y se dice.
        with tempfile.TemporaryDirectory() as temporal:
            intermedio = Path(temporal) / (ruta.stem + ".xlsx")
            resultado = _xlsx(hojas, intermedio,
                              _plantilla_valida(plantilla, ("xlsx",)) if plantilla else None,
                              estilo_hoja)
            resultado.avisos = avisos_estilo + resultado.avisos
            pdf_temporal = Path(temporal) / (ruta.stem + ".pdf")
            if _con_office("Excel", resultado.ruta, str(pdf_temporal)):
                os.replace(str(pdf_temporal), str(ruta))
                return Resultado(ruta=str(ruta), formato="pdf",
                                 detalle=resultado.detalle + "; convertido con Excel",
                                 avisos=resultado.avisos)
            html, avisos = _html_de_hoja(Path(resultado.ruta))
            _pdf_apaisado_sin_office(html, ruta)
            if any(h.get("grafico") for h in hojas):
                avisos.append("Sin Excel, el gráfico no entra en el PDF.")
            return Resultado(ruta=str(ruta), formato="pdf",
                             detalle=resultado.detalle + "; maquetado sin Office",
                             avisos=resultado.avisos + avisos)

    if extension == "csv":
        hoja = hojas[0]
        with open(ruta, "w", encoding="utf-8-sig", newline="") as f:
            # El separador que Excel abre bien en la región: ";" en español, "," en inglés.
            escritor = csv.writer(f, delimiter=convencion().separador_csv)
            if hoja["encabezados"]:
                escritor.writerow(hoja["encabezados"])
            escritor.writerows(hoja["filas"])
        return Resultado(ruta=str(ruta), formato="csv",
                         detalle=f"{len(hoja['filas'])} filas" +
                                 (", solo la primera hoja" if len(hojas) > 1 else ""))

    resultado = _xlsx(hojas, ruta, _plantilla_valida(plantilla, ("xlsx",)) if plantilla else None,
                      estilo_hoja)
    resultado.avisos = avisos_estilo + resultado.avisos
    return resultado


def _xlsx(hojas: List[Dict[str, Any]], ruta: Path, ruta_plantilla: Optional[Path],
          estilo: Optional[Dict[str, Any]] = None) -> Resultado:
    """Escribe el libro en `ruta`, que ya viene decidida. Es lo que comparten .xlsx y .pdf."""
    from openpyxl import Workbook, load_workbook

    if ruta_plantilla:
        wb = load_workbook(str(ruta_plantilla))
    else:
        wb = Workbook()
        wb.remove(wb.active)

    resumen = []
    avisos: List[str] = []
    for hoja in hojas:
        desde_fila, desde_columna = 1, 1
        escribir_encabezados = True
        celda_modelo = None
        if ruta_plantilla and hoja["nombre"] in wb.sheetnames:
            ws = wb[hoja["nombre"]]
            if hoja.get("desde"):
                # `desde` es donde va la primera fila de DATOS. Si la plantilla ya tiene
                # el encabezado ahí o justo encima, no se repite; si está justo ahí, los
                # datos bajan una fila.
                desde_columna, desde_fila = _celda_a_indices(hoja["desde"])
                fila_enc = _fila_de_encabezado_existente(ws, hoja["encabezados"],
                                                         (desde_fila, desde_fila - 1))
                if fila_enc is not None:
                    escribir_encabezados = False
                    if fila_enc == desde_fila:
                        desde_fila += 1
            elif (ws.max_row or 0) > 1 or ws.cell(row=1, column=1).value is not None:
                # La plantilla tiene contenido: se escribe debajo de lo que haya, y si el
                # encabezado ya está en alguna fila, no se repite.
                fila_enc = _fila_de_encabezado_existente(ws, hoja["encabezados"],
                                                         range(ws.max_row, 0, -1))
                desde_fila = ws.max_row + 1
                if fila_enc is not None:
                    escribir_encabezados = False
                    primera_con_texto = next(
                        (c.column for c in next(ws.iter_rows(min_row=fila_enc, max_row=fila_enc))
                         if c.value is not None), 1)
                    desde_columna = primera_con_texto
            if escribir_encabezados:
                celda_modelo = _celda_de_encabezado_con_estilo(ws, desde_fila)
        else:
            ws = wb.create_sheet(hoja["nombre"])
        hecho = _escribir_hoja(ws, hoja, desde_fila, desde_columna, escribir_encabezados,
                               estilo, celda_modelo)
        linea = f"{hoja['nombre']}: {hecho['filas']} filas × {hecho['columnas']} columnas"
        if hecho["primera_dato"] != 2:
            linea += f", datos desde la fila {hecho['primera_dato']}"
            if hecho["fila_encabezado"]:
                linea += f" (encabezado en la {hecho['fila_encabezado']})"
        resumen.append(linea)

    wb.save(str(ruta))
    detalle = "; ".join(resumen)
    if ruta_plantilla:
        detalle += f"; sobre la plantilla {ruta_plantilla.name}"
    return Resultado(ruta=str(ruta), formato="xlsx", detalle=detalle, avisos=avisos)


# ===================================================================== presentaciones

def _diapositivas_de(spec: Any) -> Tuple[Dict[str, Any], List[Dict[str, Any]]]:
    if isinstance(spec, list):
        return {}, [d if isinstance(d, dict) else {"titulo": str(d)} for d in spec]
    if not isinstance(spec, dict):
        raise DocumentoRechazado("El contenido de la presentación tiene que ser un objeto con "
                                 "'diapositivas'.")
    diapositivas = spec.get("diapositivas") or spec.get("slides") or []
    cabecera = {k: spec.get(k) for k in ("titulo", "subtitulo") if spec.get(k)}
    return cabecera, [d if isinstance(d, dict) else {"titulo": str(d)} for d in diapositivas]


def _validar_presentacion(cabecera: Dict[str, Any], diapositivas: List[Dict[str, Any]],
                          con_reemplazos: bool) -> None:
    if con_reemplazos and not diapositivas:
        return
    con_contenido = [d for d in diapositivas
                     if d.get("puntos") or d.get("texto") or d.get("tabla") or d.get("imagen")
                     or d.get("grafico")]
    if len(con_contenido) < 2:
        raise DocumentoRechazado(
            "La presentación no tiene contenido suficiente: hacen falta al menos dos "
            "diapositivas con sus puntos o su texto, no solo títulos. Escribila completa."
        )


def _layout(prs, *nombres: str, con_cuerpo: bool = True):
    """Return el diseño de diapositiva que mejor encaja, por nombre o por sus marcadores."""
    from pptx.enum.shapes import PP_PLACEHOLDER

    for nombre in nombres:
        for layout in prs.slide_layouts:
            if layout.name.lower() == nombre.lower():
                return layout
    # Sin coincidencia por nombre: el primero que tenga título y, si hace falta, cuerpo.
    for layout in prs.slide_layouts:
        tipos = {ph.placeholder_format.type for ph in layout.placeholders}
        tiene_titulo = PP_PLACEHOLDER.TITLE in tipos or PP_PLACEHOLDER.CENTER_TITLE in tipos
        tiene_cuerpo = PP_PLACEHOLDER.BODY in tipos or PP_PLACEHOLDER.OBJECT in tipos
        if tiene_titulo and (tiene_cuerpo or not con_cuerpo):
            return layout
    return prs.slide_layouts[0]


def _reemplazar_en_pptx(prs, reemplazos: Dict[str, str]) -> int:
    cuantos = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for parrafo in shape.text_frame.paragraphs:
                    texto = "".join(r.text for r in parrafo.runs)
                    if any(m in texto for m in reemplazos):
                        for m, v in reemplazos.items():
                            cuantos += texto.count(m)
                            texto = texto.replace(m, str(v))
                        if parrafo.runs:
                            parrafo.runs[0].text = texto
                            for r in parrafo.runs[1:]:
                                r.text = ""
            if shape.has_table:
                for fila in shape.table.rows:
                    for celda in fila.cells:
                        for parrafo in celda.text_frame.paragraphs:
                            texto = "".join(r.text for r in parrafo.runs)
                            if any(m in texto for m in reemplazos):
                                for m, v in reemplazos.items():
                                    cuantos += texto.count(m)
                                    texto = texto.replace(m, str(v))
                                if parrafo.runs:
                                    parrafo.runs[0].text = texto
                                    for r in parrafo.runs[1:]:
                                        r.text = ""
    return cuantos


def _quitar_diapositivas(prs) -> None:
    """Deja la presentación sin diapositivas, conservando el patrón, los diseños y el tema."""
    lista = prs.slides._sldIdLst
    for sld_id in list(lista):
        prs.part.drop_rel(sld_id.rId)
        lista.remove(sld_id)


def _agregar_diapositivas(prs, diapositivas: List[Dict[str, Any]], avisos: List[str]) -> None:
    """Añade las diapositivas al final de la presentación. Lo usan crear y editar."""
    from pptx.util import Cm, Pt

    for d in diapositivas:
        tiene_cuerpo = bool(d.get("puntos") or d.get("texto"))
        layout = _layout(prs, "Title and Content", "Título y objetos", con_cuerpo=tiene_cuerpo) \
            if tiene_cuerpo else _layout(prs, "Title Only", "Solo el título", con_cuerpo=False)
        slide = prs.slides.add_slide(layout)
        if slide.shapes.title is not None:
            slide.shapes.title.text = str(d.get("titulo") or "")

        # Todo lo que no va en un marcador —tablas, imágenes, texto suelto— se coloca
        # DEBAJO del título, donde termine en este diseño. A una altura fija, en el tema
        # Retrospect la tabla quedaba montada sobre el título.
        titulo_forma = slide.shapes.title
        arriba = (titulo_forma.top + titulo_forma.height + Cm(0.6)) if titulo_forma is not None             else Cm(4.5)
        alto_libre = max(prs.slide_height - arriba - Cm(1.5), Cm(4))

        cuerpo = next((ph for ph in slide.placeholders if ph.placeholder_format.idx == 1), None)
        if tiene_cuerpo and cuerpo is not None:
            tf = cuerpo.text_frame
            tf.clear()
            primero = True
            puntos = d.get("puntos") or ([d["texto"]] if d.get("texto") else [])
            for punto in puntos:
                nivel_items = [(0, punto)] if not isinstance(punto, (list, tuple)) \
                    else [(1, sub) for sub in punto]
                for nivel, texto in nivel_items:
                    p = tf.paragraphs[0] if primero else tf.add_paragraph()
                    primero = False
                    p.text = str(texto)
                    p.level = nivel
        elif tiene_cuerpo:
            caja = slide.shapes.add_textbox(Cm(2), arriba, prs.slide_width - Cm(4), alto_libre)
            tf = caja.text_frame
            tf.word_wrap = True
            for i, punto in enumerate(d.get("puntos") or [d.get("texto")]):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = ("• " if d.get("puntos") else "") + str(punto)
                p.font.size = Pt(18)

        tabla = d.get("tabla")
        if isinstance(tabla, dict) and (tabla.get("filas") or tabla.get("encabezados")):
            enc = [str(c) for c in (tabla.get("encabezados") or [])]
            filas = [[("" if c is None else str(c)) for c in f] for f in (tabla.get("filas") or [])]
            cols = max([len(enc)] + [len(f) for f in filas])
            n_filas = len(filas) + (1 if enc else 0)
            alto = min(Cm(0.9) * n_filas, alto_libre)
            forma = slide.shapes.add_table(n_filas, cols, Cm(1.5), arriba,
                                           prs.slide_width - Cm(3), alto)
            r = 0
            if enc:
                for c, texto in enumerate(enc):
                    forma.table.cell(0, c).text = texto
                r = 1
            for fila in filas:
                for c, texto in enumerate(fila[:cols]):
                    forma.table.cell(r, c).text = texto
                r += 1

        grafico = d.get("grafico")
        if isinstance(grafico, dict):
            if not _grafico_nativo_pptx(slide, grafico, Cm(1.5), arriba,
                                        prs.slide_width - Cm(3), alto_libre):
                avisos.append(f"El gráfico de «{d.get('titulo') or ''}» no traía datos y se omitió.")

        imagen = d.get("imagen")
        if imagen:
            if os.path.isfile(str(imagen)):
                try:
                    ancho = prs.slide_width - Cm(4)
                    slide.shapes.add_picture(str(imagen), Cm(2), arriba,
                                             width=ancho if not tiene_cuerpo else ancho // 2)
                except Exception as e:
                    avisos.append(f"No pude insertar la imagen «{os.path.basename(str(imagen))}»: {e}.")
            else:
                avisos.append(f"No encontré la imagen «{imagen}».")

        if d.get("notas"):
            slide.notes_slide.notes_text_frame.text = str(d["notas"])



def _aplicar_estilo_pptx(prs, diapositivas, estilo: Dict[str, Any], avisos: List[str]) -> List[str]:
    """Aplica a las diapositivas creadas lo que el usuario pidió de aspecto. Return qué se aplicó.

    Fuente y colores van run por run (el tema de Office se aplica después con PowerPoint y
    solo toca el patrón, así que el formato explícito se conserva). Logo, pie y número son
    formas añadidas en las esquinas.
    """
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Cm, Pt

    if not estilo or not diapositivas:
        return []
    aplicado: List[str] = []
    fuente = estilo.get("fuente")
    color_titulos = estilo.get("color_titulos")
    color_texto = estilo.get("color_texto")
    total = len(diapositivas)
    ancho, alto = prs.slide_width, prs.slide_height

    ruta_logo = str(estilo.get("logo") or "")
    if ruta_logo and not os.path.isfile(ruta_logo):
        avisos.append(f"No encontré el logo «{ruta_logo}»; la presentación salió sin él.")
        ruta_logo = ""

    for numero, slide in enumerate(diapositivas, 1):
        titulo = slide.shapes.title
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            es_titulo = titulo is not None and shape.shape_id == titulo.shape_id
            for parrafo in shape.text_frame.paragraphs:
                for run in parrafo.runs:
                    if fuente:
                        run.font.name = str(fuente)
                    color = color_titulos if es_titulo else color_texto
                    if color:
                        run.font.color.rgb = RGBColor.from_string(color)
        if ruta_logo:
            try:
                ancho_logo = Cm(float(str(estilo.get("logo_ancho_cm") or 2.5).replace(",", ".")))
                slide.shapes.add_picture(ruta_logo, ancho - ancho_logo - Cm(0.5), Cm(0.4),
                                         width=ancho_logo)
            except Exception as e:
                avisos.append(f"No pude poner el logo: {e}.")
                ruta_logo = ""
        if estilo.get("pie"):
            caja = slide.shapes.add_textbox(Cm(0.8), alto - Cm(1.2), ancho * 2 // 3, Cm(0.8))
            p = caja.text_frame.paragraphs[0]
            p.text = str(estilo["pie"])
            p.font.size = Pt(10)
            p.font.color.rgb = RGBColor.from_string(color_texto or "7F7F7F")
            if fuente:
                p.font.name = str(fuente)
        if estilo.get("numeracion"):
            caja = slide.shapes.add_textbox(ancho - Cm(3.3), alto - Cm(1.2), Cm(2.5), Cm(0.8))
            p = caja.text_frame.paragraphs[0]
            p.text = f"{numero} / {total}"
            p.alignment = PP_ALIGN.RIGHT
            p.font.size = Pt(10)
            p.font.color.rgb = RGBColor.from_string(color_texto or "7F7F7F")
            if fuente:
                p.font.name = str(fuente)

    for clave, etiqueta in (("fuente", f"fuente {fuente}"), ("color_titulos", "color de títulos"),
                            ("color_texto", "color de texto"), ("pie", "pie"),
                            ("numeracion", "número de diapositiva")):
        if estilo.get(clave):
            aplicado.append(etiqueta)
    if ruta_logo:
        aplicado.append("logo")
    return aplicado


def _pptx(cabecera: Dict[str, Any], diapositivas: List[Dict[str, Any]], ruta: Path,
          plantilla: Optional[Path], reemplazos: Dict[str, str],
          estilo: Optional[Dict[str, Any]] = None) -> Resultado:
    from pptx import Presentation
    from pptx.util import Cm, Pt

    from pptx.util import Inches

    prs = Presentation(str(plantilla)) if plantilla else Presentation()
    if not plantilla:
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)   # panorámica 16:9
    avisos: List[str] = []
    sustituidos = _reemplazar_en_pptx(prs, reemplazos) if reemplazos else 0
    if reemplazos and not sustituidos:
        avisos.append("Ojo: ninguno de los marcadores que me diste aparece en la plantilla.")
    if plantilla and diapositivas:
        _quitar_diapositivas(prs)
    if plantilla:
        restantes = _marcadores_restantes_pptx(prs)
        if restantes:
            raise DocumentoRechazado(
                "La plantilla quedó con marcadores sin rellenar: "
                + "; ".join(f"{m} ({donde})" for m, donde in restantes)
                + ". Pasá su valor en 'reemplazos' (o preguntale al usuario) y volvé a pedirlo."
            )
    habia = len(prs.slides)

    if cabecera.get("titulo"):
        slide = prs.slides.add_slide(_layout(prs, "Title Slide", "Diapositiva de título",
                                             con_cuerpo=False))
        if slide.shapes.title is not None:
            slide.shapes.title.text = str(cabecera["titulo"])
        for ph in slide.placeholders:
            if ph.placeholder_format.idx == 1 and cabecera.get("subtitulo"):
                ph.text = str(cabecera["subtitulo"])

    _agregar_diapositivas(prs, diapositivas, avisos)
    aplicado = _aplicar_estilo_pptx(prs, list(prs.slides)[habia:], estilo or {}, avisos)

    prs.save(str(ruta))
    detalle = f"{len(prs.slides)} diapositivas"
    if sustituidos:
        detalle += f", {sustituidos} marcadores rellenados"
    if plantilla:
        detalle += f", con el tema de {plantilla.name}"
    if aplicado:
        detalle += ", estilo: " + ", ".join(aplicado)
    return Resultado(ruta=str(ruta), formato="pptx", detalle=detalle, avisos=avisos)


#: Tema de diseño de Office que se aplica cuando no hay plantilla. Sin esto, python-pptx
#: produce el blanco 4:3 de hace veinte años, y una presentación que parece de 2003 no es un
#: entregable aunque el contenido esté completo. "Retrospect" es sobrio: sirve para una
#: junta y para una clase.
TEMA_POR_DEFECTO = "Retrospect"

_CARPETAS_DE_TEMAS = (
    r"C:\Program Files\Microsoft Office\root\Document Themes 16",
    r"C:\Program Files (x86)\Microsoft Office\root\Document Themes 16",
    r"C:\Program Files\Microsoft Office\Document Themes 16",
)


def temas_disponibles() -> List[str]:
    """Return los nombres de los temas de diseño que Office tiene instalados."""
    for carpeta in _CARPETAS_DE_TEMAS:
        if os.path.isdir(carpeta):
            return sorted(os.path.splitext(f)[0] for f in os.listdir(carpeta)
                          if f.lower().endswith(".thmx"))
    return []


def _ruta_de_tema(nombre: str) -> Optional[str]:
    buscado = str(nombre or "").strip().lower()
    for carpeta in _CARPETAS_DE_TEMAS:
        if not os.path.isdir(carpeta):
            continue
        for archivo in os.listdir(carpeta):
            if archivo.lower().endswith(".thmx") and os.path.splitext(archivo)[0].lower() == buscado:
                return os.path.join(carpeta, archivo)
    return None


def _aplicar_tema(ruta_pptx: str, tema: str) -> bool:
    """Aplica un tema de diseño de Office a una presentación ya guardada. Return si pudo.

    Lo hace el propio PowerPoint: aplicar un tema es recolocar cada marcador según los
    diseños del tema, y eso python-pptx no lo sabe hacer.
    """
    thmx = _ruta_de_tema(tema)
    if not thmx or os.name != "nt":
        return False
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        return False

    pythoncom.CoInitialize()
    app = None
    try:
        app = win32com.client.DispatchEx("PowerPoint.Application")
        pres = app.Presentations.Open(ruta_pptx, ReadOnly=False, WithWindow=False)
        pres.ApplyTheme(thmx)
        pres.Save()
        pres.Close()
        return True
    except Exception as e:
        logger.warning(f"no se pudo aplicar el tema «{tema}»: {e}")
        return False
    finally:
        try:
            if app is not None:
                app.Quit()
        except Exception as e:
            logger.debug(f"PowerPoint no cerró limpio: {e}")
        pythoncom.CoUninitialize()


def _html_de_diapositivas(cabecera: Dict[str, Any], diapositivas: List[Dict[str, Any]]) -> str:
    """Return las diapositivas como páginas HTML apaisadas. Para el PDF sin PowerPoint."""
    import base64
    import html as _html

    e = _html.escape
    partes = ["<!doctype html><html><head><meta charset='utf-8'><style>"
              "body{font-family:Calibri,Arial,sans-serif;color:#222;margin:0}"
              ".d{page-break-after:always;padding:0.4cm 0.8cm}h1{font-size:26pt;margin:0 0 10pt 0;"
              "border-bottom:2px solid #305496;padding-bottom:6pt}h2{font-size:34pt;margin-top:3cm;"
              "text-align:center}p.sub{text-align:center;color:#555;font-size:16pt}"
              "li{font-size:15pt;margin:4pt 0}ul ul li{font-size:13pt}p{font-size:15pt}"
              "table{border-collapse:collapse;font-size:12pt;margin-top:8pt}"
              "td,th{border:1px solid #999;padding:4pt 8pt}th{background:#305496;color:#fff}"
              "</style></head><body>"]
    if cabecera.get("titulo"):
        partes.append(f"<div class='d'><h2>{e(str(cabecera['titulo']))}</h2>")
        if cabecera.get("subtitulo"):
            partes.append(f"<p class='sub'>{e(str(cabecera['subtitulo']))}</p>")
        partes.append("</div>")
    for d in diapositivas:
        partes.append(f"<div class='d'><h1>{e(str(d.get('titulo') or ''))}</h1>")
        puntos = d.get("puntos") or ([d["texto"]] if d.get("texto") else [])
        if puntos:
            partes.append("<ul>")
            for punto in puntos:
                if isinstance(punto, (list, tuple)):
                    partes.append("<ul>" + "".join(f"<li>{e(str(s))}</li>" for s in punto) + "</ul>")
                else:
                    partes.append(f"<li>{e(str(punto))}</li>")
            partes.append("</ul>")
        tabla = d.get("tabla")
        if isinstance(tabla, dict) and (tabla.get("filas") or tabla.get("encabezados")):
            enc = "".join(f"<th>{e(str(c))}</th>" for c in (tabla.get("encabezados") or []))
            filas = "".join("<tr>" + "".join(f"<td>{e('' if c is None else str(c))}</td>" for c in f)
                            + "</tr>" for f in (tabla.get("filas") or []))
            partes.append(f"<table>{'<tr>' + enc + '</tr>' if enc else ''}{filas}</table>")
        grafico = d.get("grafico")
        if isinstance(grafico, dict):
            with tempfile.TemporaryDirectory() as temporal:
                png = os.path.join(temporal, "g.png")
                if _grafico_como_imagen(grafico, png):
                    datos = base64.b64encode(Path(png).read_bytes()).decode("ascii")
                    partes.append(f"<img src='data:image/png;base64,{datos}' style='max-height:11cm'>")
        partes.append("</div>")
    partes.append("</body></html>")
    return "".join(partes)


def crear_presentacion(nombre: str, contenido: Any, plantilla: Optional[str] = None,
                       reemplazos: Optional[Dict[str, str]] = None,
                       carpeta: Optional[str] = None, tema: Optional[str] = None,
                       estilo: Any = None) -> Resultado:
    """Produce una presentación .pptx (o su .pdf). Levanta `DocumentoRechazado`.

    `contenido` trae `titulo`, `subtitulo` y `diapositivas`; cada una `titulo`, `puntos`
    (una lista dentro de la lista es un subnivel), `texto`, `tabla` (`encabezados`,
    `filas`), `imagen` (ruta) y `notas` para el orador. `estilo`: `fuente`,
    `color_titulos`, `color_texto`, `logo`, `pie`, `numeracion`.
    """
    extension = _extension(nombre)
    if extension not in EXTENSIONES_PRESENTACION + ("pdf",):
        raise DocumentoRechazado(f"«{nombre}» no es una presentación: hace falta .pptx (o .pdf).")
    reemplazos = {str(k): str(v) for k, v in (reemplazos or {}).items()}
    spec = _como_estructura(contenido)
    rechazar_relleno(spec)
    cabecera, diapositivas = _diapositivas_de(spec)
    _validar_presentacion(cabecera, diapositivas, con_reemplazos=bool(reemplazos))
    destino = carpeta_de_salida(carpeta)
    ruta = ruta_libre(destino, nombre)
    ruta_plantilla = _plantilla_valida(plantilla, ("pptx",)) if plantilla else None
    avisos_estilo: List[str] = []
    estilo_pres = _estilo_de(estilo, avisos_estilo)

    def _producir(destino_pptx: Path) -> Resultado:
        resultado = _pptx(cabecera, diapositivas, destino_pptx, ruta_plantilla, reemplazos,
                          estilo_pres)
        resultado.avisos = avisos_estilo + resultado.avisos
        if ruta_plantilla is None:
            elegido = str(tema or TEMA_POR_DEFECTO)
            if _ruta_de_tema(elegido) is None:
                resultado.avisos.append(
                    f"No existe el tema «{elegido}» en este Office; los que hay: "
                    f"{', '.join(temas_disponibles()[:12]) or 'ninguno'}. Quedó con el diseño básico."
                )
            elif _aplicar_tema(str(destino_pptx), elegido):
                resultado.detalle += f", tema {elegido}"
            else:
                resultado.avisos.append("No pude aplicar el tema de diseño (hace falta "
                                        "PowerPoint); quedó con el diseño básico.")
        return resultado

    if extension == "pptx":
        return _producir(ruta)

    with tempfile.TemporaryDirectory() as temporal:
        intermedio = Path(temporal) / (ruta.stem + ".pptx")
        resultado = _producir(intermedio)
        pdf_temporal = Path(temporal) / (ruta.stem + ".pdf")
        if _con_office("PowerPoint", str(intermedio), str(pdf_temporal)):
            os.replace(str(pdf_temporal), str(ruta))
            return Resultado(ruta=str(ruta), formato="pdf",
                             detalle=resultado.detalle + ", convertido con PowerPoint",
                             avisos=resultado.avisos)

    # Sin PowerPoint: cada diapositiva es una página apaisada con su título, sus puntos, su
    # tabla y su gráfico dibujado. No lleva el tema, y se dice.
    _pdf_apaisado_sin_office(_html_de_diapositivas(cabecera, diapositivas), ruta)
    avisos = [a for a in resultado.avisos if "tema" not in a.lower()]
    avisos.append("Sin PowerPoint el PDF no lleva el tema de diseño: cada diapositiva va "
                  "maquetada en limpio.")
    cuantas = len(diapositivas) + (1 if cabecera.get("titulo") else 0)
    return Resultado(ruta=str(ruta), formato="pdf",
                     detalle=f"{cuantas} diapositivas, maquetado sin Office", avisos=avisos)


# ============================================================================ otros

def crear_archivo(nombre: str, contenido: Any, carpeta: Optional[str] = None) -> Resultado:
    """Produce cualquier otro archivo de texto (.json, .csv suelto, .sql, .py, lo que sea)."""
    extension = _extension(nombre)
    if not extension:
        raise DocumentoRechazado("El nombre del archivo tiene que llevar extensión.")
    spec = _como_estructura(contenido)
    rechazar_relleno(spec)
    destino = carpeta_de_salida(carpeta)
    ruta = ruta_libre(destino, nombre)

    if extension in EXTENSIONES_DATOS and not isinstance(spec, str):
        ruta.write_text(json.dumps(spec, ensure_ascii=False, indent=2), encoding="utf-8")
        return Resultado(ruta=str(ruta), formato="json", detalle="JSON con sangría")

    texto = spec if isinstance(spec, str) else json.dumps(spec, ensure_ascii=False, indent=2)
    if not texto.strip():
        raise DocumentoRechazado("El archivo quedaría vacío: no hay contenido que escribir.")
    ruta.write_text(texto, encoding="utf-8")
    return Resultado(ruta=str(ruta), formato=extension,
                     detalle=f"{len(texto.splitlines())} líneas")


# ======================================================================= PDF con Office

def _con_office(aplicacion: str, origen: str, destino_pdf: str) -> bool:
    """Convierte con Word, Excel o PowerPoint en una instancia propia e invisible. Return si pudo.

    Instancia propia (`DispatchEx`) y no la que el usuario tenga abierta: cerrar la nuestra
    al terminar no puede llevarse la suya, ni su documento a medio escribir.
    """
    if os.name != "nt":
        return False
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        return False

    pythoncom.CoInitialize()
    app = None
    try:
        if aplicacion == "Word":
            app = win32com.client.DispatchEx("Word.Application")
            app.Visible = False
            doc = app.Documents.Open(origen, ReadOnly=True)
            doc.ExportAsFixedFormat(destino_pdf, 17)        # wdExportFormatPDF
            doc.Close(False)
        elif aplicacion == "Excel":
            app = win32com.client.DispatchEx("Excel.Application")
            app.Visible = False
            app.DisplayAlerts = False
            libro = app.Workbooks.Open(origen, ReadOnly=True)
            libro.ExportAsFixedFormat(0, destino_pdf)          # xlTypePDF
            libro.Close(False)
        elif aplicacion == "PowerPoint":
            app = win32com.client.DispatchEx("PowerPoint.Application")
            pres = app.Presentations.Open(origen, ReadOnly=True, WithWindow=False)
            pres.SaveAs(destino_pdf, 32)                       # ppSaveAsPDF
            pres.Close()
        else:
            return False
        return os.path.isfile(destino_pdf)
    except Exception as e:
        logger.warning(f"no se pudo convertir a PDF con {aplicacion}: {e}")
        return False
    finally:
        try:
            if app is not None:
                app.Quit()
        except Exception as e:
            logger.debug(f"{aplicacion} no cerró limpio: {e}")
        pythoncom.CoUninitialize()


def convertir_a_pdf(ruta: str, carpeta: Optional[str] = None) -> Resultado:
    """Convierte un .docx, .xlsx o .pptx existente a PDF con Office. Levanta si no se puede."""
    origen = Path(os.path.expanduser(str(ruta or "").strip()))
    if not origen.is_file():
        raise DocumentoRechazado(f"No encuentro «{ruta}».")
    aplicacion = {"docx": "Word", "doc": "Word", "xlsx": "Excel", "xls": "Excel",
                  "pptx": "PowerPoint", "ppt": "PowerPoint"}.get(_extension(origen.name))
    if aplicacion is None:
        raise DocumentoRechazado(f"No sé convertir «{origen.name}» a PDF: solo Word, Excel y PowerPoint.")
    destino = carpeta_de_salida(carpeta) if carpeta else origen.parent
    salida = ruta_libre(destino, origen.stem + ".pdf")
    if _con_office(aplicacion, str(origen), str(salida)):
        return Resultado(ruta=str(salida), formato="pdf", detalle=f"convertido con {aplicacion}")
    if aplicacion == "Excel":
        html, avisos = _html_de_hoja(origen)
        _pdf_apaisado_sin_office(html, salida)
        return Resultado(ruta=str(salida), formato="pdf", detalle="maquetado sin Office",
                         avisos=avisos + ["Sin Excel, los gráficos no entran en el PDF."])
    raise DocumentoRechazado(f"No pude convertir «{origen.name}» con {aplicacion}. "
                             f"¿Está instalado y sin diálogos abiertos?")


# ======================================================================== inspeccionar

_MARCADOR_RE = re.compile(r"\{\{[^}]{1,60}\}\}|\[[A-ZÁÉÍÓÚÑ_ ]{2,40}\]|<<[^>]{1,60}>>|«[A-ZÁÉÍÓÚÑ_ ]{2,40}»")


def inspeccionar(ruta: str) -> str:
    """Return qué tiene una plantilla: estilos, diseños, hojas y marcadores. Para usarla bien.

    Sin esto, seguir una plantilla es adivinar cómo se llaman sus estilos y sus marcadores.
    """
    archivo = Path(os.path.expanduser(str(ruta or "").strip()))
    if not archivo.is_file():
        for base in [_escritorio(), Path.home() / "Documents", Path.home() / "OneDrive" / "Documentos",
                     Path.home() / "Downloads"]:
            if (base / archivo).is_file():
                archivo = base / archivo
                break
    if not archivo.is_file():
        return f"No encuentro «{ruta}»."

    extension = _extension(archivo.name)
    try:
        if extension == "docx":
            return _inspeccionar_docx(archivo)
        if extension == "xlsx":
            return _inspeccionar_xlsx(archivo)
        if extension == "pptx":
            return _inspeccionar_pptx(archivo)
    except Exception as e:
        logger.warning(f"no se pudo inspeccionar {archivo}: {e}")
        return f"Pude abrir «{archivo.name}» pero no leer su estructura: {type(e).__name__}."
    return f"«{archivo.name}» no es un documento de Office que sepa inspeccionar (.docx, .xlsx, .pptx)."


def _marcadores_en(texto: str) -> List[str]:
    return sorted(set(_MARCADOR_RE.findall(texto or "")))


def _inspeccionar_docx(archivo: Path) -> str:
    from docx import Document

    doc = Document(str(archivo))
    usados = []
    textos = []
    for p in doc.paragraphs:
        if p.text.strip():
            textos.append(p.text.strip())
            if p.style is not None and p.style.name not in usados:
                usados.append(p.style.name)
    for t in doc.tables:
        for fila in t.rows:
            for c in fila.cells:
                textos.append(c.text)
    encabezados = []
    for s in doc.sections:
        for parte, nombre in ((s.header, "encabezado"), (s.footer, "pie")):
            contenido = " / ".join(p.text.strip() for p in parte.paragraphs if p.text.strip())
            if contenido:
                encabezados.append(f"{nombre}: {contenido[:120]}")
                textos.append(contenido)
    marcadores = _marcadores_en("\n".join(textos))
    lineas = [f"Plantilla Word «{archivo.name}»: {len(doc.paragraphs)} párrafos, "
              f"{len(doc.tables)} tablas, {len(doc.sections)} sección(es)."]
    if usados:
        lineas.append("Estilos que usa: " + ", ".join(usados[:15]))
    if encabezados:
        lineas += ["Encabezados y pies:"] + [f"  {e}" for e in encabezados]
    if marcadores:
        lineas.append("Marcadores para rellenar (pasalos en 'reemplazos'): " + ", ".join(marcadores))
    muestra = [t for t in textos if t][:6]
    if muestra:
        lineas += ["Empieza así:"] + [f"  {t[:100]}" for t in muestra]
    return "\n".join(lineas)


def _inspeccionar_xlsx(archivo: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(archivo), read_only=True)
    try:
        return _inspeccionar_libro(wb, archivo)
    finally:
        wb.close()


def _inspeccionar_libro(wb, archivo: Path) -> str:
    lineas = [f"Plantilla Excel «{archivo.name}»: {len(wb.sheetnames)} hoja(s)."]
    textos = []
    for ws in wb.worksheets:
        encabezado = []
        for fila in ws.iter_rows(min_row=1, max_row=1, values_only=True):
            encabezado = [str(c) for c in fila if c is not None]
        for fila in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 1, 30), values_only=True):
            textos += [str(c) for c in fila if isinstance(c, str)]
        lineas.append(f"  «{ws.title}»: {ws.max_row} filas × {ws.max_column} columnas"
                      + (f"; encabezado: {', '.join(encabezado[:12])}" if encabezado else ""))
    marcadores = _marcadores_en("\n".join(textos))
    if marcadores:
        lineas.append("Marcadores: " + ", ".join(marcadores))
    lineas.append("Para escribir en una hoja existente usá su mismo 'nombre' y, si hace falta, 'desde'.")
    return "\n".join(lineas)


def _inspeccionar_pptx(archivo: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(archivo))
    lineas = [f"Plantilla PowerPoint «{archivo.name}»: {len(prs.slides)} diapositivas, "
              f"{len(prs.slide_layouts)} diseños."]
    lineas.append("Diseños disponibles: " + ", ".join(f"«{l.name}»" for l in prs.slide_layouts))
    textos = []
    for i, slide in enumerate(prs.slides, 1):
        titulo = slide.shapes.title.text if slide.shapes.title is not None else ""
        for shape in slide.shapes:
            if shape.has_text_frame:
                textos.append(shape.text_frame.text)
        if i <= 8:
            lineas.append(f"  {i}. {titulo[:80] or '(sin título)'}")
    marcadores = _marcadores_en("\n".join(textos))
    if marcadores:
        lineas.append("Marcadores para rellenar (pasalos en 'reemplazos'): " + ", ".join(marcadores))
    return "\n".join(lineas)


# ======================================================================= editar existente

def _destino_de_edicion(origen: Path, en_sitio: bool, carpeta: Optional[str]) -> Path:
    """Return dónde se guarda el documento editado.

    Por defecto, a un archivo nuevo al lado del original (`informe (2).docx`): editar el
    documento de alguien y equivocarse no tiene deshacer. Si se pide `en_sitio`, se
    sobrescribe el original — y eso es lo que muestra la confirmación.
    """
    if en_sitio:
        return origen
    if carpeta:
        return ruta_libre(carpeta_de_salida(carpeta), origen.name)
    return ruta_libre(origen.parent, origen.name)


def _origen_valido(ruta: str) -> Path:
    archivo = Path(os.path.expanduser(str(ruta or "").strip()))
    if not archivo.is_file():
        for base in [_escritorio(), Path.home() / "Documents", Path.home() / "OneDrive" / "Documentos",
                     Path.home() / "Downloads"]:
            if (base / archivo).is_file():
                archivo = base / archivo
                break
    if not archivo.is_file():
        raise DocumentoRechazado(f"No encuentro «{ruta}».")
    real = Path(os.path.realpath(archivo))
    for base in _bases_permitidas():
        base_real = str(Path(os.path.realpath(base))).lower()
        try:
            if os.path.commonpath([str(real).lower(), base_real]) == base_real:
                return real
        except ValueError:
            continue
    raise DocumentoRechazado(
        f"«{archivo.name}» está fuera de tu carpeta personal y de los espacios de trabajo: "
        f"no lo edito."
    )


def _mover_al_final_despues_de(doc, ancla: str, habia_antes: set) -> bool:
    """Mueve lo agregado al final del cuerpo para que quede justo después del párrafo `ancla`.

    python-docx solo sabe añadir al final. Para "poné esto después de la sección 3" se
    escribe al final y luego se recolocan esos elementos detrás del párrafo que contiene
    el texto ancla. Return si se encontró el ancla.

    Lo nuevo se reconoce por identidad y no por posición: python-docx inserta cada párrafo
    ANTES del `sectPr` final, así que "los elementos a partir del índice N" era solo el
    `sectPr` y el contenido nuevo se quedaba al final. Pasó en la primera prueba.
    """
    from os_integration.ui_tree import _normalizar

    buscado = _normalizar(ancla)
    cuerpo = doc.element.body
    nuevos = [h for h in cuerpo if h not in habia_antes and not h.tag.endswith("}sectPr")]
    objetivo = None
    for parrafo in doc.paragraphs:
        if parrafo._p in nuevos:
            break
        if buscado and buscado in _normalizar(parrafo.text):
            objetivo = parrafo._p
    if objetivo is None:
        return False
    for elemento in reversed(nuevos):
        objetivo.addnext(elemento)
    return True


def _editar_docx(origen: Path, destino: Path, cambios: Dict[str, Any]) -> Resultado:
    from docx import Document

    doc = Document(str(origen))
    avisos: List[str] = []
    hechos: List[str] = []

    reemplazos = {str(k): str(v) for k, v in (cambios.get("reemplazos") or {}).items()}
    if reemplazos:
        cuantos = _reemplazar_en_docx(doc, reemplazos)
        hechos.append(f"{cuantos} reemplazos")
        if not cuantos:
            avisos.append("Ninguno de los textos a reemplazar aparece en el documento.")

    quitar = [str(q) for q in (cambios.get("quitar") or []) if str(q).strip()]
    if quitar:
        from os_integration.ui_tree import _normalizar

        borrados = 0
        for parrafo in list(doc.paragraphs):
            texto = _normalizar(parrafo.text)
            if texto and any(_normalizar(q) in texto for q in quitar):
                parrafo._p.getparent().remove(parrafo._p)
                borrados += 1
        hechos.append(f"{borrados} párrafos quitados")
        if not borrados:
            avisos.append("No encontré ningún párrafo con los textos a quitar.")

    agregar = cambios.get("agregar") or []
    despues_de = str(cambios.get("despues_de") or "").strip()
    if isinstance(agregar, dict) and ("bloques" in agregar or "despues_de" in agregar):
        # `agregar: {"despues_de": ..., "bloques": [...]}` es la forma natural de pedirlo,
        # aunque la documentada lleve `despues_de` al lado. Se aceptan las dos.
        despues_de = despues_de or str(agregar.get("despues_de") or "").strip()
        agregar = agregar.get("bloques") or agregar.get("contenido") or []
    _, bloques = _bloques_de(agregar if isinstance(agregar, str) else {"bloques": agregar})
    if bloques:
        rechazar_relleno(bloques)
        antes = set(doc.element.body)
        _escribir_bloques(doc, bloques, avisos)
        if despues_de and not _mover_al_final_despues_de(doc, despues_de, antes):
            avisos.append(f"No encontré «{despues_de}» para poner el contenido después; "
                          f"quedó al final.")
        hechos.append(f"{len(bloques)} bloques agregados"
                      + (f" después de «{despues_de}»" if despues_de else " al final"))

    if not hechos:
        raise DocumentoRechazado("No me diste ningún cambio: 'reemplazos', 'agregar' o 'quitar'.")
    doc.save(str(destino))
    return Resultado(ruta=str(destino), formato="docx", detalle=", ".join(hechos), avisos=avisos)


_TOTAL_RE = re.compile(r"^(total|totales|suma|sumatoria)\b", re.IGNORECASE)
#: Una referencia de celda o rango dentro de una fórmula, con su hoja si la trae. Los
#: lookarounds evitan tomar «LOG10(» o «DAYS360(» por referencias.
_REF_EN_FORMULA_RE = re.compile(
    r"(?<![A-Za-z0-9_])(?P<hoja>(?:'[^']+'|[A-Za-z0-9_\.]+)!)?"
    r"(?P<c1>\$?[A-Z]{1,3})(?P<f1>\$?\d{1,7})"
    r"(?::(?P<c2>\$?[A-Z]{1,3})(?P<f2>\$?\d{1,7}))?(?![A-Za-z0-9_(])"
)


def _fila_de_totales(ws) -> Optional[int]:
    """Return la fila de totales de una hoja, o None: de abajo hacia arriba, la primera cuya
    primera celda con texto diga «Total» o que traiga una SUM/SUBTOTAL sobre las de arriba."""
    for fila in range((ws.max_row or 0), 1, -1):
        celdas = next(ws.iter_rows(min_row=fila, max_row=fila))
        valores = [c.value for c in celdas]
        if all(v in (None, "") for v in valores):
            continue
        primer_texto = next((v for v in valores if isinstance(v, str) and v.strip()), "")
        if _TOTAL_RE.match(primer_texto.strip()):
            return fila
        if any(isinstance(v, str) and re.match(r"=\s*(SUM|SUBTOTAL)\(", v, re.IGNORECASE)
               for v in valores):
            return fila
        return None                                # la última fila con datos no es de totales
    return None


def _copiar_formato_de_fila(ws, origen: int, destinos) -> None:
    """Copia fuente, borde, relleno, alineación y formato de número de una fila a otras."""
    from copy import copy

    if origen < 1 or origen > (ws.max_row or 0):
        return
    for celda in next(ws.iter_rows(min_row=origen, max_row=origen)):
        for fila in destinos:
            destino = ws.cell(row=fila, column=celda.column)
            destino.font = copy(celda.font)
            destino.border = copy(celda.border)
            destino.fill = copy(celda.fill)
            destino.alignment = copy(celda.alignment)
            destino.number_format = celda.number_format


def _desplazar_referencias(formula: str, hoja: str, hoja_de_la_formula: str,
                           desde_fila: int, cuantas: int) -> str:
    """Return la fórmula con las referencias a `hoja` corridas `cuantas` filas a partir de
    `desde_fila`, y los rangos que terminaban justo encima extendidos para abarcar lo nuevo.

    openpyxl inserta filas sin tocar una sola fórmula: el `=SUM(C2:C3)` de la fila de
    totales seguiría sumando dos filas con tres, y el `=Clientes!C4` de otra hoja apuntaría
    a una fila de datos. Es lo que Excel hace solo al insertar filas dentro de un rango.
    """
    def _misma_hoja(prefijo: Optional[str]) -> bool:
        if not prefijo:
            return hoja_de_la_formula == hoja
        return prefijo[:-1].strip("'") == hoja

    def _corregir(m: re.Match) -> str:
        if not _misma_hoja(m.group("hoja")):
            return m.group(0)
        f1 = int(m.group("f1").replace("$", ""))
        nuevo_f1 = f1 + cuantas if f1 >= desde_fila else f1
        salida = ((m.group("hoja") or "") + m.group("c1")
                  + m.group("f1").replace(str(f1), str(nuevo_f1)))
        if m.group("c2"):
            f2 = int(m.group("f2").replace("$", ""))
            if f2 >= desde_fila:
                nuevo_f2 = f2 + cuantas
            elif f2 == desde_fila - 1 and f1 < desde_fila:
                nuevo_f2 = f2 + cuantas        # terminaba justo encima: se extiende
            else:
                nuevo_f2 = f2
            salida += ":" + m.group("c2") + m.group("f2").replace(str(f2), str(nuevo_f2))
        return salida

    # Las cadenas entre comillas no se tocan: "C4" dentro de un texto no es una referencia.
    partes = re.split(r'("[^"]*")', formula)
    return "".join(parte if parte.startswith('"') else _REF_EN_FORMULA_RE.sub(_corregir, parte)
                   for parte in partes)


def _insertar_filas_antes_del_total(wb, ws, fila_total: int, nuevas: List[List[Any]]) -> None:
    """Mete las filas nuevas justo antes de la de totales, con el formato de la última fila
    de datos, y corrige todas las fórmulas del libro (y los gráficos de la hoja)."""
    cuantas = len(nuevas)
    ws.insert_rows(fila_total, cuantas)
    for i, fila in enumerate(nuevas):
        for j, valor in enumerate(fila, 1):
            ws.cell(row=fila_total + i, column=j, value=_valor_de_celda(valor))
    if fila_total - 1 >= 1:
        _copiar_formato_de_fila(ws, fila_total - 1, range(fila_total, fila_total + cuantas))

    for hoja in wb.worksheets:
        for fila_celdas in hoja.iter_rows():
            for celda in fila_celdas:
                if isinstance(celda.value, str) and celda.value.startswith("="):
                    celda.value = _desplazar_referencias(celda.value, ws.title, hoja.title,
                                                         fila_total, cuantas)
    for grafico in getattr(ws, "_charts", []):
        for serie in grafico.series:
            for referencia in (getattr(serie.val, "numRef", None) if serie.val else None,
                               getattr(serie.cat, "numRef", None) if serie.cat else None,
                               getattr(serie.cat, "strRef", None) if serie.cat else None):
                if referencia is not None and referencia.f:
                    referencia.f = _desplazar_referencias(referencia.f, ws.title, ws.title,
                                                          fila_total, cuantas)


def _editar_xlsx(origen: Path, destino: Path, cambios: Dict[str, Any]) -> Resultado:
    from openpyxl import load_workbook

    wb = load_workbook(str(origen))
    avisos: List[str] = []
    hechos: List[str] = []

    celdas = cambios.get("celdas") or {}
    if celdas:
        escritas = 0
        for referencia, valor in celdas.items():
            hoja, _, celda = str(referencia).rpartition("!")
            ws = wb[hoja] if hoja and hoja in wb.sheetnames else wb.worksheets[0]
            if hoja and hoja not in wb.sheetnames:
                avisos.append(f"No hay una hoja «{hoja}»; escribí en «{ws.title}».")
            try:
                ws[celda.upper()] = _valor_de_celda(valor)
                escritas += 1
            except Exception as e:
                avisos.append(f"No pude escribir en {referencia}: {e}")
        hechos.append(f"{escritas} celdas escritas")

    agregar = cambios.get("agregar") or {}
    if isinstance(agregar, list):
        agregar = {wb.worksheets[0].title: agregar}
    for nombre_hoja, filas in (agregar or {}).items():
        ws = wb[nombre_hoja] if nombre_hoja in wb.sheetnames else wb.create_sheet(str(nombre_hoja)[:31])
        nuevas = [list(f) if isinstance(f, (list, tuple)) else [f] for f in (filas or [])]
        if nuevas:
            fila_total = _fila_de_totales(ws)
            if fila_total is not None:
                _insertar_filas_antes_del_total(wb, ws, fila_total, nuevas)
                hechos.append(f"{len(nuevas)} filas agregadas a «{ws.title}» antes de la fila "
                              f"de totales")
            else:
                ultima = ws.max_row
                for fila in nuevas:
                    ws.append([_valor_de_celda(v) for v in fila])
                _copiar_formato_de_fila(ws, ultima, range(ultima + 1, ultima + 1 + len(nuevas)))
                hechos.append(f"{len(nuevas)} filas agregadas a «{ws.title}»")

    reemplazos = {str(k): str(v) for k, v in (cambios.get("reemplazos") or {}).items()}
    if reemplazos:
        cuantos = 0
        for ws in wb.worksheets:
            for fila in ws.iter_rows():
                for celda in fila:
                    if isinstance(celda.value, str) and any(m in celda.value for m in reemplazos):
                        nuevo = celda.value
                        for m, v in reemplazos.items():
                            cuantos += nuevo.count(m)
                            nuevo = nuevo.replace(m, v)
                        celda.value = _valor_de_celda(nuevo)
        hechos.append(f"{cuantos} reemplazos")

    if not hechos:
        raise DocumentoRechazado("No me diste ningún cambio: 'celdas', 'agregar' o 'reemplazos'.")
    wb.save(str(destino))
    wb.close()
    return Resultado(ruta=str(destino), formato="xlsx", detalle=", ".join(hechos), avisos=avisos)


def _renumerar_diapositivas(prs) -> None:
    """Deja las partes de las diapositivas numeradas seguidas: slide1, slide2, slide3...

    python-pptx nombra una diapositiva nueva por CANTIDAD (`slide{n+1}`), no mirando qué
    nombres ya existen. Tras borrar la 2 de tres, la siguiente que se agrega se llama
    `slide3.xml`, que es la que ya está, y el archivo sale con dos partes del mismo nombre:
    PowerPoint lo abre roto. Se vio en la primera prueba: «Cierre» desapareció y «Anexo»
    salió dos veces. Con las partes renumeradas después de borrar, no hay choque.
    """
    from pptx.opc.packuri import PackURI

    # Primero a nombres provisionales, para no chocar a mitad del cambio.
    for i, slide in enumerate(prs.slides, 1):
        slide.part.partname = PackURI(f"/ppt/slides/tmp{i}.xml")
    for i, slide in enumerate(prs.slides, 1):
        slide.part.partname = PackURI(f"/ppt/slides/slide{i}.xml")


def _editar_pptx(origen: Path, destino: Path, cambios: Dict[str, Any]) -> Resultado:
    from pptx import Presentation

    prs = Presentation(str(origen))
    avisos: List[str] = []
    hechos: List[str] = []

    reemplazos = {str(k): str(v) for k, v in (cambios.get("reemplazos") or {}).items()}
    if reemplazos:
        cuantos = _reemplazar_en_pptx(prs, reemplazos)
        hechos.append(f"{cuantos} reemplazos")
        if not cuantos:
            avisos.append("Ninguno de los textos a reemplazar aparece en la presentación.")

    quitar = [str(q) for q in (cambios.get("quitar") or []) if str(q).strip()]
    if quitar:
        from os_integration.ui_tree import _normalizar

        lista = prs.slides._sldIdLst
        borradas = 0
        for indice in reversed(range(len(prs.slides))):
            slide = prs.slides[indice]
            titulo = _normalizar(slide.shapes.title.text if slide.shapes.title is not None else "")
            numero = str(indice + 1)
            if any(q == numero or (titulo and _normalizar(q) in titulo) for q in quitar):
                sld_id = list(lista)[indice]
                prs.part.drop_rel(sld_id.rId)
                lista.remove(sld_id)
                borradas += 1
        hechos.append(f"{borradas} diapositivas quitadas")
        if borradas:
            _renumerar_diapositivas(prs)

    _, diapositivas = _diapositivas_de({"diapositivas": cambios.get("agregar") or []})
    if diapositivas:
        rechazar_relleno(diapositivas)
        _agregar_diapositivas(prs, diapositivas, avisos)
        hechos.append(f"{len(diapositivas)} diapositivas agregadas al final")

    if not hechos:
        raise DocumentoRechazado("No me diste ningún cambio: 'reemplazos', 'agregar' o 'quitar'.")
    prs.save(str(destino))
    return Resultado(ruta=str(destino), formato="pptx", detalle=", ".join(hechos), avisos=avisos)


def editar_documento(ruta: str, cambios: Any, en_sitio: bool = False,
                     carpeta: Optional[str] = None) -> Resultado:
    """Edita un .docx, .xlsx o .pptx que ya existe. Levanta `DocumentoRechazado`.

    `cambios` según el tipo:
    - Word: `reemplazos` {texto: nuevo}, `agregar` (bloques o Markdown) con `despues_de`
      (texto del párrafo tras el que va), `quitar` [textos de párrafos].
    - Excel: `celdas` {"Hoja!B4": valor}, `agregar` {hoja: [filas]}, `reemplazos`.
    - PowerPoint: `reemplazos`, `agregar` [diapositivas], `quitar` [títulos o números].
    """
    origen = _origen_valido(ruta)
    spec = _como_estructura(cambios)
    if not isinstance(spec, dict):
        raise DocumentoRechazado("Los cambios tienen que ser un objeto con 'reemplazos', "
                                 "'agregar', 'quitar' o 'celdas'.")
    rechazar_relleno({k: v for k, v in spec.items() if k != "quitar"})
    destino = _destino_de_edicion(origen, bool(en_sitio), carpeta)
    extension = _extension(origen.name)

    if extension == "docx":
        resultado = _editar_docx(origen, destino, spec)
    elif extension == "xlsx":
        resultado = _editar_xlsx(origen, destino, spec)
    elif extension == "pptx":
        resultado = _editar_pptx(origen, destino, spec)
    else:
        raise DocumentoRechazado(f"No sé editar «{origen.name}»: solo .docx, .xlsx y .pptx.")

    if not en_sitio:
        resultado.detalle += f"; el original {origen.name} quedó intacto"
    return resultado
