"""Validación con Office real de REQ-066: cambiar el aspecto de algo que ya existe."""
import sys
import logging
from pathlib import Path

logging.basicConfig(level=logging.WARNING, format="%(levelname)s: %(message)s")
sys.path.insert(0, r"C:\Users\WHOAMI\Documents\Apps\agent-ia\agente_ia")
SALIDA = Path(__file__).parent / "salida3"
SALIDA.mkdir(exist_ok=True)
from core import documentos as D
D._bases_permitidas = lambda: [SALIDA]
D.CARPETAS_VEDADAS = frozenset()


def paso(t):
    print("\n" + "=" * 78 + "\n" + t + "\n" + "=" * 78)


def render(archivo, prefijo, paginas=2):
    import fitz
    pdf = D.convertir_a_pdf(str(archivo), carpeta=str(SALIDA)).ruta
    doc = fitz.open(pdf)
    for i, page in enumerate(doc):
        if i >= paginas:
            break
        page.get_pixmap(dpi=70).save(str(SALIDA / f"{prefijo}_p{i + 1}.png"))
    doc.close()


from matplotlib import pyplot as plt
fig = plt.figure(figsize=(2, 0.8)); fig.text(0.1, 0.4, "AGRO SAS", fontsize=20, color="#1B5E20", weight="bold")
fig.savefig(str(SALIDA / "logo.png"), dpi=100); plt.close(fig)
LOGO = str(SALIDA / "logo.png")

paso("1. Word ya escrito, sin estilo → «ponele el membrete, Arial, horizontal y numerado»")
informe = {"titulo": "Informe de producción", "bloques": [
    {"tipo": "titulo", "texto": "Resumen", "nivel": 1},
    {"tipo": "parrafo", "texto": "La producción del trimestre subió 12,5 % frente al anterior, impulsada por el lote norte. El costo por hectárea se mantuvo estable."},
    {"tipo": "titulo", "texto": "Detalle", "nivel": 2},
    {"tipo": "parrafo", "texto": "El lote sur tuvo afectación por lluvias y se replantea en octubre con otra variedad."}]}
plano = D.crear_documento("informe.docx", informe, carpeta=str(SALIDA))
print("  creado sin estilo:", plano.describir())
r = D.editar_documento(plano.ruta, estilo={
    "fuente": "Arial", "tamano": 11, "fuente_titulos": "Georgia", "color_titulos": "verde oscuro",
    "orientacion": "horizontal", "margenes_cm": 2, "encabezado": "AGRO SAS · Informe interno",
    "pie": "Confidencial", "logo": LOGO, "numeracion": True, "justificado": True}, carpeta=str(SALIDA))
print(" ", r.describir())
from docx import Document
d = Document(r.ruta)
print("  Normal:", d.styles["Normal"].font.name, d.styles["Normal"].font.size.pt,
      "| Heading 1:", d.styles["Heading 1"].font.name, d.styles["Heading 1"].font.color.rgb)
print("  horizontal:", d.sections[0].page_width > d.sections[0].page_height,
      "| encabezado:", [p.text for p in d.sections[0].header.paragraphs if p.text])
render(r.ruta, "word_restilado", 1)

paso("2. Sin cambios ni estilo → se rechaza (antes reventaba con 'ningún cambio' solo por cambios)")
try:
    D.editar_documento(plano.ruta, {}, carpeta=str(SALIDA))
except D.DocumentoRechazado as e:
    print("  rechazado:", e)

paso("3. El estilo dentro de 'cambios' (donde el modelo lo pone a veces) también vale")
r = D.editar_documento(plano.ruta, {"estilo": {"color_titulos": "rojo"}}, carpeta=str(SALIDA))
print(" ", r.describir())
print("  Heading 1:", Document(r.ruta).styles["Heading 1"].font.color.rgb)

paso("4. PowerPoint con estilo → agregar diapositivas → ¿heredan y se renumera?")
pres = D.crear_presentacion("p.pptx", {"titulo": "Resultados", "diapositivas": [
    {"titulo": "Uno", "puntos": ["Primer punto", "Segundo punto"]},
    {"titulo": "Dos", "texto": "Contenido dos."}]}, carpeta=str(SALIDA), tema="Ion",
    estilo={"fuente": "Georgia", "color_titulos": "#FFD54F", "logo": LOGO,
            "pie": "Agro SAS · Confidencial", "numeracion": True})
print("  creada:", pres.describir())
from pptx import Presentation
def numeros(ruta):
    return [[s.text_frame.text for s in slide.shapes
             if s.has_text_frame and str(s.name).endswith("numero")] for slide in Presentation(ruta).slides]
print("  números:", numeros(pres.ruta))
r = D.editar_documento(pres.ruta, {"agregar": [{"titulo": "Tres", "texto": "Contenido tres."}]}, carpeta=str(SALIDA))
print(" ", r.describir())
print("  números tras agregar:", numeros(r.ruta))
ultima = Presentation(r.ruta).slides[-1]
print("  la nueva tiene pie/numero:", sorted(str(s.name) for s in ultima.shapes if str(s.name).startswith("ORION-estilo-")))
render(r.ruta, "pptx_agregada", 4)

paso("5. Reaplicar estilo a la misma presentación: ¿duplica logo/pie?")
antes = [len([s for s in sl.shapes if str(s.name).startswith("ORION-estilo-")]) for sl in Presentation(r.ruta).slides]
r2 = D.editar_documento(r.ruta, estilo={"fuente": "Georgia", "color_titulos": "azul claro",
                                        "logo": LOGO, "pie": "Agro SAS · Interno", "numeracion": True},
                        carpeta=str(SALIDA))
print(" ", r2.describir())
despues = [len([s for s in sl.shapes if str(s.name).startswith("ORION-estilo-")]) for sl in Presentation(r2.ruta).slides]
print("  formas de estilo por diapositiva antes:", antes, "después:", despues)
print("  pies:", [[s.text_frame.text for s in sl.shapes if str(s.name).endswith('pie')] for sl in Presentation(r2.ruta).slides])

paso("6. Excel ya escrito → «ponele letra Calibri 10 y el encabezado en verde»")
libro = D.crear_hoja("v.xlsx", {"hojas": [{"nombre": "Ventas", "titulo": "AGRO SAS — Ventas",
    "encabezados": ["Producto", "Cantidad", "Total"],
    "filas": [["Abono", "120", "4.260.000"], ["Semilla", "40", "34.800.000"]], "totales": True,
    "formatos": {"C": '"$" #,##0'}}]}, carpeta=str(SALIDA))
print("  creado:", libro.describir())
r = D.editar_documento(libro.ruta, estilo={"fuente": "Calibri", "tamano": 10,
                                           "color_encabezado": "verde oscuro"}, carpeta=str(SALIDA))
print(" ", r.describir())
from openpyxl import load_workbook
ws = load_workbook(r.ruta)["Ventas"]
print("  A1 (título):", ws["A1"].font.name, ws["A1"].font.size, "negrita:", ws["A1"].font.b)
print("  A3 (encabezado):", ws["A3"].value, ws["A3"].fill.fgColor.rgb, ws["A3"].font.name)
print("  A4 (dato):", ws["A4"].value, ws["A4"].font.name, ws["A4"].font.size)
print("  C6 (total):", ws["C6"].value, "negrita:", ws["C6"].font.b)
render(r.ruta, "excel_restilado", 1)

paso("7. Insertar filas con celda combinada y formato condicional debajo del total")
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
wb = Workbook(); ws = wb.active; ws.title = "Caja"
ws.append(["Concepto", "Valor"]); ws.append(["Uno", 100]); ws.append(["Dos", 200])
ws.append(["Total", "=SUM(B2:B3)"])
ws.merge_cells("A6:B6"); ws["A6"] = "Nota al pie de la tabla"
ws.conditional_formatting.add("B2:B3", CellIsRule(operator="greaterThan", formula=["150"],
                                                  stopIfTrue=False, font=None))
origen = SALIDA / "caja.xlsx"; wb.save(str(origen))
r = D.editar_documento(str(origen), {"agregar": {"Caja": [["Tres", "300"]]}}, en_sitio=True)
print(" ", r.describir())
ws = load_workbook(str(origen))["Caja"]
print("  filas:", [[c.value for c in ws[f]] for f in range(1, 8)])
print("  combinadas:", [str(x) for x in ws.merged_cells.ranges])
print("  formato condicional:", [str(f.cells) for f in ws.conditional_formatting])
