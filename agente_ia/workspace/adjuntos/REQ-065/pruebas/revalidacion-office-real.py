"""Revalidación con Office real de lo corregido en REQ-065."""
import sys
import logging
from pathlib import Path

logging.basicConfig(level=logging.WARNING, format="%(levelname)s: %(message)s")
sys.path.insert(0, r"C:\Users\WHOAMI\Documents\Apps\agent-ia\agente_ia")
SALIDA = Path(__file__).parent / "salida2"
SALIDA.mkdir(exist_ok=True)
from core import documentos as D
D._bases_permitidas = lambda: [SALIDA]
D.CARPETAS_VEDADAS = frozenset()


def paso(t):
    print("\n" + "=" * 78 + "\n" + t + "\n" + "=" * 78)


def render(pdf, prefijo, paginas=1):
    import fitz
    doc = fitz.open(str(pdf))
    for i, page in enumerate(doc):
        if i >= paginas:
            break
        page.get_pixmap(dpi=70).save(str(SALIDA / f"{prefijo}_p{i + 1}.png"))
    doc.close()


# plantilla Excel igual a la de la validación original
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
wb = Workbook(); ws = wb.active; ws.title = "Ventas"
ws.merge_cells("A1:E1"); ws["A1"] = "AGRO SAS — Reporte de ventas"; ws["A1"].font = Font(bold=True, size=16, color="1B5E20")
ws["A2"] = "Generado por el área comercial"
for i, h in enumerate(["Producto", "Cantidad", "Precio", "Subtotal", "IVA"], 1):
    c = ws.cell(row=4, column=i, value=h); c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="1B5E20"); c.alignment = Alignment(horizontal="center")
ws.column_dimensions["A"].width = 30
tpl = SALIDA / "plantilla_ventas.xlsx"; wb.save(str(tpl))

paso("1. Excel sobre plantilla, desde A5, fórmulas EN ESPAÑOL, gráfico por nombre → PDF con Excel")
hoja = {"hojas": [{
    "nombre": "Ventas", "desde": "A5",
    "encabezados": ["Producto", "Cantidad", "Precio", "Subtotal", "IVA"],
    "filas": [["Abono orgánico", "120", "35.500", "=B5*C5", "=D5*0,19"],
              ["Semilla maíz", "40", "870.000", "=B6*C6", "=REDONDEAR(D6*0,19;0)"],
              ["Riego por goteo", "3", "1.250.000,50", "=B7*C7", "=SI(D7>1000000;D7*0,19;0)"]],
    "totales": True,
    "formatos": {"C": '"$" #,##0', "D": '"$" #,##0', "E": '"$" #,##0'},
    "grafico": {"tipo": "barras", "titulo": "Subtotal por producto", "columnas": ["Subtotal"], "categorias": "Producto"},
}]}
r = D.crear_hoja("ventas.pdf", hoja, plantilla=str(tpl), carpeta=str(SALIDA))
print(r.describir())
render(r.ruta, "ventas")
r = D.crear_hoja("ventas.xlsx", hoja, plantilla=str(tpl), carpeta=str(SALIDA))
print(r.describir())
ws = load_workbook(r.ruta)["Ventas"]
for row in ws.iter_rows(min_row=4, max_row=ws.max_row, values_only=True):
    print("  ", row)

paso("2. Valores calculados por Excel de verdad")
import win32com.client, pythoncom
pythoncom.CoInitialize()
app = win32com.client.DispatchEx("Excel.Application"); app.Visible = False; app.DisplayAlerts = False
libro = app.Workbooks.Open(r.ruta, ReadOnly=True); h = libro.Worksheets("Ventas")
for celda in ("D5", "E5", "E6", "E7", "D8", "E8"):
    print(f"  {celda}: {h.Range(celda).Formula}  =  {h.Range(celda).Text}")
libro.Close(False); app.Quit()

paso("3. Editar: agregar cliente antes del total, con otra hoja apuntando al total")
origen = D.crear_hoja("clientes.xlsx", {"hojas": [
    {"nombre": "Clientes", "encabezados": ["Nombre", "Ciudad", "Deuda"],
     "filas": [["Finca La Esperanza", "Ibagué", "4.350.000"], ["Hacienda El Roble", "Neiva", "250.000"]],
     "totales": True, "formatos": {"C": '"$" #,##0'},
     "grafico": {"tipo": "barras", "titulo": "Deuda", "columnas": ["Deuda"], "categorias": "Nombre"}},
    {"nombre": "Resumen", "encabezados": ["Concepto", "Valor"], "filas": [["Cartera total", "=Clientes!C4"]]}]},
    carpeta=str(SALIDA)).ruta
r = D.editar_documento(origen, {"agregar": {"Clientes": [["Vivero Los Pinos", "Cali", "1.100.000"]]}}, carpeta=str(SALIDA))
print(r.describir())
pythoncom.CoInitialize()
app = win32com.client.DispatchEx("Excel.Application"); app.Visible = False; app.DisplayAlerts = False
libro = app.Workbooks.Open(r.ruta, ReadOnly=True)
h = libro.Worksheets("Clientes")
for f in range(1, 6):
    print("  ", [h.Cells(f, c).Text for c in range(1, 4)])
print("  Resumen!B2:", libro.Worksheets("Resumen").Range("B2").Formula, "=", libro.Worksheets("Resumen").Range("B2").Text)
libro.Close(False); app.Quit()
r2 = D.convertir_a_pdf(r.ruta, carpeta=str(SALIDA)); render(r2.ruta, "clientes")

paso("4. Presentación con estilo (Georgia, verde, logo, pie, numeración) sobre tema Ion → PDF")
from matplotlib import pyplot as plt
fig = plt.figure(figsize=(2, 0.8)); fig.text(0.1, 0.4, "AGRO SAS", fontsize=20, color="#1B5E20", weight="bold")
fig.savefig(str(SALIDA / "logo.png"), dpi=100); plt.close(fig)
r = D.crear_presentacion("resultados.pdf", {"titulo": "Resultados Trimestre 3", "subtitulo": "Agro SAS", "diapositivas": [
    {"titulo": "Qué pasó", "puntos": ["Ventas +12,5 %", "Nuevo cliente: Finca La Esperanza", ["Contrato a 12 meses"]]},
    {"titulo": "Cifras", "tabla": {"encabezados": ["Producto", "Ventas"], "filas": [["Abono", "$ 4.260.000"], ["Semilla", "$ 34.800.000"]]}},
]}, tema="Ion", carpeta=str(SALIDA), estilo={"fuente": "Georgia", "color_titulos": "#FFD54F", "logo": str(SALIDA / "logo.png"), "pie": "Agro SAS · Confidencial", "numeracion": True})
print(r.describir())
render(r.ruta, "pres", 2)

paso("5. Word con plantilla que tiene {{cliente}} en el pie, sin reemplazos → rechazo")
from docx import Document
doc = Document(); doc.sections[0].footer.paragraphs[0].text = "Cliente: {{cliente}}"; doc.add_paragraph("x")
doc.save(str(SALIDA / "membrete.docx"))
informe = {"titulo": "Informe", "bloques": [{"tipo": "parrafo", "texto": "Contenido real del informe " * 8}]}
try:
    D.crear_documento("i.docx", informe, plantilla=str(SALIDA / "membrete.docx"), carpeta=str(SALIDA))
except D.DocumentoRechazado as e:
    print("  rechazado:", e)
print("  ", D.crear_documento("i.docx", informe, plantilla=str(SALIDA / "membrete.docx"),
                             reemplazos={"{{cliente}}": "Palmar"}, carpeta=str(SALIDA)).describir())
