"""Validación real de REQ-040: plantillas, estilos, PDF con Office, y render a PNG."""
import os
import sys
import json
import logging
from pathlib import Path

logging.basicConfig(level=logging.WARNING)
sys.path.insert(0, r"C:\Users\WHOAMI\Documents\Apps\agent-ia\agente_ia")

SALIDA = Path(__file__).parent / "salida"
SALIDA.mkdir(exist_ok=True)

from core import documentos as D

# El scratchpad está bajo AppData: se abre solo para esta validación (no toca el proyecto).
D._bases_permitidas = lambda: [SALIDA]
D.CARPETAS_VEDADAS = frozenset()


def paso(titulo):
    print("\n" + "=" * 78 + "\n" + titulo + "\n" + "=" * 78)


def render_pdf(pdf, prefijo, paginas=2):
    import fitz
    doc = fitz.open(str(pdf))
    for i, page in enumerate(doc):
        if i >= paginas:
            break
        pix = page.get_pixmap(dpi=70)
        out = SALIDA / f"{prefijo}_p{i + 1}.png"
        pix.save(str(out))
        print("  render:", out.name)
    doc.close()


# --------------------------------------------------------------------------- plantilla Word
paso("A. Plantilla Word corporativa (membrete, estilos propios, marcadores partidos en runs)")
from docx import Document
from docx.shared import Pt, RGBColor

tpl = Document()
tpl.styles["Normal"].font.name = "Arial"
tpl.styles["Normal"].font.size = Pt(11)
tpl.styles["Heading 1"].font.color.rgb = RGBColor(0x2E, 0x7D, 0x32)   # verde corporativo
tpl.styles["Heading 1"].font.name = "Georgia"
sec = tpl.sections[0]
hp = sec.header.paragraphs[0]
hp.text = "AGRO SAS · NIT 900.123.456-7"
hp.runs[0].bold = True
sec.footer.paragraphs[0].text = "Calle 10 # 5-20, Bogotá · www.agrosas.co · Cliente: {{cliente}}"
p = tpl.add_paragraph()
r1 = p.add_run("Cotización N.º {{")
r1.bold = True
r2 = p.add_run("numero")
r2.bold = True
r2.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
p.add_run("}}")
t = tpl.add_table(rows=3, cols=2)
t.style = "Table Grid"
t.cell(0, 0).text = "Cliente"
t.cell(0, 1).text = "{{cliente}}"
t.cell(1, 0).text = "Fecha"
t.cell(1, 1).text = "[FECHA]"
t.cell(2, 0).text = "Total"
t.cell(2, 1).text = "{{total}}"
tpl.add_paragraph("Estimado {{cliente}}: adjuntamos la cotización solicitada, válida por 15 días.")
tpl.add_heading("Condiciones", level=1)
tpl.add_paragraph("Pago a 30 días. Precios sin IVA.")
ruta_tpl = SALIDA / "plantilla_cotizacion.docx"
tpl.save(str(ruta_tpl))
print("plantilla:", ruta_tpl)

paso("B. document_inspect sobre la plantilla")
print(D.inspeccionar(str(ruta_tpl)))

paso("C. Rellenar marcadores (cotización) — solo reemplazos, sin bloques")
r = D.crear_documento("cotizacion.docx", {"bloques": []}, plantilla=str(ruta_tpl),
                      reemplazos={"{{cliente}}": "Finca La Esperanza", "{{numero}}": "2026-0147",
                                  "[FECHA]": "21 de septiembre de 2026", "{{total}}": "$ 4.350.000"},
                      carpeta=str(SALIDA))
print(r.describir())
d = Document(r.ruta)
print("  párrafos cuerpo:", [p.text for p in d.paragraphs if p.text])
print("  runs del 1er párrafo:", [(x.text, x.bold, str(x.font.color.rgb) if x.font.color and x.font.color.rgb else None) for x in d.paragraphs[0].runs])
print("  tabla:", [[c.text for c in row.cells] for row in d.tables[0].rows])
print("  header:", d.sections[0].header.paragraphs[0].text, "| footer:", d.sections[0].footer.paragraphs[0].text)
restos = [p.text for p in d.paragraphs if "{{" in p.text or "[FECHA]" in p.text]
print("  marcadores sin rellenar en cuerpo:", restos or "ninguno")

paso("D. Contenido nuevo dentro de la plantilla (membrete + estilos heredados)")
informe = {
    "titulo": "Informe de producción — Trimestre 3",
    "subtitulo": "Finca La Esperanza",
    "autor": "Área técnica", "fecha": "21/09/2026",
    "bloques": [
        {"tipo": "titulo", "texto": "Resumen", "nivel": 1},
        {"tipo": "parrafo", "texto": "La producción del trimestre subió 12,5 % frente al anterior, "
                                     "impulsada por el lote norte. El costo por hectárea se mantuvo "
                                     "estable gracias a la renegociación de insumos."},
        {"tipo": "lista", "items": ["Lote norte: 48 t", "Lote sur: 31 t", ["Afectado por lluvias", "Replanteo en octubre"]]},
        {"tipo": "titulo", "texto": "Cifras", "nivel": 1},
        {"tipo": "tabla", "titulo": "Producción por lote", "encabezados": ["Lote", "Toneladas", "Costo/ha"],
         "filas": [["Norte", "48", "$ 1.250.000"], ["Sur", "31", "$ 1.310.000"]]},
        {"tipo": "grafico", "tipo_grafico": "barras", "titulo": "Toneladas por lote",
         "categorias": ["Norte", "Sur"], "series": [{"nombre": "Toneladas", "valores": [48, 31]}], "pie": "Figura 1"},
        {"tipo": "cita", "texto": "Lo que no se mide no se mejora."},
    ],
}
r = D.crear_documento("informe.docx", informe, plantilla=str(ruta_tpl), carpeta=str(SALIDA))
print(r.describir())
d = Document(r.ruta)
print("  header conservado:", d.sections[0].header.paragraphs[0].text)
print("  footer (con marcador sin rellenar?):", d.sections[0].footer.paragraphs[0].text)
print("  Normal font:", d.styles["Normal"].font.name, d.styles["Normal"].font.size)
print("  Heading 1 color/font:", d.styles["Heading 1"].font.color.rgb, d.styles["Heading 1"].font.name)
print("  primeros párrafos:", [(p.text[:40], p.style.name) for p in d.paragraphs[:8]])
print("  ¿quedó la tabla vieja de la plantilla?:", len(d.tables), "tablas")

paso("E. Mismo informe a PDF con Word y render")
r = D.crear_documento("informe.pdf", informe, plantilla=str(ruta_tpl), carpeta=str(SALIDA))
print(r.describir())
render_pdf(r.ruta, "informe_docx", paginas=2)

paso("E2. Documento SIN plantilla, guiado por el usuario (¿puede pedir fuente/colores?)")
r = D.crear_documento("sin_plantilla.pdf", informe, carpeta=str(SALIDA))
print(r.describir())
render_pdf(r.ruta, "sin_plantilla", paginas=1)

paso("F. Markdown → docx")
md = """# Acta de reunión

## Asistentes
- Johan
- María

## Acuerdos
1. Entregar el informe el viernes.
2. Revisar la cotización de insumos con dos proveedores más antes de decidir.

| Tema | Responsable |
|---|---|
| Informe | Johan |
| Cotización | María |
"""
r = D.crear_documento("acta.docx", md, carpeta=str(SALIDA))
print(r.describir())
d = Document(r.ruta)
print("  ", [(p.text[:30], p.style.name) for p in d.paragraphs[:8]], "tablas:", len(d.tables))

# --------------------------------------------------------------------------- Excel
paso("G. Plantilla Excel con encabezado corporativo propio y hoja 'Ventas'")
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
wb = Workbook()
ws = wb.active
ws.title = "Ventas"
ws.merge_cells("A1:E1")
ws["A1"] = "AGRO SAS — Reporte de ventas"
ws["A1"].font = Font(bold=True, size=16, color="1B5E20")
ws["A2"] = "Generado por el área comercial"
for i, h in enumerate(["Producto", "Cantidad", "Precio", "Subtotal", "IVA"], 1):
    c = ws.cell(row=4, column=i, value=h)
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="1B5E20")
    c.alignment = Alignment(horizontal="center")
ws["C5"] = 1000; ws["C5"].number_format = '"$" #,##0'
ws["B5"] = 1; ws["A5"] = "(fila de ejemplo de la plantilla)"; ws["D5"] = "=B5*C5"; ws["E5"] = "=D5*0.19"
ws.column_dimensions["A"].width = 30
ruta_xtpl = SALIDA / "plantilla_ventas.xlsx"
wb.save(str(ruta_xtpl))
print(D.inspeccionar(str(ruta_xtpl)))

paso("H. spreadsheet_create sobre la plantilla (desde A5) con fórmulas, totales, formatos y gráfico")
hoja = {"hojas": [{
    "nombre": "Ventas", "desde": "A5",
    "encabezados": ["Producto", "Cantidad", "Precio", "Subtotal", "IVA"],
    "filas": [["Abono orgánico", "120", "35.500", "=B5*C5", "=D5*0,19"],
              ["Semilla maíz", "40", "870.000", "=B6*C6", "=D6*0.19"],
              ["Riego por goteo", "3", "1.250.000,50", "=B7*C7", "=D7*0.19"]],
    "totales": True,
    "formatos": {"C": '"$" #,##0', "D": '"$" #,##0', "E": '"$" #,##0'},
    "grafico": {"tipo": "barras", "titulo": "Subtotal por producto", "columnas": ["Subtotal"], "categorias": "Producto"},
}]}
r = D.crear_hoja("ventas.xlsx", hoja, plantilla=str(ruta_xtpl), carpeta=str(SALIDA))
print(r.describir())
wb2 = load_workbook(r.ruta)
ws2 = wb2["Ventas"]
for row in ws2.iter_rows(min_row=1, max_row=ws2.max_row, values_only=True):
    print("  ", row)
print("  fuente A5 (dato nuevo) vs A4 (encabezado plantilla):", ws2["A5"].font.name, ws2["A5"].font.b, "|", ws2["A4"].font.b, ws2["A4"].fill.fgColor.rgb)
print("  formato C5:", ws2["C5"].number_format, " charts:", len(ws2._charts))
wb2.close()

paso("H2. Mismo libro a PDF con Excel")
r = D.crear_hoja("ventas.pdf", hoja, plantilla=str(ruta_xtpl), carpeta=str(SALIDA))
print(r.describir())
render_pdf(r.ruta, "ventas_xlsx", paginas=1)

paso("H3. Excel sin plantilla, 2 hojas")
r = D.crear_hoja("clientes.xlsx", {"hojas": [
    {"nombre": "Clientes", "encabezados": ["Nombre", "Ciudad", "Deuda"],
     "filas": [["Finca La Esperanza", "Ibagué", "4.350.000"], ["Hacienda El Roble", "Neiva", "0"]], "totales": True,
     "formatos": {"C": '"$" #,##0'}},
    {"nombre": "Resumen", "encabezados": ["Concepto", "Valor"], "filas": [["Cartera total", "=Clientes!C4"]]},
]}, carpeta=str(SALIDA))
print(r.describir())

# --------------------------------------------------------------------------- PowerPoint
paso("I. presentation_create con tema Office 'Ion' → PDF con PowerPoint → render")
pres = {
    "titulo": "Resultados Trimestre 3", "subtitulo": "Agro SAS · Área comercial",
    "diapositivas": [
        {"titulo": "Qué pasó", "puntos": ["Ventas +12,5 %", "Nuevo cliente: Finca La Esperanza", ["Contrato a 12 meses", "Riego por goteo"]],
         "notas": "Recordar mencionar el lote norte."},
        {"titulo": "Cifras", "tabla": {"encabezados": ["Producto", "Ventas"], "filas": [["Abono", "$ 4.260.000"], ["Semilla", "$ 34.800.000"]]}},
        {"titulo": "Ventas por producto", "grafico": {"tipo": "torta", "categorias": ["Abono", "Semilla", "Riego"], "series": [{"nombre": "Ventas", "valores": [4.26, 34.8, 3.75]}]}},
        {"titulo": "Próximos pasos", "texto": "Cerrar dos cotizaciones pendientes antes del 15 de octubre y revisar precios de insumos."},
    ],
}
r = D.crear_presentacion("resultados.pptx", pres, tema="Ion", carpeta=str(SALIDA))
print(r.describir())
from pptx import Presentation
prs = Presentation(r.ruta)
print("  tamaño:", prs.slide_width, "x", prs.slide_height, "| diapositivas:", len(prs.slides))
print("  layouts:", [l.name for l in prs.slide_layouts][:6])
for i, s in enumerate(prs.slides, 1):
    print(f"   {i}. {[sh.shape_type for sh in s.shapes]} notas={s.has_notes_slide and s.notes_slide.notes_text_frame.text[:30]!r}")
ruta_pptx = r.ruta
r = D.convertir_a_pdf(ruta_pptx, carpeta=str(SALIDA))
print(r.describir())
render_pdf(r.ruta, "resultados_pptx", paginas=4)

paso("J. La presentación anterior como PLANTILLA (tema Ion heredado) + reemplazos + diapositivas nuevas")
r = D.crear_presentacion("propuesta.pptx", {"titulo": "Propuesta comercial 2027", "diapositivas": [
    {"titulo": "Alcance", "puntos": ["Suministro de insumos", "Asistencia técnica mensual"]}]},
    plantilla=ruta_pptx, carpeta=str(SALIDA))
print(r.describir())
prs = Presentation(r.ruta)
print("  diapositivas:", len(prs.slides), "| tamaño:", prs.slide_width)

# --------------------------------------------------------------------------- edición
paso("K. document_edit: Word (reemplazar + agregar después de un párrafo), en copia")
r = D.editar_documento(str(SALIDA / "informe.docx"), {
    "reemplazos": {"Trimestre 3": "Trimestre 4"},
    "agregar": {"despues_de": "Cifras", "bloques": [{"tipo": "parrafo", "texto": "Nota: cifras preliminares sujetas a cierre contable."}]},
}, carpeta=str(SALIDA))
print(r.describir())
d = Document(r.ruta)
print("  ", [p.text[:45] for p in d.paragraphs if p.text][:8])

# --------------------------------------------------------------------------- rechazos
paso("L. Rechazos: relleno, vacío, extensión rara, carpeta vedada")
for nombre, cont in [("x.docx", {"titulo": "Informe", "bloques": [{"tipo": "parrafo", "texto": "[insertar aquí el resumen]"}]}),
                     ("x.docx", {"titulo": "Solo título", "bloques": []}),
                     ("x.exe", "hola")]:
    try:
        D.crear_documento(nombre, cont, carpeta=str(SALIDA)); print("  ¡aceptó!", nombre)
    except D.DocumentoRechazado as e:
        print("  rechazado:", nombre, "→", str(e)[:90])
try:
    D.crear_archivo("script.py", "print('hola')\n", carpeta=str(SALIDA)); print("  file_create .py OK")
except D.DocumentoRechazado as e:
    print("  file_create:", e)

print("\nArchivos en salida:")
for f in sorted(SALIDA.iterdir()):
    print("  ", f.name, f.stat().st_size)
